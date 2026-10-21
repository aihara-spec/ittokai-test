import logging
from sqlalchemy.exc import IntegrityError
import os, io, csv, json, base64, hashlib, hmac, secrets, urllib.request, urllib.parse, urllib.error
from PIL import Image as PILImage
from openpyxl import Workbook, load_workbook
from pypdf import PdfReader, PdfWriter
from datetime import datetime, timedelta, timezone, date, time
from typing import Optional

from fastapi import FastAPI, Request, HTTPException, UploadFile, File, Form
from fastapi.responses import JSONResponse, FileResponse, Response, StreamingResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel
from sqlalchemy import create_engine, Column, Integer, String, Text, DateTime, Float, LargeBinary, ForeignKey, UniqueConstraint, select, func
from sqlalchemy.orm import declarative_base, sessionmaker, relationship
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image as RLImage, PageBreak
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors

ROOT = os.path.dirname(os.path.abspath(__file__))
STATIC = os.path.join(ROOT, 'static')
DATA_DIR = os.path.join(ROOT, 'data')
os.makedirs(DATA_DIR, exist_ok=True)
JST = timezone(timedelta(hours=9))

# Supabase/PostgreSQL: set DATABASE_URL=postgresql+psycopg://...
DATABASE_URL = os.environ.get('DATABASE_URL', f"sqlite:///{os.path.join(DATA_DIR, 'ittokai_v20.db')}")
if DATABASE_URL.startswith('postgres://'):
    DATABASE_URL = DATABASE_URL.replace('postgres://', 'postgresql+psycopg://', 1)
if DATABASE_URL.startswith('postgresql://') and '+psycopg' not in DATABASE_URL:
    DATABASE_URL = DATABASE_URL.replace('postgresql://', 'postgresql+psycopg://', 1)

engine_kwargs = {'future': True, 'pool_pre_ping': True}
if DATABASE_URL.startswith('sqlite:'):
    engine_kwargs['connect_args'] = {'check_same_thread': False}
engine = create_engine(DATABASE_URL, **engine_kwargs)
SessionLocal = sessionmaker(bind=engine, autoflush=False, autocommit=False, expire_on_commit=False)
Base = declarative_base()


def now_dt(): return datetime.now(JST).replace(tzinfo=None)
def now_iso(): return datetime.now(JST).isoformat(timespec='seconds')
def sha256_bytes(b: bytes): return hashlib.sha256(b).hexdigest()
def sha256_text(s: str): return hashlib.sha256((s or '').encode('utf-8')).hexdigest()

def add_years(d: date, years:int):
    try: return d.replace(year=d.year+years)
    except ValueError: return d.replace(month=2,day=28,year=d.year+years)

def safe_date(v):
    try: return date.fromisoformat(str(v)[:10])
    except Exception: return None

def month_key(d: date): return f'{d.year:04d}-{d.month:02d}'



def hash_password(pw, salt=None):
    salt = salt or secrets.token_bytes(16)
    dk = hashlib.pbkdf2_hmac('sha256', pw.encode(), salt, 160000)
    return base64.b64encode(salt).decode()+':' + base64.b64encode(dk).decode()

def verify_password(pw, stored):
    try:
        s,h=stored.split(':',1); salt=base64.b64decode(s); expected=base64.b64decode(h)
        got=hashlib.pbkdf2_hmac('sha256', pw.encode(), salt, 160000)
        return hmac.compare_digest(got, expected)
    except Exception: return False

# ------------------ Models ------------------
class Facility(Base):
    __tablename__='facilities'
    id=Column(Integer, primary_key=True); code=Column(String(64), unique=True, index=True); name=Column(String(200), nullable=False)
    corporation=Column(String(200), default='社会福祉法人 一燈会'); address=Column(String(400), default=''); phone=Column(String(80), default=''); active=Column(Integer, default=1)

class Department(Base):
    __tablename__='departments'
    id=Column(Integer, primary_key=True); facility_id=Column(Integer, ForeignKey('facilities.id'), index=True); name=Column(String(160)); code=Column(String(64)); active=Column(Integer, default=1)
    __table_args__=(UniqueConstraint('facility_id','code', name='uq_department_code'),)

class User(Base):
    __tablename__='users'
    id=Column(Integer, primary_key=True); employee_code=Column(String(64), default=''); name=Column(String(160)); email=Column(String(240), unique=True, index=True); phone=Column(String(80), default='')
    password_hash=Column(Text); role=Column(String(40), index=True); facility_id=Column(Integer, ForeignKey('facilities.id'), nullable=True, index=True); department_id=Column(Integer, ForeignKey('departments.id'), nullable=True)
    line_user_id=Column(String(120), default=''); active=Column(Integer, default=1); created_at=Column(DateTime, default=now_dt)

class Resident(Base):
    __tablename__='residents'
    id=Column(Integer, primary_key=True); resident_code=Column(String(64), unique=True, index=True); facility_id=Column(Integer, ForeignKey('facilities.id'), index=True); name=Column(String(160)); service_type=Column(String(80), default='入所'); active=Column(Integer, default=1)

class FamilyLink(Base):
    __tablename__='family_links'
    id=Column(Integer, primary_key=True); user_id=Column(Integer, ForeignKey('users.id'), index=True); resident_id=Column(Integer, ForeignKey('residents.id'), index=True); relation=Column(String(80), default='家族'); is_primary=Column(Integer, default=0)
    __table_args__=(UniqueConstraint('user_id','resident_id', name='uq_family_resident'),)

class SessionToken(Base):
    __tablename__='sessions'
    token=Column(String(128), primary_key=True); user_id=Column(Integer, ForeignKey('users.id'), index=True); created_at=Column(DateTime, default=now_dt); expires_at=Column(DateTime)

class AuditLog(Base):
    __tablename__='audit_logs'
    id=Column(Integer, primary_key=True); facility_id=Column(Integer, nullable=True, index=True); actor_id=Column(Integer, nullable=True, index=True); actor_name=Column(String(160)); module=Column(String(80), index=True); action=Column(String(120)); entity_type=Column(String(80)); entity_id=Column(String(80)); metadata_json=Column(Text, default='{}'); ip=Column(String(100), default=''); user_agent=Column(Text, default=''); created_at=Column(DateTime, default=now_dt, index=True)

class Notification(Base):
    __tablename__='notifications'
    id=Column(Integer, primary_key=True); user_id=Column(Integer, ForeignKey('users.id'), index=True); title=Column(String(240)); body=Column(Text); module=Column(String(80)); entity_id=Column(String(80), default=''); read_at=Column(DateTime, nullable=True); created_at=Column(DateTime, default=now_dt)

class NotificationEvent(Base):
    __tablename__='notification_events'
    id=Column(Integer, primary_key=True); event_key=Column(String(240), unique=True, index=True); event_type=Column(String(100), index=True); module=Column(String(80), index=True); entity_id=Column(String(80), default='', index=True); facility_id=Column(Integer, nullable=True, index=True); actor_id=Column(Integer, nullable=True); title=Column(String(240)); body=Column(Text); metadata_json=Column(Text, default='{}'); created_at=Column(DateTime, default=now_dt, index=True)

class NotificationRule(Base):
    __tablename__='notification_rules'
    id=Column(Integer, primary_key=True); event_type=Column(String(100), index=True); recipient_role=Column(String(40), default='*', index=True); app_enabled=Column(Integer, default=1); line_enabled=Column(Integer, default=1); repeat_after_hours=Column(Integer, default=24); deadline_before_hours=Column(Integer, default=72); active=Column(Integer, default=1); updated_at=Column(DateTime, default=now_dt)
    __table_args__=(UniqueConstraint('event_type','recipient_role', name='uq_notification_rule_event_role'),)

class NotificationDelivery(Base):
    __tablename__='notification_deliveries'
    id=Column(Integer, primary_key=True); event_id=Column(Integer, ForeignKey('notification_events.id'), index=True); user_id=Column(Integer, ForeignKey('users.id'), index=True); channel=Column(String(20), index=True); status=Column(String(40), index=True); attempts=Column(Integer, default=0); deep_link=Column(Text, default=''); last_error=Column(Text, default=''); sent_at=Column(DateTime, nullable=True); created_at=Column(DateTime, default=now_dt); updated_at=Column(DateTime, default=now_dt)
    __table_args__=(UniqueConstraint('event_id','user_id','channel', name='uq_notification_delivery_event_user_channel'),)

class LineAccountLink(Base):
    __tablename__='line_account_links'
    id=Column(Integer, primary_key=True); user_id=Column(Integer, ForeignKey('users.id'), unique=True, index=True); line_user_id=Column(String(120), unique=True, index=True); verified=Column(Integer, default=1); active=Column(Integer, default=1); linked_at=Column(DateTime, default=now_dt); updated_at=Column(DateTime, default=now_dt)

class LineLinkAttempt(Base):
    __tablename__='line_link_attempts'
    id=Column(Integer, primary_key=True)
    user_id=Column(Integer, ForeignKey('users.id'), index=True)
    code=Column(String(80)); authenticated=Column(Integer, default=0)
    created_at=Column(DateTime, default=now_dt)

class CustomerOrder(Base):
    __tablename__='customer_orders'
    id=Column(Integer, primary_key=True); order_no=Column(String(80), unique=True, index=True); family_user_id=Column(Integer, ForeignKey('users.id'), index=True); facility_id=Column(Integer, ForeignKey('facilities.id'), index=True); resident_id=Column(Integer, ForeignKey('residents.id'), nullable=True, index=True); merchant_type=Column(String(20), index=True); supplier_id=Column(Integer, nullable=True, index=True); store_id=Column(Integer, nullable=True, index=True); status=Column(String(40), default='submitted', index=True); total_amount=Column(Float, default=0); gift_message=Column(Text, default=''); delivery_date=Column(String(10), default=''); created_at=Column(DateTime, default=now_dt); accepted_at=Column(DateTime, nullable=True); shipped_at=Column(DateTime, nullable=True); delivered_at=Column(DateTime, nullable=True); cancelled_at=Column(DateTime, nullable=True); updated_at=Column(DateTime, default=now_dt)

class CustomerOrderItem(Base):
    __tablename__='customer_order_items'
    id=Column(Integer, primary_key=True); order_id=Column(Integer, ForeignKey('customer_orders.id'), index=True); product_kind=Column(String(20)); product_id=Column(Integer, index=True); sku=Column(String(80)); name=Column(String(240)); qty=Column(Float, default=1); unit_price=Column(Float, default=0)

class Message(Base):
    __tablename__='messages'
    id=Column(Integer, primary_key=True); facility_id=Column(Integer, index=True); family_user_id=Column(Integer, index=True); sender_id=Column(Integer, nullable=True); sender_name=Column(String(160)); sender_role=Column(String(40)); body=Column(Text); created_at=Column(DateTime, default=now_dt); read_at=Column(DateTime, nullable=True)

class News(Base):
    __tablename__='news'
    id=Column(Integer, primary_key=True); facility_id=Column(Integer, index=True); category=Column(String(80)); title=Column(String(240)); body=Column(Text); important=Column(Integer, default=0); image_url=Column(Text, default=''); created_at=Column(DateTime, default=now_dt)

# Documents v2
class Document(Base):
    __tablename__='documents'
    id=Column(Integer, primary_key=True); facility_id=Column(Integer, index=True); title=Column(String(300)); category=Column(String(100)); body=Column(Text, default=''); version=Column(Integer, default=1); status=Column(String(40), default='draft', index=True); due_at=Column(DateTime, nullable=True)
    requires_signature=Column(Integer, default=1); required_checks_json=Column(Text, default='["内容を確認しました"]'); file_name=Column(String(260), default=''); file_blob=Column(LargeBinary, nullable=True); original_hash=Column(String(64), default=''); created_by=Column(Integer); created_at=Column(DateTime, default=now_dt); sent_at=Column(DateTime, nullable=True); voided_at=Column(DateTime, nullable=True); previous_document_id=Column(Integer, nullable=True)

class DocumentRecipient(Base):
    __tablename__='document_recipients'
    id=Column(Integer, primary_key=True); document_id=Column(Integer, ForeignKey('documents.id'), index=True); user_id=Column(Integer, ForeignKey('users.id'), index=True); signer_role=Column(String(80), default='家族'); sign_order=Column(Integer, default=1); viewed_at=Column(DateTime, nullable=True); consent_at=Column(DateTime, nullable=True); signed_at=Column(DateTime, nullable=True); status=Column(String(40), default='sent')
    __table_args__=(UniqueConstraint('document_id','user_id', name='uq_document_recipient'),)

class DocumentResidentLink(Base):
    __tablename__='document_resident_links'
    id=Column(Integer, primary_key=True); document_id=Column(Integer, ForeignKey('documents.id'), index=True); resident_id=Column(Integer, ForeignKey('residents.id'), index=True)
    __table_args__=(UniqueConstraint('document_id','resident_id', name='uq_document_resident'),)

class DocumentSignature(Base):
    __tablename__='document_signatures'
    id=Column(Integer, primary_key=True); document_id=Column(Integer, index=True); recipient_id=Column(Integer, index=True); user_id=Column(Integer, index=True); signer_name=Column(String(160)); signature_data=Column(Text); signature_hash=Column(String(64)); ip=Column(String(100)); user_agent=Column(Text); created_at=Column(DateTime, default=now_dt)

# Staff workflow
class WorkflowRequest(Base):
    __tablename__='workflow_requests'
    id=Column(Integer, primary_key=True); facility_id=Column(Integer, index=True); department_id=Column(Integer, nullable=True); applicant_id=Column(Integer, index=True); request_type=Column(String(40), index=True); status=Column(String(40), default='pending_manager', index=True); request_date=Column(String(20)); start_time=Column(String(10), default=''); end_time=Column(String(10), default=''); payload_json=Column(Text, default='{}'); applicant_comment=Column(Text, default=''); manager_id=Column(Integer, nullable=True); admin_id=Column(Integer, nullable=True); created_at=Column(DateTime, default=now_dt); updated_at=Column(DateTime, default=now_dt); exported_at=Column(DateTime, nullable=True)

class WorkflowApproval(Base):
    __tablename__='workflow_approvals'
    id=Column(Integer, primary_key=True); request_id=Column(Integer, index=True); approver_id=Column(Integer); approver_role=Column(String(40)); action=Column(String(40)); comment=Column(Text, default=''); created_at=Column(DateTime, default=now_dt)

class LeaveBalance(Base):
    __tablename__='leave_balances'
    id=Column(Integer, primary_key=True); user_id=Column(Integer, unique=True, index=True); fiscal_year=Column(String(10), default='2026'); granted_days=Column(Float, default=20); remaining_days=Column(Float, default=20); updated_at=Column(DateTime, default=now_dt)

# Visit reservations
class LeaveGrant(Base):
    __tablename__='leave_grants'
    id=Column(Integer, primary_key=True); user_id=Column(Integer, ForeignKey('users.id'), index=True); grant_date=Column(String(10), index=True); granted_days=Column(Float, default=0); remaining_days=Column(Float, default=0); taken_days=Column(Float, default=0); expires_on=Column(String(10), index=True); source=Column(String(40), default='manual'); note=Column(Text, default=''); created_at=Column(DateTime, default=now_dt)
    __table_args__=(UniqueConstraint('user_id','grant_date', name='uq_leave_grant_date'),)

class LeaveUsage(Base):
    __tablename__='leave_usages'
    id=Column(Integer, primary_key=True); request_id=Column(Integer, unique=True, index=True); user_id=Column(Integer, index=True); used_days=Column(Float, default=0); allocation_json=Column(Text, default='[]'); created_at=Column(DateTime, default=now_dt)

class WorkTimeRecord(Base):
    __tablename__='work_time_records'
    id=Column(Integer, primary_key=True); user_id=Column(Integer, ForeignKey('users.id'), index=True); work_date=Column(String(10), index=True); overtime_minutes=Column(Integer, default=0); holiday_minutes=Column(Integer, default=0); source=Column(String(40), default='chronos'); imported_at=Column(DateTime, default=now_dt)
    __table_args__=(UniqueConstraint('user_id','work_date','source', name='uq_work_time_source_date'),)

class LaborPolicy(Base):
    __tablename__='labor_policies'
    id=Column(Integer, primary_key=True); facility_id=Column(Integer, unique=True, index=True); agreement_year_start=Column(String(5), default='04-01'); monthly_overtime_limit=Column(Float, default=45); annual_overtime_limit=Column(Float, default=360); special_annual_limit=Column(Float, default=720); combined_month_limit=Column(Float, default=100); rolling_average_limit=Column(Float, default=80); over45_months_limit=Column(Integer, default=6); annual_leave_required_days=Column(Float, default=5); time_leave_max_days=Column(Float, default=5); updated_at=Column(DateTime, default=now_dt)

class ProductImage(Base):
    __tablename__='product_images'
    id=Column(Integer, primary_key=True); product_id=Column(Integer, ForeignKey('products.id'), index=True); mime_type=Column(String(80), default='image/jpeg'); image_blob=Column(LargeBinary); image_hash=Column(String(64), default=''); is_primary=Column(Integer, default=1); created_at=Column(DateTime, default=now_dt)

class VisitRule(Base):
    __tablename__='visit_rules'
    id=Column(Integer, primary_key=True); facility_id=Column(Integer, index=True); place_name=Column(String(120), default='面会室'); weekday=Column(Integer); start_time=Column(String(5)); end_time=Column(String(5)); slot_minutes=Column(Integer, default=30); capacity=Column(Integer, default=1); max_visitors=Column(Integer, default=4); booking_days_ahead=Column(Integer, default=30); cutoff_hours=Column(Integer, default=24); approval_mode=Column(String(20), default='auto'); active=Column(Integer, default=1)

class VisitBlock(Base):
    __tablename__='visit_blocks'
    id=Column(Integer, primary_key=True); facility_id=Column(Integer, index=True); block_date=Column(String(10), index=True); start_time=Column(String(5), default='00:00'); end_time=Column(String(5), default='23:59'); reason=Column(String(240), default='')

class VisitReservation(Base):
    __tablename__='visit_reservations'
    id=Column(Integer, primary_key=True); facility_id=Column(Integer, index=True); resident_id=Column(Integer, index=True); family_user_id=Column(Integer, index=True); visit_date=Column(String(10), index=True); start_time=Column(String(5)); end_time=Column(String(5)); place_name=Column(String(120)); visitor_count=Column(Integer, default=1); representative=Column(String(160)); phone=Column(String(80), default=''); notes=Column(Text, default=''); status=Column(String(40), default='confirmed'); checked_in_at=Column(DateTime, nullable=True); created_at=Column(DateTime, default=now_dt)

# CareDo procurement
class Supplier(Base):
    __tablename__='suppliers'
    id=Column(Integer, primary_key=True); supplier_type=Column(String(30), index=True); name=Column(String(200)); code=Column(String(64), unique=True); email=Column(String(240), default=''); phone=Column(String(80), default=''); site_url=Column(Text, default=''); order_url=Column(Text, default=''); active=Column(Integer, default=1)

class Product(Base):
    __tablename__='products'
    id=Column(Integer, primary_key=True); supplier_id=Column(Integer, ForeignKey('suppliers.id'), index=True); sku=Column(String(80), unique=True, index=True); name=Column(String(240)); specification=Column(String(240), default=''); category=Column(String(100)); unit=Column(String(40), default='個'); pack_size=Column(Integer, default=1); tax_rate=Column(Float, default=10.0); standard_price=Column(Float, default=0); min_order_qty=Column(Integer, default=1); lead_time_days=Column(Integer, default=2); substitute_sku=Column(String(80), default=''); inventory_managed=Column(Integer, default=1); active=Column(Integer, default=1)

class ProductMeta(Base):
    __tablename__='product_meta'
    id=Column(Integer, primary_key=True); product_id=Column(Integer, ForeignKey('products.id'), unique=True, index=True); short_message=Column(String(300), default=''); description=Column(Text, default=''); purchase_price=Column(Float, default=0); sale_price=Column(Float, default=0); recommended=Column(Integer, default=0); sale_start=Column(String(10), default=''); sale_end=Column(String(10), default=''); display_order=Column(Integer, default=0); updated_at=Column(DateTime, default=now_dt)

class FacilityStockPolicy(Base):
    __tablename__='facility_stock_policies'
    id=Column(Integer, primary_key=True); facility_id=Column(Integer, index=True); product_id=Column(Integer, index=True); reorder_qty=Column(Float, default=0); max_stock=Column(Float, default=0); auto_order_mode=Column(String(20), default='approval'); last_ordered_at=Column(DateTime, nullable=True); last_received_at=Column(DateTime, nullable=True); updated_at=Column(DateTime, default=now_dt)
    __table_args__=(UniqueConstraint('facility_id','product_id', name='uq_facility_stock_policy'),)

class AutoOrderEvent(Base):
    __tablename__='auto_order_events'
    id=Column(Integer, primary_key=True); facility_id=Column(Integer, index=True); product_id=Column(Integer, index=True); order_id=Column(Integer, nullable=True, index=True); mode=Column(String(20), default='approval'); status=Column(String(40), default='pending_approval'); suggested_qty=Column(Float, default=0); on_hand_at_trigger=Column(Float, default=0); reorder_point=Column(Float, default=0); note=Column(Text, default=''); created_at=Column(DateTime, default=now_dt); resolved_at=Column(DateTime, nullable=True)

class PartnerStore(Base):
    __tablename__='partner_stores'
    id=Column(Integer, primary_key=True); owner_user_id=Column(Integer, ForeignKey('users.id'), index=True); company_name=Column(String(200), default=''); store_name=Column(String(200)); category=Column(String(100), default='その他'); image_url=Column(Text, default=''); address=Column(String(400), default=''); phone=Column(String(80), default=''); hours=Column(String(160), default=''); closed_days=Column(String(160), default=''); delivery_area=Column(String(300), default=''); min_order_amount=Column(Float, default=0); order_cutoff=Column(String(80), default=''); delivery_methods=Column(String(240), default='施設配送'); message=Column(Text, default=''); active=Column(Integer, default=1); created_at=Column(DateTime, default=now_dt); updated_at=Column(DateTime, default=now_dt)

class PartnerProduct(Base):
    __tablename__='partner_products'
    id=Column(Integer, primary_key=True); store_id=Column(Integer, ForeignKey('partner_stores.id'), index=True); sku=Column(String(80), index=True); name=Column(String(240)); short_message=Column(String(300), default=''); description=Column(Text, default=''); category=Column(String(100), default='ギフト'); specification=Column(String(240), default=''); unit=Column(String(40), default='個'); price=Column(Float, default=0); image_url=Column(Text, default=''); recommended=Column(Integer, default=0); active=Column(Integer, default=1); sale_start=Column(String(10), default=''); sale_end=Column(String(10), default=''); created_at=Column(DateTime, default=now_dt); updated_at=Column(DateTime, default=now_dt)
    __table_args__=(UniqueConstraint('store_id','sku', name='uq_partner_store_sku'),)

class FacilityPrice(Base):
    __tablename__='facility_prices'
    id=Column(Integer, primary_key=True); facility_id=Column(Integer, index=True); product_id=Column(Integer, index=True); purchase_price=Column(Float, default=0); sale_price=Column(Float, default=0); effective_from=Column(String(10), default=''); effective_to=Column(String(10), default='')
    __table_args__=(UniqueConstraint('facility_id','product_id', name='uq_facility_product_price'),)

class Inventory(Base):
    __tablename__='inventory'
    id=Column(Integer, primary_key=True); facility_id=Column(Integer, index=True); product_id=Column(Integer, index=True); location=Column(String(120), default='主倉庫'); on_hand=Column(Float, default=0); safety_stock=Column(Float, default=0); reorder_point=Column(Float, default=0); target_stock=Column(Float, default=0); avg_daily_usage=Column(Float, default=0); updated_at=Column(DateTime, default=now_dt)
    __table_args__=(UniqueConstraint('facility_id','product_id','location', name='uq_inventory_location'),)

class InventoryTxn(Base):
    __tablename__='inventory_txns'
    id=Column(Integer, primary_key=True); facility_id=Column(Integer, index=True); product_id=Column(Integer, index=True); txn_type=Column(String(40)); qty=Column(Float); reference_type=Column(String(60), default=''); reference_id=Column(String(80), default=''); note=Column(Text, default=''); actor_id=Column(Integer, nullable=True); created_at=Column(DateTime, default=now_dt)

class PurchaseOrder(Base):
    __tablename__='purchase_orders'
    id=Column(Integer, primary_key=True); order_no=Column(String(80), unique=True, index=True); facility_id=Column(Integer, index=True); supplier_id=Column(Integer, index=True); status=Column(String(40), default='draft'); total_amount=Column(Float, default=0); approved_by=Column(Integer, nullable=True); ordered_at=Column(DateTime, nullable=True); confirmed_at=Column(DateTime, nullable=True); shipped_at=Column(DateTime, nullable=True); received_at=Column(DateTime, nullable=True); created_at=Column(DateTime, default=now_dt)

class PurchaseOrderItem(Base):
    __tablename__='purchase_order_items'
    id=Column(Integer, primary_key=True); order_id=Column(Integer, index=True); product_id=Column(Integer, index=True); qty=Column(Float); unit_price=Column(Float); received_qty=Column(Float, default=0)

class Invoice(Base):
    __tablename__='invoices'
    id=Column(Integer, primary_key=True); invoice_no=Column(String(80), unique=True, index=True); supplier_id=Column(Integer, index=True); facility_id=Column(Integer, index=True); order_id=Column(Integer, index=True); invoice_date=Column(String(10)); total_amount=Column(Float, default=0); pdf_blob=Column(LargeBinary, nullable=True); status=Column(String(40), default='received'); match_status=Column(String(40), default='unchecked'); difference_json=Column(Text, default='{}'); created_at=Column(DateTime, default=now_dt)

class InvoiceItem(Base):
    __tablename__='invoice_items'
    id=Column(Integer, primary_key=True); invoice_id=Column(Integer, index=True); product_id=Column(Integer, index=True); qty=Column(Float); unit_price=Column(Float); amount=Column(Float)

class IntegrationProfile(Base):
    __tablename__='integration_profiles'
    id=Column(Integer, primary_key=True); system_name=Column(String(80), unique=True); display_name=Column(String(160)); mode=Column(String(40), default='CSV'); direction=Column(String(40), default='bidirectional'); status=Column(String(40), default='template_ready'); notes=Column(Text, default='')

# ------------------ Helpers ------------------
ROLES = ['family','staff','middle_manager','admin','hq','caredo','sponsor']
ROLE_LABEL = {'family':'家族','staff':'一般職員','middle_manager':'中間管理職','admin':'施設管理者','hq':'法人本部','caredo':'ケア・ドゥ','sponsor':'外部協賛企業'}


def client_ip(request: Request):
    return (request.headers.get('x-forwarded-for') or request.client.host if request.client else '')[:100]

def get_token(request: Request):
    auth=request.headers.get('authorization','')
    if auth.lower().startswith('bearer '): return auth[7:].strip()
    return request.headers.get('x-session-token','')

def current_user(request: Request, s):
    token=get_token(request)
    if not token: raise HTTPException(401,'ログインが必要です')
    st=s.get(SessionToken, token)
    if not st or st.expires_at < now_dt(): raise HTTPException(401,'セッションが無効です')
    u=s.get(User, st.user_id)
    if not u or not u.active: raise HTTPException(401,'ユーザーが無効です')
    return u

def require_roles(u, *roles):
    if u.role not in roles: raise HTTPException(403,'権限がありません')

def facility_scope(u, facility_id):
    if u.role=='hq': return
    if u.role in ('caredo','sponsor'): return
    if u.facility_id != facility_id: raise HTTPException(403,'他施設のデータにはアクセスできません')

def audit(s, request, u, module, action, entity_type='', entity_id='', meta=None, facility_id=None):
    s.add(AuditLog(facility_id=facility_id or getattr(u,'facility_id',None), actor_id=getattr(u,'id',None), actor_name=getattr(u,'name','system'), module=module, action=action, entity_type=entity_type, entity_id=str(entity_id or ''), metadata_json=json.dumps(meta or {}, ensure_ascii=False), ip=client_ip(request), user_agent=request.headers.get('user-agent','')[:1000]))

APP_BASE_URL=(os.environ.get('APP_BASE_URL') or 'http://localhost:8000').rstrip('/')
LIFF_BASE_URL=(os.environ.get('LIFF_BASE_URL') or '').rstrip('/')
LINE_API_BASE='https://api.line.me'

def _view_for_module(module,role=''):
    if module=='documents': return 'dashboard' if role=='hq' else 'documents'
    if module=='workflow': return 'requests'
    if module=='visits': return 'visits'
    if module in ('procurement','inventory'): return 'procurement'
    if module=='shop': return 'partner' if role=='sponsor' else ('procurement' if role=='caredo' else 'shop')
    if module=='chat': return 'talk'
    if module=='news': return 'news'
    return 'home'

def deep_link_for_user(u,module,entity_id=''):
    params={'view':_view_for_module(module,getattr(u,'role',''))}
    if entity_id not in ('',None): params['entity']=str(entity_id)
    qs=urllib.parse.urlencode(params)
    base=LIFF_BASE_URL or APP_BASE_URL
    return f'{base}?{qs}'

def line_user_id_for(s,user_id):
    link=s.query(LineAccountLink).filter_by(user_id=user_id,active=1,verified=1).first()
    if link and link.line_user_id: return link.line_user_id
    u=s.get(User,user_id)
    return (u.line_user_id or '') if u else ''

def _line_push(line_user_id,title,body,deep_link=''):
    if not line_user_id: return False,'LINE未連携'
    if os.environ.get('LINE_MOCK_MODE','')=='1': return True,''
    token=os.environ.get('LINE_CHANNEL_ACCESS_TOKEN','')
    if not token: return False,'LINE_CHANNEL_ACCESS_TOKEN未設定'
    text=(body or '').strip() or 'ITTOKAIに新しいお知らせがあります。'
    if deep_link:
        message={'type':'template','altText':title[:400], 'template':{'type':'buttons','text':text[:160], 'title':title[:40], 'actions':[{'type':'uri','label':'ITTOKAIで確認','uri':deep_link[:1000]}]}}
    else:
        message={'type':'text','text':f'{title}\n{text}'[:5000]}
    payload=json.dumps({'to':line_user_id,'messages':[message]},ensure_ascii=False).encode('utf-8')
    try:
        req=urllib.request.Request(LINE_API_BASE+'/v2/bot/message/push',data=payload,headers={'Authorization':'Bearer '+token,'Content-Type':'application/json'},method='POST')
        with urllib.request.urlopen(req,timeout=8) as res: res.read()
        return True,''
    except urllib.error.HTTPError as e:
        try: detail=e.read().decode('utf-8','ignore')[:1000]
        except Exception: detail=str(e)
        return False,f'HTTP {e.code}: {detail}'
    except Exception as e:
        return False,str(e)[:1000]

def notification_rule_for(s,event_type,role):
    exact=s.query(NotificationRule).filter_by(event_type=event_type,recipient_role=role).first()
    if exact: return exact
    return s.query(NotificationRule).filter_by(event_type=event_type,recipient_role='*').first()

def emit_notification_event(s,event_type,recipients,title,body,module,entity_id='',facility_id=None,actor_id=None,event_key=None,metadata=None):
    user_ids=[]
    for x in recipients or []:
        try: uid=int(getattr(x,'id',x))
        except Exception: continue
        if uid not in user_ids: user_ids.append(uid)
    event_key=event_key or f'{event_type}:{module}:{entity_id}:{secrets.token_hex(10)}'
    ev=s.query(NotificationEvent).filter_by(event_key=event_key).first()
    if ev: return {'event_id':ev.id,'deduplicated':True,'recipients':user_ids}
    ev=NotificationEvent(event_key=event_key,event_type=event_type,module=module,entity_id=str(entity_id or ''),facility_id=facility_id,actor_id=actor_id,title=title,body=body,metadata_json=json.dumps(metadata or {},ensure_ascii=False)); s.add(ev); s.flush()
    for uid in user_ids:
        u=s.get(User,uid)
        if not u or not u.active: continue
        rule=notification_rule_for(s,event_type,u.role)
        enabled=True if not rule else bool(rule.active)
        app_enabled=enabled and (True if not rule else bool(rule.app_enabled))
        line_enabled=enabled and (True if not rule else bool(rule.line_enabled))
        link=deep_link_for_user(u,module,entity_id)
        if app_enabled:
            s.add(Notification(user_id=uid,title=title,body=body,module=module,entity_id=str(entity_id or '')))
            s.add(NotificationDelivery(event_id=ev.id,user_id=uid,channel='app',status='sent',attempts=1,deep_link=link,sent_at=now_dt(),updated_at=now_dt()))
        if line_enabled:
            # Transactional outbox: persist the intended LINE delivery first.
            # Actual external transmission is performed only after the business
            # transaction has committed, so a rollback cannot leave a ghost LINE alert.
            s.add(NotificationDelivery(event_id=ev.id,user_id=uid,channel='line',status='pending',attempts=0,deep_link=link,last_error='',sent_at=None,updated_at=now_dt()))
    return {'event_id':ev.id,'deduplicated':False,'recipients':user_ids}

def dispatch_pending_line_deliveries(limit=100):
    # Dispatch committed LINE outbox rows. Safe to call repeatedly.
    claimed=[]
    with SessionLocal() as s:
        q=(s.query(NotificationDelivery)
           .filter(NotificationDelivery.channel=='line',NotificationDelivery.status=='pending')
           .order_by(NotificationDelivery.id.asc()))
        try: rows=q.with_for_update(skip_locked=True).limit(limit).all()
        except Exception: rows=q.limit(limit).all()
        for d in rows:
            d.status='sending'; d.attempts=(d.attempts or 0)+1; d.updated_at=now_dt(); claimed.append(d.id)
        if claimed: s.commit()
    if not claimed: return {'processed':0,'sent':0,'failed':0,'skipped_unlinked':0}
    result={'processed':0,'sent':0,'failed':0,'skipped_unlinked':0}
    for did in claimed:
        with SessionLocal() as s:
            d=s.get(NotificationDelivery,did)
            if not d or d.status!='sending': continue
            e=s.get(NotificationEvent,d.event_id); usr=s.get(User,d.user_id)
            if not e or not usr:
                d.status='failed'; d.last_error='通知イベントまたは利用者が見つかりません'; d.updated_at=now_dt(); s.commit(); result['processed']+=1; result['failed']+=1; continue
            ok,err=_line_push(line_user_id_for(s,d.user_id),e.title,e.body,d.deep_link)
            d.status='sent' if ok else ('skipped_unlinked' if err=='LINE未連携' else 'failed')
            d.last_error=err; d.sent_at=now_dt() if ok else None; d.updated_at=now_dt(); s.commit()
            result['processed']+=1; result[d.status]=result.get(d.status,0)+1
    return result

def notify(s, user_id, title, body, module, entity_id=''):
    # Backward-compatible one-recipient notification; business-critical flows use explicit event keys.
    return emit_notification_event(s,f'{module}.generic',[user_id],title,body,module,entity_id,event_key=f'{module}.generic:{entity_id}:{user_id}:{secrets.token_hex(8)}')

def users_by_roles(s,roles,facility_id=None,department_id=None):
    q=s.query(User).filter(User.role.in_(list(roles)),User.active==1)
    if facility_id is not None: q=q.filter(User.facility_id==facility_id)
    if department_id is not None: q=q.filter((User.department_id==department_id)|(User.department_id.is_(None)))
    return q.all()

def ensure_notification_rules(s):
    event_types=['document.requested','document.signed','document.reminder','workflow.submitted','workflow.forwarded','workflow.result','visit.requested','visit.confirmed','visit.cancelled','procurement.order_submitted','procurement.order_status','procurement.invoice_created','inventory.auto_order_candidate','inventory.auto_order_created','shop.order_submitted','shop.order_status','chat.new_message']
    for et in event_types:
        if not s.query(NotificationRule).filter_by(event_type=et,recipient_role='*').first(): s.add(NotificationRule(event_type=et,recipient_role='*',app_enabled=1,line_enabled=1,repeat_after_hours=24,deadline_before_hours=72,active=1))
    s.flush()

def user_json(u):
    return {'id':u.id,'name':u.name,'email':u.email,'role':u.role,'role_label':ROLE_LABEL.get(u.role,u.role),'facility_id':u.facility_id,'department_id':u.department_id,'employee_code':u.employee_code}

def facility_json(f): return {'id':f.id,'code':f.code,'name':f.name,'corporation':f.corporation,'address':f.address,'phone':f.phone}

def dt_iso(v): return v.isoformat(timespec='seconds') if v else None

# ------------------ Seed ------------------
def seed_data():
    with SessionLocal() as s:
        if s.query(Facility).count(): return
        fac=Facility(code='MAISON-NINOMIYA',name='メゾン・二宮',address='神奈川県中郡二宮町',phone='0463-00-0000'); s.add(fac); s.flush()
        deps=[]
        for code,name in [('ADMISSION','入所'),('DAY','デイサービス'),('SHORT','ショートステイ'),('ADMIN','事務')]:
            d=Department(facility_id=fac.id,code=code,name=name); s.add(d); deps.append(d)
        s.flush(); dep=deps[0]
        users=[
            User(name='相原 良太',email='family@ittokai.local',password_hash=hash_password('demo1234'),role='family',facility_id=fac.id),
            User(name='一般職員テスト',email='staff@ittokai.local',employee_code='E0001',password_hash=hash_password('staff1234'),role='staff',facility_id=fac.id,department_id=dep.id),
            User(name='中間管理職テスト',email='manager@ittokai.local',employee_code='M0001',password_hash=hash_password('manager1234'),role='middle_manager',facility_id=fac.id,department_id=dep.id),
            User(name='施設管理者',email='admin@ittokai.local',employee_code='A0001',password_hash=hash_password('admin1234'),role='admin',facility_id=fac.id,department_id=dep.id),
            User(name='法人本部',email='hq@ittokai.local',employee_code='HQ001',password_hash=hash_password('hq1234'),role='hq',facility_id=None),
            User(name='ケア・ドゥ担当',email='caredo@ittokai.local',password_hash=hash_password('caredo1234'),role='caredo',facility_id=None),
            User(name='協賛企業担当',email='sponsor@ittokai.local',password_hash=hash_password('sponsor1234'),role='sponsor',facility_id=None),
        ]
        s.add_all(users); s.flush(); family,staff,manager,admin,hq,caredo,sponsor=users
        res=Resident(resident_code='R0001',facility_id=fac.id,name='相原 太郎',service_type='入所'); s.add(res); s.flush(); s.add(FamilyLink(user_id=family.id,resident_id=res.id,relation='家族',is_primary=1))
        s.add_all([
            News(facility_id=fac.id,category='重要',title='面会についてのお知らせ',body='面会予約はITTOKAIからオンラインで行えます。',important=1,image_url='/assets/cards/news-important.svg'),
            News(facility_id=fac.id,category='施設だより',title='ITTOKAI統合版テスト開始',body='電子書類・職員申請・面会予約・在庫発注の統合テストを開始します。',important=0,image_url='/assets/cards/news-event.svg')
        ])
        s.add_all([
            Message(facility_id=fac.id,family_user_id=family.id,sender_id=admin.id,sender_name='メゾン・二宮',sender_role='admin',body='ITTOKAI統合版へようこそ。',read_at=now_dt()),
            Message(facility_id=fac.id,family_user_id=family.id,sender_id=family.id,sender_name=family.name,sender_role='family',body='よろしくお願いいたします。',read_at=now_dt())
        ])
        # Document sample
        doc=Document(facility_id=fac.id,title='写真・広報掲載に関する同意確認（テスト）',category='同意書',body='本書類は電子書類機能のテストです。内容を確認し、同意のうえ署名してください。',version=1,status='sent',due_at=now_dt()+timedelta(days=7),requires_signature=1,created_by=admin.id,sent_at=now_dt()); doc.original_hash=sha256_text(doc.title+doc.body); s.add(doc); s.flush(); s.add(DocumentRecipient(document_id=doc.id,user_id=family.id,signer_role='家族',sign_order=1,status='sent'))
        # Visit rules Monday-Sunday 14-17 30min
        for wd in range(7): s.add(VisitRule(facility_id=fac.id,place_name='面会室A',weekday=wd,start_time='14:00',end_time='17:00',slot_minutes=30,capacity=2,max_visitors=4,booking_days_ahead=30,cutoff_hours=24,approval_mode='auto'))
        # Suppliers/products/inventory
        sup=Supplier(supplier_type='caredo',name='株式会社ケア・ドゥ',code='CAREDO',site_url='https://caredo-zakka.jp/shop/',order_url='https://caredo-zakka.jp/shop/'); s.add(sup)
        sp=Supplier(supplier_type='sponsor',name='外部協賛企業A',code='SPONSOR-A',site_url='https://example.com',order_url='https://example.com'); s.add(sp); s.flush()
        products=[
            Product(supplier_id=sup.id,sku='CD-OM-M',name='リハビリパンツ M',specification='20枚',category='紙おむつ',unit='ケース',standard_price=2980,min_order_qty=1,lead_time_days=2),
            Product(supplier_id=sup.id,sku='CD-TISSUE',name='BOXティッシュ',specification='150組×5箱',category='日用品',unit='ケース',standard_price=1980,min_order_qty=1,lead_time_days=2),
            Product(supplier_id=sup.id,sku='CD-GLOVE',name='使い捨て手袋 M',specification='100枚',category='衛生',unit='箱',standard_price=780,min_order_qty=5,lead_time_days=2),
        ]; s.add_all(products); s.flush()
        for i,p in enumerate(products):
            s.add(FacilityPrice(facility_id=fac.id,product_id=p.id,purchase_price=p.standard_price*0.8,sale_price=p.standard_price))
            s.add(Inventory(facility_id=fac.id,product_id=p.id,on_hand=[2,8,3][i],safety_stock=[3,4,5][i],reorder_point=[4,5,6][i],target_stock=[8,10,12][i],avg_daily_usage=[0.3,0.2,0.5][i]))
            s.add(ProductMeta(product_id=p.id,short_message=['毎日のケアに。必要な分を施設へ直接お届け。','毎日使うものだから、施設へのまとめ買いがおすすめ。','衛生管理の定番。必要なタイミングで確実に補充。'][i],purchase_price=p.standard_price*0.8,sale_price=p.standard_price,recommended=1 if i<2 else 0,display_order=i+1))
            s.add(FacilityStockPolicy(facility_id=fac.id,product_id=p.id,reorder_qty=[6,10,10][i],max_stock=[10,20,15][i],auto_order_mode='approval'))
        flower=PartnerStore(owner_user_id=sponsor.id,company_name='協賛企業サンプル',store_name='ITTOKAIフラワー',category='花屋',delivery_area='神奈川県内の対象施設',delivery_methods='施設配送',message='離れていても、大切な気持ちをお花に込めて。'); s.add(flower); s.flush()
        s.add_all([
            PartnerProduct(store_id=flower.id,sku='FLOWER-BIRTHDAY',name='季節のお誕生日花束',short_message='離れていても、“おめでとう”の気持ちをお花に込めて。',category='記念日・花',specification='季節のお花／メッセージカード付',price=3980,recommended=1),
            PartnerProduct(store_id=flower.id,sku='FLOWER-ARRANGE',name='季節のフラワーアレンジ',short_message='お部屋を彩る、やさしい季節の贈りもの。',category='記念日・花',specification='施設配送対応',price=4980,recommended=1),
        ])
        cake=PartnerStore(owner_user_id=sponsor.id,company_name='協賛企業サンプル',store_name='ITTOKAI洋菓子店',category='ケーキ屋',delivery_area='対象施設へ配送',delivery_methods='施設配送',message='今年のお誕生日も、施設で一緒にお祝い。'); s.add(cake); s.flush()
        s.add(PartnerProduct(store_id=cake.id,sku='CAKE-BIRTHDAY',name='記念日ホールケーキ',short_message='今年のお誕生日も、施設で一緒にお祝いしませんか？',category='記念日・ケーキ',specification='プレート・ろうそく対応',price=4500,recommended=1))
        # Integration profiles
        s.add_all([
            IntegrationProfile(system_name='honobono',display_name='ほのぼの（介護・施設請求）',mode='CSV/API',direction='bidirectional',status='template_ready',notes='利用者・契約・請求・文書保存のCSV/APIアダプタ'),
            IntegrationProfile(system_name='chronos',display_name='クロノス（勤怠・給与）',mode='CSV/API',direction='bidirectional',status='template_ready',notes='職員・勤務・年休・残業・承認結果'),
            IntegrationProfile(system_name='yayoi',display_name='弥生（ケア・ドゥ販売／請求）',mode='CSV/API',direction='bidirectional',status='template_ready',notes='受注・売上・請求・入金・仕訳データ')
        ])
        s.commit()

# ------------------ PDF ------------------
def pdf_styles():
    try: pdfmetrics.registerFont(UnicodeCIDFont('HeiseiKakuGo-W5'))
    except Exception: pass
    styles=getSampleStyleSheet(); font='HeiseiKakuGo-W5'
    return {
        'title':ParagraphStyle('jp_title',parent=styles['Title'],fontName=font,fontSize=16,leading=22,spaceAfter=8),
        'body':ParagraphStyle('jp_body',parent=styles['BodyText'],fontName=font,fontSize=9.5,leading=15),
        'small':ParagraphStyle('jp_small',parent=styles['BodyText'],fontName=font,fontSize=7.5,leading=11,textColor=colors.HexColor('#555555')),
    }

def make_document_pdf(s, doc: Document, rec: Optional[DocumentRecipient]=None):
    buf=io.BytesIO(); st=pdf_styles(); story=[]
    story.append(Paragraph(doc.title,st['title'])); story.append(Paragraph(f'文書ID: {doc.id} / 版: {doc.version} / 区分: {doc.category}',st['small'])); story.append(Spacer(1,5*mm))
    if doc.file_blob:
        story.append(Paragraph('原本PDFが添付されています。ITTOKAIでは署名証跡ページを付与して保存します。',st['body'])); story.append(Spacer(1,5*mm))
    for para in (doc.body or '').split('\n'):
        story.append(Paragraph(para or '　',st['body']))
    story.append(Spacer(1,8*mm)); story.append(Paragraph('電子署名・監査証跡',st['title']))
    recs=[rec] if rec else s.query(DocumentRecipient).filter_by(document_id=doc.id).order_by(DocumentRecipient.sign_order,DocumentRecipient.id).all()
    for idx,rr in enumerate(recs):
        u=s.get(User,rr.user_id); sig=s.query(DocumentSignature).filter_by(document_id=doc.id,recipient_id=rr.id).order_by(DocumentSignature.id.desc()).first()
        if idx: story.append(Spacer(1,6*mm))
        data=[['署名者',u.name if u else ''],['ステータス',rr.status],['閲覧日時',dt_iso(rr.viewed_at) or '-'],['同意日時',dt_iso(rr.consent_at) or '-'],['署名日時',dt_iso(rr.signed_at) or '-'],['原本SHA-256',doc.original_hash or '-']]
        if sig: data += [['署名SHA-256',sig.signature_hash],['IP',sig.ip],['User-Agent',sig.user_agent[:80]]]
        t=Table(data,colWidths=[35*mm,145*mm]); t.setStyle(TableStyle([('FONTNAME',(0,0),(-1,-1),'HeiseiKakuGo-W5'),('FONTSIZE',(0,0),(-1,-1),7.5),('GRID',(0,0),(-1,-1),0.3,colors.HexColor('#dddddd')),('BACKGROUND',(0,0),(0,-1),colors.HexColor('#fff3e8')),('VALIGN',(0,0),(-1,-1),'TOP'),('PADDING',(0,0),(-1,-1),4)])); story.append(t)
        if sig and sig.signature_data and ',' in sig.signature_data:
            try:
                raw=base64.b64decode(sig.signature_data.split(',',1)[1]); bio=io.BytesIO(raw); PILImage.open(bio).verify(); bio.seek(0); story.append(Spacer(1,5*mm)); story.append(Paragraph('手書きサイン',st['body'])); story.append(RLImage(bio,width=55*mm,height=22*mm))
            except Exception: pass
    SimpleDocTemplate(buf,pagesize=A4,rightMargin=15*mm,leftMargin=15*mm,topMargin=15*mm,bottomMargin=15*mm).build(story); buf.seek(0); evidence=buf.getvalue()
    if doc.file_blob and bytes(doc.file_blob).startswith(b'%PDF'):
        try:
            writer=PdfWriter()
            for page in PdfReader(io.BytesIO(bytes(doc.file_blob))).pages: writer.add_page(page)
            for page in PdfReader(io.BytesIO(evidence)).pages: writer.add_page(page)
            merged=io.BytesIO(); writer.write(merged); return merged.getvalue()
        except Exception:
            return evidence
    return evidence

# ------------------ App ------------------
app=FastAPI(title='ITTOKAI v20.7.1 Integrated',version='20.7.1')

@app.middleware('http')
async def no_store_api_responses(request:Request, call_next):
    response=await call_next(request)
    # Process only committed notification outbox rows. If the endpoint failed or
    # rolled back, no outbox row exists and no external LINE message is sent.
    try: dispatch_pending_line_deliveries()
    except Exception: pass  # business response must not fail because LINE is unavailable
    if request.url.path.startswith('/api/'):
        response.headers['Cache-Control']='no-store, no-cache, must-revalidate, max-age=0'
        response.headers['Pragma']='no-cache'
    return response

Base.metadata.create_all(bind=engine); seed_data()
with SessionLocal() as _s:
    ensure_notification_rules(_s); _s.commit()

@app.get('/health')
def health(): return {'ok':True,'version':'20.7.1','database':'postgresql' if DATABASE_URL.startswith('postgresql') else 'sqlite'}

class LoginIn(BaseModel): email:str; password:str
@app.post('/api/login')
def login(inp:LoginIn, request:Request):
    with SessionLocal() as s:
        u=s.query(User).filter(func.lower(User.email)==inp.email.lower()).first()
        if not u or not verify_password(inp.password,u.password_hash): raise HTTPException(401,'IDまたはパスワードが違います')
        token=secrets.token_urlsafe(32); s.add(SessionToken(token=token,user_id=u.id,expires_at=now_dt()+timedelta(days=30))); audit(s,request,u,'auth','login','user',u.id); s.commit(); return {'token':token,'user':user_json(u)}

@app.post('/api/logout')
def logout(request:Request):
    with SessionLocal() as s:
        token=get_token(request); st=s.get(SessionToken,token)
        if st: s.delete(st); s.commit()
    return {'ok':True}

def document_revision_value(s,u):
    """Return a lightweight revision number for document-related changes visible to the user.

    AuditLog IDs are monotonic and every document send/view/sign/remind/reissue action creates a log.
    Using a facility-scoped max ID lets clients detect remote changes without reloading the entire app.
    """
    q=s.query(func.max(AuditLog.id)).filter(AuditLog.module=='documents')
    if u.role!='hq':
        if not u.facility_id:
            return 0
        q=q.filter(AuditLog.facility_id==u.facility_id)
    return int(q.scalar() or 0)

@app.get('/api/documents/revision')
def document_revision(request:Request):
    with SessionLocal() as s:
        u=current_user(request,s)
        if u.role not in ('family','admin','hq'):
            raise HTTPException(403,'電子書類の同期権限がありません')
        return Response(
            content=json.dumps({'revision':document_revision_value(s,u)}, ensure_ascii=False),
            media_type='application/json',
            headers={'Cache-Control':'no-store, no-cache, must-revalidate, max-age=0','Pragma':'no-cache'}
        )

SYNC_ALLOWED = {
    'family': {'documents','chat','visits','news','all'},
    'staff': {'chat','workflow','visits','news','all'},
    'middle_manager': {'chat','workflow','visits','news','procurement','all'},
    'admin': {'documents','chat','workflow','visits','news','procurement','all'},
    'hq': {'documents','workflow','visits','news','procurement','all'},
    'caredo': {'procurement','all'},
    'sponsor': {'procurement','all'},
}

def module_revision_value(s,u,module):
    if module not in SYNC_ALLOWED.get(u.role,set()):
        raise HTTPException(403,'同期権限がありません')
    q=s.query(func.max(AuditLog.id))
    if module!='all': q=q.filter(AuditLog.module==module)
    # HQ is corporation-wide. CareDo sees procurement across all facilities.
    if u.role=='caredo':
        q=q.filter(AuditLog.module=='procurement')
    elif u.role!='hq':
        if not u.facility_id: return 0
        q=q.filter(AuditLog.facility_id==u.facility_id)
    return int(q.scalar() or 0)

@app.get('/api/sync/revision')
def sync_revision(request:Request,module:str='all'):
    with SessionLocal() as s:
        u=current_user(request,s)
        return {'module':module,'revision':module_revision_value(s,u,module)}

def chat_target(s,u,user_id=None):
    require_roles(u,'family','staff','middle_manager','admin')
    target=u.id if u.role=='family' else int(user_id or 0)
    if not target: raise HTTPException(400,'相手を選択してください')
    tu=s.get(User,target)
    if not tu or tu.role!='family' or not tu.active: raise HTTPException(404,'家族アカウントがありません')
    facility_scope(u,tu.facility_id)
    return target,tu

def chat_revision_value(s,u,user_id=None):
    require_roles(u,'family','staff','middle_manager','admin')
    if u.role!='family' and not user_id:
        q=s.query(Message).filter_by(facility_id=u.facility_id)
    else:
        target,tu=chat_target(s,u,user_id)
        q=s.query(Message).filter_by(family_user_id=target,facility_id=tu.facility_id)
    max_id=int(q.with_entities(func.max(Message.id)).scalar() or 0)
    read_count=int(q.filter(Message.read_at.is_not(None)).count())
    return f'{max_id}:{read_count}'

@app.get('/api/messages/revision')
def messages_revision(request:Request,user_id:Optional[int]=None):
    with SessionLocal() as s:
        u=current_user(request,s)
        return {'revision':chat_revision_value(s,u,user_id)}

def line_account_status(s,uid):
    linked=bool(line_user_id_for(s,uid))
    attempt=s.query(LineLinkAttempt).filter_by(user_id=uid).order_by(LineLinkAttempt.id.desc()).first()
    failed=s.query(NotificationDelivery).filter_by(user_id=uid,channel='line',status='failed').first()
    status='認証済み保存失敗' if attempt and attempt.authenticated and attempt.code not in ('LINKED','UNLINKED') else '送信失敗' if linked and failed else '連携済み' if linked else '未連携'
    return {'status':status,'cause_code':attempt.code if attempt else ''}

def notification_health_summary(s,u):
    q=s.query(NotificationDelivery,NotificationEvent).join(NotificationEvent,NotificationDelivery.event_id==NotificationEvent.id).filter(NotificationDelivery.channel=='line')
    if u.role=='admin': q=q.filter(NotificationEvent.facility_id==u.facility_id)
    rows=q.order_by(NotificationDelivery.id.desc()).limit(1000).all()
    counts={'sent':0,'failed':0,'skipped_unlinked':0,'other':0}
    failed=[]
    for d,e in rows:
        if d.status in counts: counts[d.status]+=1
        else: counts['other']+=1
        if d.status=='failed' and len(failed)<10:
            failed.append({'id':d.id,'event_type':e.event_type,'title':e.title,'last_error':d.last_error,'attempts':d.attempts,'created_at':dt_iso(d.created_at)})
    users=s.query(User).filter(User.active==1)
    if u.role=='admin': users=users.filter(User.facility_id==u.facility_id)
    accounts=[{'user_id':x.id,'name':x.name,**line_account_status(s,x.id)} for x in users.all()]
    return {'accounts':accounts,'counts':counts,'failed':failed,'note':'LINEのsentはMessaging API受付成功を表し、端末通知・閲覧までを保証するものではありません。'}

@app.get('/api/bootstrap')
def bootstrap(request:Request):
    with SessionLocal() as s:
        u=current_user(request,s)
        fac=s.get(Facility,u.facility_id) if u.facility_id else None
        data={'user':user_json(u),'facility':facility_json(fac) if fac else None,'role_labels':ROLE_LABEL,'version':'20.7.1','targets':{'facilities':50,'staff':1000,'middle_managers':200,'admins':50,'families':2000,'sponsors':100}}
        if u.role in ('family','admin','hq'): data['document_revision']=document_revision_value(s,u)
        data['notifications']=[{'id':n.id,'title':n.title,'body':n.body,'module':n.module,'entity_id':n.entity_id,'created_at':dt_iso(n.created_at),'read_at':dt_iso(n.read_at),'deep_link':deep_link_for_user(u,n.module,n.entity_id)} for n in s.query(Notification).filter_by(user_id=u.id).order_by(Notification.id.desc()).limit(30).all()]
        data['line_notification']={**line_account_status(s,u.id),'linked':bool(line_user_id_for(s,u.id)),'line_user_id_masked':(('***'+line_user_id_for(s,u.id)[-6:]) if line_user_id_for(s,u.id) else ''),'liff_id':os.environ.get('LIFF_ID',''),'mock_mode':os.environ.get('LINE_MOCK_MODE','')=='1'}
        if u.role in ('admin','hq'): data['notification_health']=notification_health_summary(s,u)
        if u.role=='family':
            links=s.query(FamilyLink,Resident).join(Resident,FamilyLink.resident_id==Resident.id).filter(FamilyLink.user_id==u.id).all(); resident_ids=[r.id for _,r in links]
            data['residents']=[{'id':r.id,'name':r.name,'facility_id':r.facility_id,'service_type':r.service_type,'relation':l.relation} for l,r in links]
            data['news']=[{'id':n.id,'category':n.category,'title':n.title,'body':n.body,'important':n.important,'image_url':n.image_url,'created_at':dt_iso(n.created_at)} for n in s.query(News).filter_by(facility_id=u.facility_id).order_by(News.id.desc()).limit(30)]
            data['documents']=document_list_for_user(s,u)
            data['contract_cards']=contract_cards_for_family(s,u)
            data['visits']=visit_list_for_user(s,u)
            data['messages']=messages_for_family(s,u.id,u.facility_id)
            sup=s.query(Supplier).filter_by(code='CAREDO').first()
            if sup:
                prods=s.query(Product).filter_by(supplier_id=sup.id,active=1).limit(8).all()
                data['shop_recommendations']=[product_public_json(s,p) for p in prods]
            data['gift_recommendations']=partner_public_products(s,limit=12)
            data['customer_orders']=customer_orders_for_family(s,u.id)
        elif u.role in ('staff','middle_manager','admin'):
            data['news']=[{'id':n.id,'category':n.category,'title':n.title,'body':n.body,'important':n.important,'image_url':n.image_url,'created_at':dt_iso(n.created_at)} for n in s.query(News).filter_by(facility_id=u.facility_id).order_by(News.id.desc()).limit(30)]
            data['family_users']=family_users_for_facility(s,u.facility_id)
            data['resident_options']=residents_for_facility(s,u.facility_id)
            data['my_requests']=workflow_list(s,u,own=True)
            data['labor_summary']=labor_summary(s,u)
            data['visits']=visit_list_for_facility(s,u.facility_id)
            if u.role in ('middle_manager','admin'):
                data['approval_queue']=workflow_list(s,u,approval=True)
                data['facility_labor']=facility_labor_summary(s,u)
                data['procurement']=procurement_snapshot(s,u.facility_id)
            if u.role=='admin':
                data['documents']=document_list_for_facility(s,u.facility_id); data['contract_cards']=contract_cards_for_facility(s,u.facility_id); data['visit_rules']=visit_rules_for_facility(s,u.facility_id)
        elif u.role=='hq':
            data['dashboard']=hq_dashboard(s); data['facilities']=[facility_json(f) for f in s.query(Facility).filter_by(active=1).all()]; data['integrations']=integration_list(s)
        elif u.role=='caredo':
            data['caredo']=caredo_snapshot(s)
        elif u.role=='sponsor':
            data['sponsor']=sponsor_snapshot(s,u)
        data['sync_revisions']={m:module_revision_value(s,u,m) for m in SYNC_ALLOWED.get(u.role,set()) if m!='chat'}
        return data

# -------- unified notifications / LINE --------
class LineLinkIn(BaseModel):
    id_token:str

def _verify_line_id_token(raw):
    if os.environ.get('LINE_MOCK_MODE','')=='1' and raw.startswith('mock:'):
        return {'sub':raw.split(':',1)[1] or 'U-MOCK'}
    client_id=os.environ.get('LINE_LOGIN_CHANNEL_ID') or os.environ.get('LIFF_CHANNEL_ID')
    if not client_id: raise HTTPException(503,'LINE_LOGIN_CHANNEL_IDが未設定です')
    data=urllib.parse.urlencode({'id_token':raw,'client_id':client_id}).encode('utf-8')
    try:
        req=urllib.request.Request(LINE_API_BASE+'/oauth2/v2.1/verify',data=data,headers={'Content-Type':'application/x-www-form-urlencoded'},method='POST')
        with urllib.request.urlopen(req,timeout=8) as res: out=json.loads(res.read().decode('utf-8'))
    except Exception as e:
        raise HTTPException(400,'LINE IDトークンを検証できません') from e
    if out.get('aud') != client_id or out.get('iss') != 'https://access.line.me' or not isinstance(out.get('exp'), (int,float)) or out['exp'] <= datetime.now(timezone.utc).timestamp(): raise HTTPException(400,'LINE_TOKEN_INVALID')
    if not isinstance(out.get('sub'),str) or not out.get('sub'): raise HTTPException(400,'LINEユーザーIDを確認できません')
    return out

@app.get('/api/line/config')
def line_config():
    return {'liff_id':os.environ.get('LIFF_ID',''),'enabled':bool(os.environ.get('LINE_CHANNEL_ACCESS_TOKEN')),'mock_mode':os.environ.get('LINE_MOCK_MODE','')=='1'}

def record_line_attempt(user_id, code, authenticated):
    logging.getLogger('ittokai.line').warning('line_link user_id=%s code=%s authenticated=%s', user_id, code, authenticated)
    try:
        with SessionLocal() as log_session:
            log_session.add(LineLinkAttempt(user_id=user_id,code=code,authenticated=int(authenticated)))
            log_session.commit()
    except Exception:
        logging.getLogger('ittokai.line').error('line_link diagnostic persistence failed user_id=%s code=%s',user_id,code)

@app.post('/api/line/link')
def line_link(inp:LineLinkIn,request:Request):
    with SessionLocal() as s:
        u=current_user(request,s); uid=u.id; authenticated=False
        try:
            profile=_verify_line_id_token(inp.id_token); authenticated=True; line_id=profile['sub']
            # Reserve LINE identities even after unlink, preventing accidental account transfer.
            conflict=s.query(LineAccountLink).filter(LineAccountLink.line_user_id==line_id,LineAccountLink.user_id!=uid).first()
            legacy=s.query(User).filter(User.line_user_id==line_id,User.id!=uid).first()
            if conflict or legacy: raise HTTPException(409,'LINE_ACCOUNT_CONFLICT')
            link=s.query(LineAccountLink).filter_by(user_id=uid).first() or LineAccountLink(user_id=uid,line_user_id=line_id)
            link.line_user_id=line_id; link.verified=1; link.active=1; link.updated_at=now_dt(); s.add(link); u.line_user_id=line_id
            s.add(LineLinkAttempt(user_id=uid,code='LINKED',authenticated=1))
            audit(s,request,u,'notifications','line_link','user',uid,{},u.facility_id)
            s.commit()
            return {'ok':True,'linked':True}
        except Exception as exc:
            s.rollback()
            code=('LINE_ACCOUNT_CONFLICT' if isinstance(exc,IntegrityError) else str(exc.detail) if isinstance(exc,HTTPException) and exc.status_code==409 else 'LINE_SAVE_FAILED' if authenticated else 'LINE_AUTH_FAILED')
            record_line_attempt(uid,code,authenticated)
            message='LINE認証には成功しましたが、ITTOKAIとの連携保存に失敗しました' if authenticated else 'LINE認証を確認できませんでした'
            raise HTTPException(409 if code=='LINE_ACCOUNT_CONFLICT' else 503 if authenticated else 400, message+' ['+code+']') from exc

@app.post('/api/line/unlink')
def line_unlink(request:Request):
    with SessionLocal() as s:
        u=current_user(request,s); link=s.query(LineAccountLink).filter_by(user_id=u.id).first()
        if link: link.active=0; link.updated_at=now_dt()
        s.add(LineLinkAttempt(user_id=u.id,code='UNLINKED',authenticated=0))
        u.line_user_id=''; audit(s,request,u,'notifications','line_unlink','user',u.id,{},u.facility_id); s.commit(); return {'ok':True,'linked':False}

@app.post('/api/notifications/{nid}/read')
def notification_read(nid:int,request:Request):
    with SessionLocal() as s:
        u=current_user(request,s); n=s.get(Notification,nid)
        if not n or n.user_id!=u.id: raise HTTPException(404,'通知がありません')
        if not n.read_at: n.read_at=now_dt(); s.commit()
        return {'ok':True}

@app.get('/api/notifications/deliveries')
def notification_deliveries(request:Request,status:Optional[str]=None):
    with SessionLocal() as s:
        u=current_user(request,s); require_roles(u,'admin','hq')
        q=s.query(NotificationDelivery,NotificationEvent,User).join(NotificationEvent,NotificationDelivery.event_id==NotificationEvent.id).join(User,NotificationDelivery.user_id==User.id)
        if u.role=='admin': q=q.filter(NotificationEvent.facility_id==u.facility_id)
        if status: q=q.filter(NotificationDelivery.status==status)
        items=[]
        for d,e,usr in q.order_by(NotificationDelivery.id.desc()).limit(500):
            items.append({'id':d.id,'event_id':e.id,'event_type':e.event_type,'title':e.title,'user_id':usr.id,'user_name':usr.name,'role':usr.role,'channel':d.channel,'status':d.status,'attempts':d.attempts,'deep_link':d.deep_link,'last_error':d.last_error,'sent_at':dt_iso(d.sent_at),'created_at':dt_iso(d.created_at)})
        return {'items':items}

@app.post('/api/notifications/retry-failed')
def retry_failed_notifications(request:Request):
    with SessionLocal() as s:
        u=current_user(request,s); require_roles(u,'admin','hq')
        q=s.query(NotificationDelivery).filter(NotificationDelivery.channel=='line',NotificationDelivery.status=='failed')
        if u.role=='admin': q=q.join(NotificationEvent,NotificationDelivery.event_id==NotificationEvent.id).filter(NotificationEvent.facility_id==u.facility_id)
        count=0
        for d in q.limit(100).all():
            e=s.get(NotificationEvent,d.event_id); usr=s.get(User,d.user_id); ok,err=_line_push(line_user_id_for(s,d.user_id),e.title,e.body,d.deep_link)
            d.attempts=(d.attempts or 0)+1; d.updated_at=now_dt(); d.last_error=err; d.status='sent' if ok else ('skipped_unlinked' if err=='LINE未連携' else 'failed'); d.sent_at=now_dt() if ok else d.sent_at; count+=1
        audit(s,request,u,'notifications','retry_failed','notification_delivery','',{'count':count},u.facility_id); s.commit(); return {'ok':True,'retried':count}

class NotificationRuleIn(BaseModel):
    event_type:str; recipient_role:str='*'; app_enabled:bool=True; line_enabled:bool=True; repeat_after_hours:int=24; deadline_before_hours:int=72; active:bool=True

@app.get('/api/notifications/rules')
def notification_rules(request:Request):
    with SessionLocal() as s:
        u=current_user(request,s); require_roles(u,'admin','hq')
        rows=s.query(NotificationRule).order_by(NotificationRule.event_type,NotificationRule.recipient_role).all()
        return {'items':[{'id':x.id,'event_type':x.event_type,'recipient_role':x.recipient_role,'app_enabled':bool(x.app_enabled),'line_enabled':bool(x.line_enabled),'repeat_after_hours':x.repeat_after_hours,'deadline_before_hours':x.deadline_before_hours,'active':bool(x.active)} for x in rows]}

@app.post('/api/notifications/rules')
def notification_rule_save(inp:NotificationRuleIn,request:Request):
    with SessionLocal() as s:
        u=current_user(request,s); require_roles(u,'admin','hq')
        if inp.repeat_after_hours<0 or inp.deadline_before_hours<0: raise HTTPException(400,'通知時間の設定が不正です')
        r=s.query(NotificationRule).filter_by(event_type=inp.event_type,recipient_role=inp.recipient_role).first() or NotificationRule(event_type=inp.event_type,recipient_role=inp.recipient_role)
        r.app_enabled=1 if inp.app_enabled else 0; r.line_enabled=1 if inp.line_enabled else 0; r.repeat_after_hours=inp.repeat_after_hours; r.deadline_before_hours=inp.deadline_before_hours; r.active=1 if inp.active else 0; r.updated_at=now_dt(); s.add(r)
        audit(s,request,u,'notifications','rule_save','notification_rule',r.id or '',{'event_type':inp.event_type,'recipient_role':inp.recipient_role},u.facility_id); s.commit(); return {'ok':True}

def _scheduled_reminder_allowed(s,event_type,module,entity_id,user_id,repeat_hours):
    repeat_hours=max(1,int(repeat_hours or 24))
    last=(s.query(NotificationEvent).join(NotificationDelivery,NotificationDelivery.event_id==NotificationEvent.id)
          .filter(NotificationEvent.event_type==event_type,NotificationEvent.module==module,NotificationEvent.entity_id==str(entity_id),NotificationDelivery.user_id==user_id)
          .order_by(NotificationEvent.created_at.desc()).first())
    return (not last) or ((now_dt()-last.created_at)>=timedelta(hours=repeat_hours))

def run_due_reminders(s):
    result={'documents':0,'workflows':0}
    now=now_dt(); today=now.date()
    for r,d in s.query(DocumentRecipient,Document).join(Document,DocumentRecipient.document_id==Document.id).filter(DocumentRecipient.signed_at.is_(None),Document.status.in_(['sent','partial_signed'])).all():
        due=safe_date(d.due_at)
        if not due: continue
        usr=s.get(User,r.user_id); rule=notification_rule_for(s,'document.reminder',usr.role if usr else 'family')
        if rule and not rule.active: continue
        deadline_hours=(rule.deadline_before_hours if rule else 72) or 72
        repeat_hours=(rule.repeat_after_hours if rule else 24) or 24
        hours_to=(datetime.combine(due,time.max)-now).total_seconds()/3600
        if 0<=hours_to<=deadline_hours and _scheduled_reminder_allowed(s,'document.reminder','documents',d.id,r.user_id,repeat_hours):
            key=f'document.scheduled_reminder:{d.id}:{r.user_id}:{int(now.timestamp()//(max(1,repeat_hours)*3600))}'
            z=emit_notification_event(s,'document.reminder',[r.user_id],'書類の確認期限が近づいています','ITTOKAIで未処理の書類をご確認ください。','documents',d.id,d.facility_id,None,key,{'hours_to_due':round(hours_to,1),'scheduled':True})
            if not z['deduplicated']: result['documents']+=1
    for w in s.query(WorkflowRequest).filter(WorkflowRequest.status.in_(['pending_manager','pending_admin'])).all():
        role='middle_manager' if w.status=='pending_manager' else 'admin'; rec=users_by_roles(s,[role],w.facility_id,w.department_id if role=='middle_manager' else None)
        if not rec and role=='middle_manager': rec=users_by_roles(s,['admin'],w.facility_id)
        for usr in rec:
            rule=notification_rule_for(s,'workflow.forwarded',usr.role); repeat_hours=(rule.repeat_after_hours if rule else 24) or 24
            if rule and not rule.active: continue
            if (now-w.created_at)<timedelta(hours=repeat_hours): continue
            if not _scheduled_reminder_allowed(s,'workflow.forwarded','workflow',w.id,usr.id,repeat_hours): continue
            key=f'workflow.scheduled_reminder:{w.id}:{usr.id}:{int(now.timestamp()//(max(1,repeat_hours)*3600))}'
            z=emit_notification_event(s,'workflow.forwarded',[usr.id],'承認待ち申請があります','ITTOKAIで承認待ちの申請をご確認ください。','workflow',w.id,w.facility_id,None,key,{'scheduled':True})
            if not z['deduplicated']: result['workflows']+=1
    return result

@app.post('/api/notifications/run-reminders')
def notification_run_reminders(request:Request):
    with SessionLocal() as s:
        secret=os.environ.get('NOTIFICATION_CRON_SECRET',''); supplied=request.headers.get('x-ittokai-cron-secret','')
        actor=None
        if not (secret and hmac.compare_digest(secret,supplied)):
            actor=current_user(request,s); require_roles(actor,'admin','hq')
        result=run_due_reminders(s); s.commit(); return {'ok':True,**result}

# -------- chat/news --------
def family_users_for_facility(s,fid):
    out=[]
    for x in s.query(User).filter_by(role='family',facility_id=fid,active=1).all():
        links=s.query(FamilyLink,Resident).join(Resident,FamilyLink.resident_id==Resident.id).filter(FamilyLink.user_id==x.id,Resident.facility_id==fid).all()
        out.append({'id':x.id,'name':x.name,'email':x.email,'residents':[{'id':r.id,'name':r.name,'relation':l.relation,'primary':bool(l.is_primary)} for l,r in links]})
    return out

def residents_for_facility(s,fid):
    return [{'id':r.id,'resident_code':r.resident_code,'name':r.name,'service_type':r.service_type} for r in s.query(Resident).filter_by(facility_id=fid,active=1).order_by(Resident.name).all()]

def messages_for_family(s,uid,fid): return [{'id':m.id,'sender_id':m.sender_id,'sender_name':m.sender_name,'sender_role':m.sender_role,'body':m.body,'created_at':dt_iso(m.created_at),'read_at':dt_iso(m.read_at)} for m in s.query(Message).filter_by(family_user_id=uid,facility_id=fid).order_by(Message.id.asc()).all()]

@app.get('/api/messages')
def get_messages(request:Request,user_id:Optional[int]=None):
    with SessionLocal() as s:
        u=current_user(request,s); target,tu=chat_target(s,u,user_id)
        changed=False
        for m in s.query(Message).filter_by(family_user_id=target,facility_id=tu.facility_id).all():
            if m.sender_id!=u.id and not m.read_at:
                m.read_at=now_dt(); changed=True
        if changed:
            audit(s,request,u,'chat','read','message_thread',target,{'target':target},tu.facility_id)
            s.commit()
        return {'messages':messages_for_family(s,target,tu.facility_id),'revision':chat_revision_value(s,u,target if u.role!='family' else None)}

class MsgIn(BaseModel): body:str; user_id:Optional[int]=None
@app.post('/api/messages')
def post_message(inp:MsgIn,request:Request):
    with SessionLocal() as s:
        u=current_user(request,s); target,tu=chat_target(s,u,inp.user_id)
        body=inp.body.strip()
        if not body: raise HTTPException(400,'メッセージを入力してください')
        if len(body)>4000: raise HTTPException(400,'メッセージは4000文字以内にしてください')
        m=Message(facility_id=tu.facility_id,family_user_id=target,sender_id=u.id,sender_name=u.name,sender_role=u.role,body=body); s.add(m); s.flush()
        if u.role=='family':
            recipients=s.query(User).filter(User.facility_id==tu.facility_id,User.role.in_(['middle_manager','admin']),User.active==1).all()
            emit_notification_event(s,'chat.new_message',recipients,'新着トークがあります','ITTOKAIでメッセージをご確認ください。','chat',target,tu.facility_id,u.id,f'chat.new_message:{m.id}:facility')
        else:
            emit_notification_event(s,'chat.new_message',[target],'施設から新着トークがあります','ITTOKAIでメッセージをご確認ください。','chat',target,tu.facility_id,u.id,f'chat.new_message:{m.id}:family')
        audit(s,request,u,'chat','send','message',m.id,{'target':target},tu.facility_id); s.commit(); return {'ok':True,'id':m.id}

@app.post('/api/messages/read')
def read_messages(inp:dict,request:Request):
    with SessionLocal() as s:
        u=current_user(request,s); target,tu=chat_target(s,u,inp.get('user_id'))
        changed=False
        for m in s.query(Message).filter_by(family_user_id=target,facility_id=tu.facility_id).all():
            if m.sender_id!=u.id and not m.read_at:
                m.read_at=now_dt(); changed=True
        if changed:
            audit(s,request,u,'chat','read','message_thread',target,{'target':target},tu.facility_id)
            s.commit()
        return {'ok':True}


class NewsIn(BaseModel):
    category:str='お知らせ'; title:str; body:str; important:bool=False; image_url:str=''
@app.post('/api/news')
def create_news(inp:NewsIn,request:Request):
    with SessionLocal() as s:
        u=current_user(request,s); require_roles(u,'admin','hq'); fid=u.facility_id or 1
        n=News(facility_id=fid,category=inp.category,title=inp.title,body=inp.body,important=1 if inp.important else 0,image_url=inp.image_url); s.add(n); s.flush()
        for fam in s.query(User).filter_by(facility_id=fid,role='family',active=1).all(): notify(s,fam.id,'施設からのお知らせ',inp.title,'news',n.id)
        audit(s,request,u,'news','publish','news',n.id,{},fid); s.commit(); return {'id':n.id,'ok':True}

# -------- documents --------
def document_list_for_user(s,u):
    rows=s.query(DocumentRecipient,Document).join(Document,DocumentRecipient.document_id==Document.id).filter(DocumentRecipient.user_id==u.id).order_by(Document.id.desc()).all()
    return [doc_json(d,r) for r,d in rows]
def document_list_for_facility(s,fid):
    rows=s.query(DocumentRecipient,Document,User).join(Document,DocumentRecipient.document_id==Document.id).join(User,DocumentRecipient.user_id==User.id).filter(Document.facility_id==fid).order_by(Document.id.desc()).all()
    return [dict(doc_json(d,r),recipient_name=u.name,recipient_user_id=u.id) for r,d,u in rows]
def doc_json(d,r=None):
    return {'id':d.id,'facility_id':d.facility_id,'title':d.title,'category':d.category,'body':d.body,'version':d.version,'status':r.status if r else d.status,'due_at':dt_iso(d.due_at),'requires_signature':d.requires_signature,'file_name':d.file_name,'sent_at':dt_iso(d.sent_at),'viewed_at':dt_iso(r.viewed_at) if r else None,'consent_at':dt_iso(r.consent_at) if r else None,'signed_at':dt_iso(r.signed_at) if r else None}

def link_document_resident(s,doc_id,resident_id):
    if not resident_id: return
    if not s.query(DocumentResidentLink).filter_by(document_id=doc_id,resident_id=resident_id).first():
        s.add(DocumentResidentLink(document_id=doc_id,resident_id=resident_id))

def ensure_document_resident_links(s):
    linked={x[0] for x in s.query(DocumentResidentLink.document_id).all()}
    for d in s.query(Document).all():
        if d.id in linked: continue
        for rec in s.query(DocumentRecipient).filter_by(document_id=d.id).all():
            fl=s.query(FamilyLink).filter_by(user_id=rec.user_id,is_primary=1).first() or s.query(FamilyLink).filter_by(user_id=rec.user_id).first()
            if fl: link_document_resident(s,d.id,fl.resident_id)
    s.flush()

def contract_card_for_resident(s,resident_id):
    r=s.get(Resident,resident_id)
    if not r: return None
    links=s.query(DocumentResidentLink).filter_by(resident_id=resident_id).all(); doc_ids=[x.document_id for x in links]
    docs=[]
    for d in (s.query(Document).filter(Document.id.in_(doc_ids)).order_by(Document.category,Document.title,Document.version.desc()).all() if doc_ids else []):
        recs=s.query(DocumentRecipient).filter_by(document_id=d.id).all()
        status='signed' if recs and all(x.signed_at for x in recs) else ('partial_signed' if any(x.signed_at for x in recs) else (recs[0].status if recs else d.status))
        docs.append({'id':d.id,'title':d.title,'category':d.category,'version':d.version,'status':status,'due_at':dt_iso(d.due_at),'file_name':d.file_name,'original_hash':d.original_hash,'sent_at':dt_iso(d.sent_at)})
    latest={}
    for d in docs:
        k=(d['category'],d['title'])
        if k not in latest or d['version']>latest[k]['version']: latest[k]=d
    current=list(latest.values())
    return {'resident_id':r.id,'resident_code':r.resident_code,'resident_name':r.name,'service_type':r.service_type,'documents':docs,'current_documents':current,'summary':{'total':len(current),'signed':sum(1 for x in current if x['status']=='signed'),'pending':sum(1 for x in current if x['status']!='signed'),'expiring':sum(1 for x in current if x['due_at'] and safe_date(x['due_at']) and 0 <= (safe_date(x['due_at'])-datetime.now(JST).date()).days <= 30)}}

def contract_cards_for_family(s,u):
    ensure_document_resident_links(s)
    return [contract_card_for_resident(s,l.resident_id) for l in s.query(FamilyLink).filter_by(user_id=u.id).all() if contract_card_for_resident(s,l.resident_id)]

def contract_cards_for_facility(s,fid):
    ensure_document_resident_links(s)
    return [contract_card_for_resident(s,r.id) for r in s.query(Resident).filter_by(facility_id=fid,active=1).all()]

class DocumentCreate(BaseModel): title:str; category:str='書類'; body:str=''; recipient_user_ids:list[int]; resident_id:Optional[int]=None; due_days:int=7; requires_signature:bool=True; required_checks:list[str]=['内容を確認しました']
@app.post('/api/documents')
def create_document(inp:DocumentCreate,request:Request):
    with SessionLocal() as s:
        u=current_user(request,s); require_roles(u,'admin','hq'); fid=u.facility_id or 1
        d=Document(facility_id=fid,title=inp.title,category=inp.category,body=inp.body,status='sent',due_at=now_dt()+timedelta(days=max(1,inp.due_days)),requires_signature=1 if inp.requires_signature else 0,required_checks_json=json.dumps(inp.required_checks,ensure_ascii=False),created_by=u.id,sent_at=now_dt()); d.original_hash=sha256_text(d.title+d.body+now_iso()); s.add(d); s.flush()
        resident_id=inp.resident_id
        if resident_id:
            rr=s.get(Resident,resident_id)
            if not rr or rr.facility_id!=fid: raise HTTPException(400,'対象利用者が不正です')
            link_document_resident(s,d.id,resident_id)
        for uid in inp.recipient_user_ids:
            ru=s.get(User,uid)
            if not ru or ru.role!='family' or ru.facility_id!=fid: continue
            s.add(DocumentRecipient(document_id=d.id,user_id=uid,signer_role='家族',sign_order=1,status='sent')); emit_notification_event(s,'document.requested',[uid],'確認・署名が必要な書類があります','ITTOKAIで書類をご確認ください。','documents',d.id,fid,u.id,f'document.requested:{d.id}:{uid}')
            if not resident_id:
                fl=s.query(FamilyLink).filter_by(user_id=uid,is_primary=1).first() or s.query(FamilyLink).filter_by(user_id=uid).first()
                if fl: link_document_resident(s,d.id,fl.resident_id)
        audit(s,request,u,'documents','create_send','document',d.id,{'recipients':inp.recipient_user_ids},fid); s.commit(); return {'id':d.id,'ok':True}

@app.post('/api/documents/upload')
async def upload_document(request:Request,title:str=Form(...),category:str=Form('契約書'),recipient_user_ids:str=Form(...),resident_id:Optional[int]=Form(None),due_days:int=Form(7),file:UploadFile=File(...)):
    data=await file.read()
    if len(data)>12*1024*1024: raise HTTPException(400,'PDFは12MB以下にしてください')
    if not data.startswith(b'%PDF'): raise HTTPException(400,'PDF形式のファイルを選択してください')
    with SessionLocal() as s:
        u=current_user(request,s); require_roles(u,'admin','hq'); fid=u.facility_id or 1
        try: ids=[int(x) for x in recipient_user_ids.split(',') if x.strip()]
        except Exception: ids=[]
        if not ids: raise HTTPException(400,'送付先が必要です')
        d=Document(facility_id=fid,title=title,category=category,body='',status='sent',due_at=now_dt()+timedelta(days=due_days),requires_signature=1,file_name=file.filename or 'document.pdf',file_blob=data,original_hash=sha256_bytes(data),created_by=u.id,sent_at=now_dt()); s.add(d); s.flush()
        if resident_id:
            rr=s.get(Resident,resident_id)
            if not rr or rr.facility_id!=fid: raise HTTPException(400,'対象利用者が不正です')
            link_document_resident(s,d.id,resident_id)
        order=1
        for uid in ids:
            ru=s.get(User,uid)
            if not ru or ru.facility_id!=fid or ru.role!='family': continue
            s.add(DocumentRecipient(document_id=d.id,user_id=uid,signer_role='家族',sign_order=order,status='sent')); emit_notification_event(s,'document.requested',[uid],'確認・署名が必要な契約書があります','ITTOKAIで契約書をご確認ください。','documents',d.id,fid,u.id,f'document.requested:{d.id}:{uid}'); order+=1
            if not resident_id:
                fl=s.query(FamilyLink).filter_by(user_id=uid,is_primary=1).first() or s.query(FamilyLink).filter_by(user_id=uid).first()
                if fl: link_document_resident(s,d.id,fl.resident_id)
        audit(s,request,u,'documents','upload_send','document',d.id,{'file':file.filename,'recipients':ids},fid); s.commit(); return {'id':d.id,'ok':True}

@app.post('/api/documents/{doc_id}/view')
def view_document(doc_id:int,request:Request):
    with SessionLocal() as s:
        u=current_user(request,s); d=s.get(Document,doc_id)
        if not d: raise HTTPException(404,'書類がありません')
        r=s.query(DocumentRecipient).filter_by(document_id=doc_id,user_id=u.id).first()
        if not r:
            require_roles(u,'admin','hq'); facility_scope(u,d.facility_id)
        if r and not r.viewed_at:
            r.viewed_at=now_dt(); r.status='viewed'; audit(s,request,u,'documents','view','document',doc_id,facility_id=d.facility_id); s.commit()
        return {'ok':True}

class SignIn(BaseModel): signer_name:str; signature_data:str; consent:bool=True
@app.post('/api/documents/{doc_id}/sign')
def sign_document(doc_id:int,inp:SignIn,request:Request):
    with SessionLocal() as s:
        u=current_user(request,s); r=s.query(DocumentRecipient).filter_by(document_id=doc_id,user_id=u.id).first(); d=s.get(Document,doc_id)
        if not r or not d: raise HTTPException(404,'書類がありません')
        if r.signed_at: raise HTTPException(400,'署名済みです')
        if not inp.consent: raise HTTPException(400,'同意が必要です')
        sig_hash=sha256_text(inp.signature_data); sig=DocumentSignature(document_id=doc_id,recipient_id=r.id,user_id=u.id,signer_name=inp.signer_name,signature_data=inp.signature_data,signature_hash=sig_hash,ip=client_ip(request),user_agent=request.headers.get('user-agent','')); s.add(sig); r.consent_at=now_dt(); r.signed_at=now_dt(); r.status='signed'; s.flush(); remaining=s.query(DocumentRecipient).filter(DocumentRecipient.document_id==doc_id,DocumentRecipient.signed_at.is_(None)).count(); d.status='signed' if remaining==0 else 'partial_signed'; audit(s,request,u,'documents','sign','document',doc_id,{'signature_hash':sig_hash});
        recipients=users_by_roles(s,['admin'],d.facility_id)
        creator=s.get(User,d.created_by) if d.created_by else None
        if creator and creator.active and creator.id not in [x.id for x in recipients]: recipients.append(creator)
        emit_notification_event(s,'document.signed',recipients,'ご家族の書類署名が完了しました','ITTOKAIで署名済み書類をご確認ください。','documents',doc_id,d.facility_id,u.id,f'document.signed:{doc_id}:{u.id}')
        s.commit(); return {'ok':True}

@app.get('/api/documents/{doc_id}/original')
def original_document(doc_id:int,request:Request):
    with SessionLocal() as s:
        u=current_user(request,s); d=s.get(Document,doc_id)
        if not d or not d.file_blob: raise HTTPException(404,'原本PDFがありません')
        r=s.query(DocumentRecipient).filter_by(document_id=doc_id,user_id=u.id).first()
        if not r:
            require_roles(u,'admin','hq'); facility_scope(u,d.facility_id)
        return Response(content=bytes(d.file_blob),media_type='application/pdf',headers={'Content-Disposition':f'inline; filename="original_{doc_id}.pdf"'})

class ReissueIn(BaseModel):
    title:Optional[str]=None; body:Optional[str]=None; due_days:int=7
@app.post('/api/documents/{doc_id}/reissue')
def reissue_document(doc_id:int,inp:ReissueIn,request:Request):
    with SessionLocal() as s:
        u=current_user(request,s); require_roles(u,'admin','hq'); old=s.get(Document,doc_id)
        if not old: raise HTTPException(404,'書類がありません')
        facility_scope(u,old.facility_id)
        d=Document(facility_id=old.facility_id,title=inp.title or old.title,category=old.category,body=inp.body if inp.body is not None else old.body,version=old.version+1,status='sent',due_at=now_dt()+timedelta(days=inp.due_days),requires_signature=old.requires_signature,required_checks_json=old.required_checks_json,file_name=old.file_name,file_blob=old.file_blob,original_hash=old.original_hash,created_by=u.id,sent_at=now_dt(),previous_document_id=old.id); s.add(d); s.flush()
        for idx,r in enumerate(s.query(DocumentRecipient).filter_by(document_id=old.id).order_by(DocumentRecipient.sign_order).all(),1):
            s.add(DocumentRecipient(document_id=d.id,user_id=r.user_id,signer_role=r.signer_role,sign_order=idx,status='sent')); emit_notification_event(s,'document.requested',[r.user_id],'書類が更新されました','ITTOKAIで再発行された書類をご確認ください。','documents',d.id,d.facility_id,u.id,f'document.reissue:{d.id}:{r.user_id}')
        for l in s.query(DocumentResidentLink).filter_by(document_id=old.id).all(): link_document_resident(s,d.id,l.resident_id)
        old.status='superseded'; audit(s,request,u,'documents','reissue','document',d.id,{'previous':old.id},old.facility_id); s.commit(); return {'id':d.id,'version':d.version,'ok':True}

@app.post('/api/documents/{doc_id}/remind')
def remind_document(doc_id:int,request:Request):
    with SessionLocal() as s:
        u=current_user(request,s); require_roles(u,'admin','hq'); d=s.get(Document,doc_id)
        if not d: raise HTTPException(404,'書類がありません')
        facility_scope(u,d.facility_id)
        n=0
        for r in s.query(DocumentRecipient).filter_by(document_id=d.id).all():
            if not r.signed_at: emit_notification_event(s,'document.reminder',[r.user_id],'書類の確認をお願いします','ITTOKAIで未処理の書類をご確認ください。','documents',d.id,d.facility_id,u.id,f'document.manual_reminder:{d.id}:{r.user_id}:{datetime.now(JST).date().isoformat()}'); n+=1
        audit(s,request,u,'documents','remind','document',d.id,{'count':n},d.facility_id); s.commit(); return {'ok':True,'count':n}

@app.get('/api/documents/{doc_id}/pdf')
def signed_pdf(doc_id:int,request:Request):
    with SessionLocal() as s:
        u=current_user(request,s); d=s.get(Document,doc_id)
        if not d: raise HTTPException(404,'書類がありません')
        r=s.query(DocumentRecipient).filter_by(document_id=doc_id,user_id=u.id).first()
        if not r:
            require_roles(u,'admin','hq'); facility_scope(u,d.facility_id)
        pdf=make_document_pdf(s,d,r); return Response(content=pdf,media_type='application/pdf',headers={'Content-Disposition':f'inline; filename="ITTOKAI_document_{doc_id}.pdf"'})

# -------- workflow / labor management --------
def _minutes_between(start,end):
    try:
        a=datetime.strptime(start,'%H:%M'); b=datetime.strptime(end,'%H:%M'); m=int((b-a).total_seconds()/60)
        return m if m>=0 else m+24*60
    except Exception: return 0

def _ensure_demo_leave_grants(s,u):
    grants=s.query(LeaveGrant).filter_by(user_id=u.id).order_by(LeaveGrant.grant_date).all()
    if grants: return grants
    # Demo/test fallback only. Production should import exact grant dates/balances from HR/Chronos.
    if (u.email or '').endswith('@ittokai.local'):
        today=datetime.now(JST).date(); y=today.year
        current=date(y,4,1) if today>=date(y,4,1) else date(y-1,4,1); previous=add_years(current,-1)
        for gd,days,rem,note in [(previous,20,6,'前回付与・デモ'),(current,20,18,'今回付与・デモ')]:
            s.add(LeaveGrant(user_id=u.id,grant_date=gd.isoformat(),granted_days=days,remaining_days=rem,taken_days=max(0,days-rem),expires_on=(add_years(gd,2)-timedelta(days=1)).isoformat(),source='demo',note=note))
        s.flush(); grants=s.query(LeaveGrant).filter_by(user_id=u.id).order_by(LeaveGrant.grant_date).all()
    return grants

def _apply_leave_to_grants(s,w):
    if w.request_type!='有給申請' or w.status!='approved' or s.query(LeaveUsage).filter_by(request_id=w.id).first(): return
    payload=json.loads(w.payload_json or '{}') or {}; days=float(payload.get('leave_days',1) or 1)
    if days<=0: return
    grants=_ensure_demo_leave_grants(s,s.get(User,w.applicant_id))
    reqd=safe_date(w.request_date) or datetime.now(JST).date(); remaining=days; alloc=[]
    for g in sorted(grants,key=lambda x:(x.expires_on,x.grant_date)):
        if remaining<=1e-9: break
        if safe_date(g.expires_on) and safe_date(g.expires_on)<reqd: continue
        take=min(float(g.remaining_days or 0),remaining)
        if take<=0: continue
        g.remaining_days=round(float(g.remaining_days or 0)-take,3); g.taken_days=round(float(g.taken_days or 0)+take,3); remaining=round(remaining-take,3); alloc.append({'grant_id':g.id,'grant_date':g.grant_date,'days':take})
    if remaining>1e-9: raise HTTPException(400,f'有給残数が不足しています（不足 {remaining:g} 日）')
    s.add(LeaveUsage(request_id=w.id,user_id=w.applicant_id,used_days=days,allocation_json=json.dumps(alloc,ensure_ascii=False)))

def leave_ledger_summary(s,u):
    grants=_ensure_demo_leave_grants(s,u); today=datetime.now(JST).date()
    active=[g for g in grants if not safe_date(g.expires_on) or safe_date(g.expires_on)>=today]
    active=sorted(active,key=lambda x:x.grant_date)
    total=round(sum(float(g.remaining_days or 0) for g in active),1)
    lots=[]
    for g in active:
        exp=safe_date(g.expires_on); days_to=(exp-today).days if exp else None
        lots.append({'id':g.id,'grant_date':g.grant_date,'granted_days':round(float(g.granted_days or 0),1),'remaining_days':round(float(g.remaining_days or 0),1),'taken_days':round(float(g.taken_days or 0),1),'expires_on':g.expires_on,'days_to_expire':days_to,'source':g.source})
    # Current/previous are grant lots, not fiscal-year buckets. This reflects the statutory 2-year right.
    current=lots[-1] if lots else None; previous=lots[-2] if len(lots)>1 else None
    expiring=[x for x in lots if x['remaining_days']>0 and x['days_to_expire'] is not None and x['days_to_expire']<=180]
    obligation=None
    if current and current['granted_days']>=10:
        gd=safe_date(current['grant_date']); end=add_years(gd,1)-timedelta(days=1); elapsed=(today-gd).days; left=max(0,(end-today).days)
        taken=float(current['taken_days'] or 0); obligation={'period_start':gd.isoformat(),'period_end':end.isoformat(),'taken_days':round(taken,1),'required_days':5,'remaining_required':round(max(0,5-taken),1),'days_left':left,'compliant':taken>=5 or today<gd}
    return {'total_remaining':total,'current_grant':current,'previous_grant':previous,'lots':lots,'expiring':expiring,'five_day_obligation':obligation,'legal_note':'年休は付与日から2年で時効。年10日以上付与者は基準日から1年以内に5日取得が必要。複雑な前倒し付与・基準日統一は個別設定が必要。'}

def labor_policy(s,u):
    if not u.facility_id:
        return {'agreement_year_start':'04-01','monthly_overtime_limit':45,'annual_overtime_limit':360,'special_annual_limit':720,'combined_month_limit':100,'rolling_average_limit':80,'over45_months_limit':6,'annual_leave_required_days':5,'time_leave_max_days':5}
    p=s.query(LaborPolicy).filter_by(facility_id=u.facility_id).first()
    if not p:
        p=LaborPolicy(facility_id=u.facility_id); s.add(p); s.flush()
    return {'agreement_year_start':p.agreement_year_start,'monthly_overtime_limit':float(p.monthly_overtime_limit),'annual_overtime_limit':float(p.annual_overtime_limit),'special_annual_limit':float(p.special_annual_limit),'combined_month_limit':float(p.combined_month_limit),'rolling_average_limit':float(p.rolling_average_limit),'over45_months_limit':int(p.over45_months_limit),'annual_leave_required_days':float(p.annual_leave_required_days),'time_leave_max_days':float(p.time_leave_max_days)}

def agreement_period(today:date,start_md='04-01'):
    try: mm,dd=[int(x) for x in start_md.split('-',1)]; start=date(today.year,mm,dd)
    except Exception: start=date(today.year,4,1)
    if today<start: start=add_years(start,-1)
    return start,add_years(start,1)-timedelta(days=1)

def _work_time_months(s,u):
    rows=s.query(WorkTimeRecord).filter_by(user_id=u.id).all(); out={}
    for r in rows:
        d=safe_date(r.work_date)
        if not d: continue
        k=month_key(d); x=out.setdefault(k,{'overtime':0,'holiday':0,'days':0}); x['overtime']+=int(r.overtime_minutes or 0); x['holiday']+=int(r.holiday_minutes or 0); x['days']+=1
    return out

def labor_summary(s,u):
    own=s.query(WorkflowRequest).filter_by(applicant_id=u.id).all(); approved=[w for w in own if w.status=='approved']; today=datetime.now(JST).date(); policy=labor_policy(s,u)
    period_start,period_end=agreement_period(today,policy['agreement_year_start'])
    changes=sum(1 for w in approved if w.request_type=='勤務変更' and safe_date(w.request_date) and period_start<=safe_date(w.request_date)<=period_end)
    ledger=leave_ledger_summary(s,u)
    months=_work_time_months(s,u); actual_months=set(months.keys()); fallback_used=False
    for w in approved:
        if w.request_type!='残業申請': continue
        wd=safe_date(w.request_date)
        if not wd: continue
        k=month_key(wd)
        if k in actual_months: continue
        x=months.setdefault(k,{'overtime':0,'holiday':0,'days':0}); x['overtime']+=_minutes_between(w.start_time,w.end_time); fallback_used=True
    current_key=month_key(today)
    if current_key in actual_months and not fallback_used: source='クロノス実績'
    elif actual_months: source='クロノス実績＋申請補完（参考）'
    else: source='申請ベース（参考）'
    cur=months.get(current_key,{'overtime':0,'holiday':0}); hours=round(cur['overtime']/60,1); legal_combined=round((cur['overtime']+cur['holiday'])/60,1)
    annual_minutes=0; exceed45=0
    for k,v in months.items():
        try: kd=date.fromisoformat(k+'-01')
        except Exception: continue
        if date(period_start.year,period_start.month,1)<=kd<=date(period_end.year,period_end.month,1):
            annual_minutes+=v['overtime']; exceed45 += 1 if v['overtime']/60>policy['monthly_overtime_limit'] else 0
    annual_ot=round(annual_minutes/60,1)
    # rolling averages 2-6 calendar months ending current month
    rolling=[]
    y,m=today.year,today.month
    keys=[]
    for i in range(6):
        yy=y; mm=m-i
        while mm<=0: mm+=12; yy-=1
        keys.append(f'{yy:04d}-{mm:02d}')
    for n in range(2,7):
        vals=[(months.get(k,{}).get('overtime',0)+months.get(k,{}).get('holiday',0))/60 for k in keys[:n]]
        rolling.append({'months':n,'average':round(sum(vals)/n,1)})
    maxavg=max((x['average'] for x in rolling),default=0)
    alerts=[]
    if legal_combined>=policy['combined_month_limit']: alerts.append({'level':'danger','title':'時間外＋休日労働100時間以上','body':'法令上限に抵触する可能性があります。即時確認してください'})
    elif legal_combined>=policy['rolling_average_limit']: alerts.append({'level':'danger','title':'時間外＋休日労働80時間以上','body':'複数月平均80時間以内の確認が必要です'})
    if maxavg>=policy['rolling_average_limit']: alerts.append({'level':'danger','title':'複数月平均80時間以上','body':f'2〜6か月平均の最大 {maxavg:g}h。即時確認が必要です'})
    if hours>=policy['monthly_overtime_limit']: alerts.append({'level':'danger','title':'時間外労働45時間超','body':'原則上限を超えています。36協定・特別条項と回数を確認してください'})
    elif hours>=max(0,policy['monthly_overtime_limit']-5): alerts.append({'level':'warn','title':'時間外労働40時間超','body':'45時間接近。業務調整を検討してください'})
    elif hours>=30: alerts.append({'level':'notice','title':'時間外労働30時間超','body':'月間時間外労働に注意してください'})
    if annual_ot>=policy['special_annual_limit']: alerts.append({'level':'danger','title':'年間時間外720時間以上','body':'特別条項があっても上限超過の可能性があります'})
    elif annual_ot>=policy['annual_overtime_limit']: alerts.append({'level':'warn','title':'年間時間外360時間超','body':'原則上限超過。特別条項の適用状況を確認してください'})
    if exceed45>policy['over45_months_limit']: alerts.append({'level':'danger','title':'月45時間超が年6回を超過','body':f'今年 {exceed45}回。特別条項の回数上限を確認してください'})
    for x in ledger['expiring']:
        level='danger' if x['days_to_expire']<=30 else 'warn'
        alerts.append({'level':level,'title':'有給の時効が近づいています','body':f"{x['remaining_days']:g}日が {x['expires_on']} に消滅予定"})
    ob=ledger.get('five_day_obligation')
    if ob and ob['remaining_required']>0 and ob['days_left']<=90:
        alerts.append({'level':'danger' if ob['days_left']<=30 else 'warn','title':'年5日取得義務の確認','body':f"あと {ob['remaining_required']:g}日 / 期限 {ob['period_end']}"})
    return {'overtime_hours':hours,'overtime_plus_holiday_hours':legal_combined,'annual_overtime_hours':annual_ot,'months_over_45':exceed45,'rolling_max_average':maxavg,'labor_source':source,'agreement_period':{'start':period_start.isoformat(),'end':period_end.isoformat()},'policy':policy,'leave_remaining':ledger['total_remaining'],'leave_ledger':ledger,'shift_changes':changes,'pending':sum(1 for w in own if w.status in ('pending_manager','pending_admin')),'alerts':alerts}

def facility_labor_summary(s,u):
    q=s.query(User).filter(User.facility_id==u.facility_id,User.role.in_(['staff','middle_manager','admin']),User.active==1)
    rows=[]
    for usr in q.all():
        x=labor_summary(s,usr); rows.append({'user_id':usr.id,'name':usr.name,'employee_code':usr.employee_code,**x})
    return {'staff':rows,'alerts':[{'name':r['name'],**a} for r in rows for a in r['alerts']], 'totals':{'overtime_hours':round(sum(r['overtime_hours'] for r in rows),1),'annual_overtime_hours':round(sum(r['annual_overtime_hours'] for r in rows),1),'shift_changes':sum(r['shift_changes'] for r in rows),'pending':sum(r['pending'] for r in rows),'expiring_leave_days':round(sum(sum(x['remaining_days'] for x in r['leave_ledger']['expiring']) for r in rows),1)}}

# -------- workflow --------
def workflow_json(s,w):
    applicant=s.get(User,w.applicant_id); return {'id':w.id,'request_type':w.request_type,'status':w.status,'request_date':w.request_date,'start_time':w.start_time,'end_time':w.end_time,'payload':json.loads(w.payload_json or '{}'),'comment':w.applicant_comment,'applicant_id':w.applicant_id,'applicant_name':applicant.name if applicant else '','created_at':dt_iso(w.created_at),'updated_at':dt_iso(w.updated_at)}
def workflow_list(s,u,own=False,approval=False):
    q=s.query(WorkflowRequest)
    if own: q=q.filter(WorkflowRequest.applicant_id==u.id)
    elif approval:
        q=q.filter(WorkflowRequest.facility_id==u.facility_id)
        if u.role=='middle_manager':
            q=q.filter(WorkflowRequest.status=='pending_manager')
            if u.department_id: q=q.filter(WorkflowRequest.department_id==u.department_id)
        elif u.role=='admin': q=q.filter(WorkflowRequest.status.in_(['pending_admin','pending_manager']))
    return [workflow_json(s,w) for w in q.order_by(WorkflowRequest.id.desc()).limit(100)]

def validate_leave_request(s,u,inp):
    if inp.request_type!='有給申請': return
    reqd=safe_date(inp.request_date)
    if not reqd: raise HTTPException(400,'対象日が不正です')
    days=float((inp.payload or {}).get('leave_days',1) or 1)
    if days<=0: raise HTTPException(400,'有給取得日数が不正です')
    ledger=leave_ledger_summary(s,u)
    usable=sum(x['remaining_days'] for x in ledger['lots'] if not safe_date(x['expires_on']) or safe_date(x['expires_on'])>=reqd)
    if days>usable+1e-9: raise HTTPException(400,f'有給残数が不足しています（利用可能 {usable:g} 日）')
    # 1日未満は時間単位/半日単位を含む。0.125は8h勤務1時間相当のUI既定値。
    if days not in (1.0,0.5) and days<1:
        policy=labor_policy(s,u); gd=ledger.get('current_grant') or {}; start=safe_date(gd.get('grant_date')) if gd else date(reqd.year,1,1); end=add_years(start,1)-timedelta(days=1) if start else date(reqd.year,12,31)
        used=0.0
        for w in s.query(WorkflowRequest).filter_by(applicant_id=u.id,request_type='有給申請',status='approved').all():
            d=safe_date(w.request_date); p=json.loads(w.payload_json or '{}') or {}; x=float(p.get('leave_days',0) or 0)
            if d and start<=d<=end and x not in (1.0,0.5) and x<1: used+=x
        if used+days>policy['time_leave_max_days']+1e-9: raise HTTPException(400,'時間単位年休は労使協定に基づく年5日相当の上限を超えるため申請できません')

class WorkflowIn(BaseModel): request_type:str; request_date:str; start_time:str=''; end_time:str=''; payload:dict={}; comment:str=''
@app.post('/api/workflows')
def create_workflow(inp:WorkflowIn,request:Request):
    with SessionLocal() as s:
        u=current_user(request,s); require_roles(u,'staff','middle_manager','admin')
        validate_leave_request(s,u,inp)
        status='pending_manager' if u.role=='staff' else ('pending_admin' if u.role=='middle_manager' else 'approved')
        w=WorkflowRequest(facility_id=u.facility_id,department_id=u.department_id,applicant_id=u.id,request_type=inp.request_type,status=status,request_date=inp.request_date,start_time=inp.start_time,end_time=inp.end_time,payload_json=json.dumps(inp.payload,ensure_ascii=False),applicant_comment=inp.comment); s.add(w); s.flush();
        if status=='approved': _apply_leave_to_grants(s,w)
        audit(s,request,u,'workflow','submit','workflow_request',w.id,{'type':inp.request_type},u.facility_id)
        if status=='pending_manager':
            approvers=users_by_roles(s,['middle_manager'],u.facility_id,u.department_id)
            if not approvers: approvers=users_by_roles(s,['admin'],u.facility_id)
        elif status=='pending_admin': approvers=users_by_roles(s,['admin'],u.facility_id)
        else: approvers=[]
        emit_notification_event(s,'workflow.submitted',approvers,'承認が必要な職員申請があります',f'{inp.request_type}の申請が届いています。ITTOKAIで確認してください。','workflow',w.id,u.facility_id,u.id,f'workflow.submitted:{w.id}:{status}')
        s.commit(); return {'id':w.id,'status':w.status}

class ApprovalIn(BaseModel): action:str; comment:str=''
@app.post('/api/workflows/{rid}/action')
def workflow_action(rid:int,inp:ApprovalIn,request:Request):
    with SessionLocal() as s:
        u=current_user(request,s); require_roles(u,'middle_manager','admin'); w=s.get(WorkflowRequest,rid)
        if not w: raise HTTPException(404,'申請がありません')
        facility_scope(u,w.facility_id)
        if u.role=='middle_manager' and u.department_id and w.department_id!=u.department_id: raise HTTPException(403,'他部署の申請は承認できません')
        if w.applicant_id==u.id: raise HTTPException(400,'自分の申請は承認できません')
        if w.status not in ('pending_manager','pending_admin'): raise HTTPException(400,'この申請はすでに処理済みです')
        if u.role=='middle_manager' and w.status!='pending_manager': raise HTTPException(400,'中間管理職の承認対象ではありません')
        if inp.action not in ('approve','reject','return'): raise HTTPException(400,'操作が不正です')
        if inp.action in ('reject','return') and not inp.comment.strip(): raise HTTPException(400,'否認・差戻しはコメント必須です')
        if inp.action=='approve':
            if u.role=='middle_manager': w.status='pending_admin'; w.manager_id=u.id
            else: w.status='approved'; w.admin_id=u.id
            _apply_leave_to_grants(s,w)
        elif inp.action=='reject': w.status='rejected'
        else: w.status='returned'
        w.updated_at=now_dt(); s.add(WorkflowApproval(request_id=w.id,approver_id=u.id,approver_role=u.role,action=inp.action,comment=inp.comment))
        if w.status=='pending_admin':
            admins=users_by_roles(s,['admin'],w.facility_id)
            emit_notification_event(s,'workflow.forwarded',admins,'承認待ち申請があります',f'{w.request_type}が中間管理職から回付されました。','workflow',w.id,w.facility_id,u.id,f'workflow.forwarded:{w.id}:admin')
        else:
            body='申請が承認されました。' if w.status=='approved' else ('申請が差し戻されました。' if w.status=='returned' else '申請が否認されました。')
            emit_notification_event(s,'workflow.result',[w.applicant_id],'申請結果が更新されました',body,'workflow',w.id,w.facility_id,u.id,f'workflow.result:{w.id}:{w.status}')
        audit(s,request,u,'workflow',inp.action,'workflow_request',w.id,{'comment':inp.comment},w.facility_id); s.commit(); return {'ok':True,'status':w.status}

@app.get('/api/workflows/export.csv')
def workflow_export(request:Request,scope:str='mine'):
    with SessionLocal() as s:
        u=current_user(request,s)
        q=s.query(WorkflowRequest,User).join(User,WorkflowRequest.applicant_id==User.id)
        if scope=='facility':
            require_roles(u,'middle_manager','admin','hq')
            if u.role!='hq': q=q.filter(WorkflowRequest.facility_id==u.facility_id)
        else: q=q.filter(WorkflowRequest.applicant_id==u.id)
        out=io.StringIO(); w=csv.writer(out)
        w.writerow(['申請ID','職員番号','氏名','申請種別','対象日','開始','終了','状態','申請理由','申請日時','更新日時','承認履歴'])
        for r,usr in q.order_by(WorkflowRequest.id.desc()).all():
            hist=s.query(WorkflowApproval).filter_by(request_id=r.id).order_by(WorkflowApproval.id).all()
            hs=' | '.join(f'{h.approver_role}:{h.action}:{h.comment}:{dt_iso(h.created_at)}' for h in hist)
            w.writerow([r.id,usr.employee_code,usr.name,r.request_type,r.request_date,r.start_time,r.end_time,r.status,r.applicant_comment,dt_iso(r.created_at),dt_iso(r.updated_at),hs])
        audit(s,request,u,'workflow','export_csv','workflow_request','',{'scope':scope},u.facility_id); s.commit()
        return Response(content=('\ufeff'+out.getvalue()).encode('utf-8'),media_type='text/csv; charset=utf-8',headers={'Content-Disposition':'attachment; filename="ittokai_workflow_audit.csv"'})

@app.get('/api/workflows/audit-summary')
def workflow_audit_summary(request:Request):
    with SessionLocal() as s:
        u=current_user(request,s); require_roles(u,'middle_manager','admin','hq')
        if u.role=='hq':
            return {'note':'法人横断は施設選択後に集計','facilities':hq_dashboard(s)['facilities']}
        return facility_labor_summary(s,u)

class LaborPolicyIn(BaseModel):
    agreement_year_start:str='04-01'; monthly_overtime_limit:float=45; annual_overtime_limit:float=360; special_annual_limit:float=720; combined_month_limit:float=100; rolling_average_limit:float=80; over45_months_limit:int=6; annual_leave_required_days:float=5; time_leave_max_days:float=5

@app.get('/api/labor/policy')
def get_labor_policy(request:Request):
    with SessionLocal() as s:
        u=current_user(request,s); require_roles(u,'middle_manager','admin','hq'); return labor_policy(s,u)

@app.post('/api/labor/policy')
def set_labor_policy(inp:LaborPolicyIn,request:Request):
    with SessionLocal() as s:
        u=current_user(request,s); require_roles(u,'admin','hq')
        if not u.facility_id: raise HTTPException(400,'法人本部からは対象施設を指定する設定画面を今後追加してください')
        p=s.query(LaborPolicy).filter_by(facility_id=u.facility_id).first() or LaborPolicy(facility_id=u.facility_id)
        for k,v in inp.model_dump().items(): setattr(p,k,v)
        p.updated_at=now_dt(); s.add(p); audit(s,request,u,'workflow','labor_policy_update','facility',u.facility_id,inp.model_dump(),u.facility_id); s.commit(); return {'ok':True,'policy':labor_policy(s,u)}

@app.get('/api/labor/leave-ledger.csv')
def leave_ledger_export(request:Request,scope:str='mine'):
    with SessionLocal() as s:
        u=current_user(request,s); users=[u]
        if scope=='facility':
            require_roles(u,'middle_manager','admin','hq')
            q=s.query(User).filter(User.role.in_(['staff','middle_manager','admin']),User.active==1)
            if u.role!='hq': q=q.filter(User.facility_id==u.facility_id)
            users=q.all()
        out=io.StringIO(); w=csv.writer(out); w.writerow(['職員番号','氏名','付与日','付与日数','取得済日数','残日数','時効予定日','消滅まで日数','データソース'])
        today=datetime.now(JST).date()
        for usr in users:
            for g in s.query(LeaveGrant).filter_by(user_id=usr.id).order_by(LeaveGrant.grant_date).all():
                exp=safe_date(g.expires_on); days_to=(exp-today).days if exp else ''
                w.writerow([usr.employee_code,usr.name,g.grant_date,round(float(g.granted_days or 0),1),round(float(g.taken_days or 0),1),round(float(g.remaining_days or 0),1),g.expires_on,days_to,g.source])
        audit(s,request,u,'workflow','export_leave_ledger','leave_grant','',{'scope':scope},u.facility_id); s.commit()
        return Response(content=('\ufeff'+out.getvalue()).encode('utf-8'),media_type='text/csv; charset=utf-8',headers={'Content-Disposition':'attachment; filename="ittokai_leave_ledger.csv"'})

@app.get('/api/labor/compliance.csv')
def labor_compliance_export(request:Request):
    with SessionLocal() as s:
        u=current_user(request,s); require_roles(u,'middle_manager','admin','hq')
        q=s.query(User).filter(User.role.in_(['staff','middle_manager','admin']),User.active==1)
        if u.role!='hq': q=q.filter(User.facility_id==u.facility_id)
        out=io.StringIO(); w=csv.writer(out); w.writerow(['職員番号','氏名','今月時間外','今月時間外+休日','年間時間外','45h超月数','2-6か月最大平均','有給残数','前回付与残','今回付与残','消滅予定日','5日義務残','データソース','アラート'])
        for usr in q.all():
            x=labor_summary(s,usr); l=x['leave_ledger']; prev=l.get('previous_grant') or {}; cur=l.get('current_grant') or {}; ob=l.get('five_day_obligation') or {}
            w.writerow([usr.employee_code,usr.name,x['overtime_hours'],x['overtime_plus_holiday_hours'],x['annual_overtime_hours'],x['months_over_45'],x['rolling_max_average'],x['leave_remaining'],prev.get('remaining_days',''),cur.get('remaining_days',''),prev.get('expires_on',''),ob.get('remaining_required',''),x['labor_source'],' | '.join(a['title'] for a in x['alerts'])])
        audit(s,request,u,'workflow','export_compliance','work_time','',{},u.facility_id); s.commit()
        return Response(content=('\ufeff'+out.getvalue()).encode('utf-8'),media_type='text/csv; charset=utf-8',headers={'Content-Disposition':'attachment; filename="ittokai_labor_compliance.csv"'})

# -------- visits --------
def visit_rules_for_facility(s,fid): return [{'id':r.id,'place_name':r.place_name,'weekday':r.weekday,'start_time':r.start_time,'end_time':r.end_time,'slot_minutes':r.slot_minutes,'capacity':r.capacity,'max_visitors':r.max_visitors,'booking_days_ahead':r.booking_days_ahead,'cutoff_hours':r.cutoff_hours,'approval_mode':r.approval_mode,'active':r.active} for r in s.query(VisitRule).filter_by(facility_id=fid).all()]
def visit_json(s,v):
    r=s.get(Resident,v.resident_id); u=s.get(User,v.family_user_id); return {'id':v.id,'facility_id':v.facility_id,'resident_id':v.resident_id,'resident_name':r.name if r else '','family_user_id':v.family_user_id,'family_name':u.name if u else '','visit_date':v.visit_date,'start_time':v.start_time,'end_time':v.end_time,'place_name':v.place_name,'visitor_count':v.visitor_count,'representative':v.representative,'phone':v.phone,'notes':v.notes,'status':v.status,'checked_in_at':dt_iso(v.checked_in_at)}
def visit_list_for_user(s,u): return [visit_json(s,v) for v in s.query(VisitReservation).filter_by(family_user_id=u.id).order_by(VisitReservation.visit_date.desc()).limit(100)]
def visit_list_for_facility(s,fid): return [visit_json(s,v) for v in s.query(VisitReservation).filter_by(facility_id=fid).order_by(VisitReservation.visit_date.desc()).limit(200)]

@app.get('/api/visits/availability')
def visit_availability(request:Request,facility_id:int,visit_date:str):
    with SessionLocal() as s:
        u=current_user(request,s); require_roles(u,'family','staff','middle_manager','admin','hq'); facility_scope(u,facility_id)
        try: d=date.fromisoformat(visit_date)
        except ValueError: raise HTTPException(400,'日付が不正です')
        if d < datetime.now(JST).date(): raise HTTPException(400,'過去の日付は予約できません')
        rules=s.query(VisitRule).filter_by(facility_id=facility_id,weekday=d.weekday(),active=1).all(); slots=[]
        blocks=s.query(VisitBlock).filter_by(facility_id=facility_id,block_date=visit_date).all()
        for r in rules:
            sh,sm=map(int,r.start_time.split(':')); eh,em=map(int,r.end_time.split(':')); cur=datetime.combine(d,time(sh,sm)); end=datetime.combine(d,time(eh,em))
            while cur+timedelta(minutes=r.slot_minutes)<=end:
                st=cur.strftime('%H:%M'); et=(cur+timedelta(minutes=r.slot_minutes)).strftime('%H:%M'); blocked=any(not (et<=b.start_time or st>=b.end_time) for b in blocks)
                # Apply booking horizon/cutoff server-side.
                days_ahead=(d-datetime.now(JST).date()).days
                cutoff_at=datetime.combine(d,time(sh,sm),tzinfo=JST)-timedelta(hours=r.cutoff_hours)
                outside=days_ahead>r.booking_days_ahead or datetime.now(JST)>cutoff_at
                count=s.query(VisitReservation).filter_by(facility_id=facility_id,visit_date=visit_date,start_time=st,place_name=r.place_name).filter(VisitReservation.status.in_(['confirmed','pending'])).count()
                slots.append({'place_name':r.place_name,'start_time':st,'end_time':et,'capacity':r.capacity,'remaining':0 if blocked or outside else max(0,r.capacity-count),'approval_mode':r.approval_mode,'max_visitors':r.max_visitors}); cur+=timedelta(minutes=r.slot_minutes)
        return {'date':visit_date,'slots':slots}

class VisitIn(BaseModel): resident_id:int; visit_date:str; start_time:str; end_time:str; place_name:str; visitor_count:int=1; representative:str; phone:str=''; notes:str=''
@app.post('/api/visits')
def create_visit(inp:VisitIn,request:Request):
    with SessionLocal() as s:
        u=current_user(request,s); require_roles(u,'family'); res=s.get(Resident,inp.resident_id)
        if not res: raise HTTPException(404,'利用者がいません')
        link=s.query(FamilyLink).filter_by(user_id=u.id,resident_id=res.id).first()
        if not link: raise HTTPException(403,'この利用者の予約権限がありません')
        avail=visit_availability(request,res.facility_id,inp.visit_date)['slots']; slot=next((x for x in avail if x['start_time']==inp.start_time and x['place_name']==inp.place_name),None)
        if not slot or slot['remaining']<=0: raise HTTPException(400,'この時間枠は予約できません')
        if inp.visitor_count<1 or inp.visitor_count>int(slot.get('max_visitors') or 4): raise HTTPException(400,'来訪人数が上限を超えています')
        if not inp.representative.strip(): raise HTTPException(400,'来訪代表者名が必要です')
        status='confirmed' if slot['approval_mode']=='auto' else 'pending'; v=VisitReservation(facility_id=res.facility_id,resident_id=res.id,family_user_id=u.id,visit_date=inp.visit_date,start_time=inp.start_time,end_time=slot['end_time'],place_name=inp.place_name,visitor_count=inp.visitor_count,representative=inp.representative.strip(),phone=inp.phone,notes=inp.notes,status=status); s.add(v); s.flush()
        admins=users_by_roles(s,['admin'],res.facility_id)
        emit_notification_event(s,'visit.requested',admins,'面会予約が入りました',f'{inp.visit_date} {inp.start_time} の面会予約です。','visits',v.id,res.facility_id,u.id,f'visit.requested:{v.id}')
        if status=='confirmed': emit_notification_event(s,'visit.confirmed',[u.id],'面会予約が確定しました',f'{inp.visit_date} {inp.start_time} の予約が確定しています。','visits',v.id,res.facility_id,None,f'visit.confirmed:{v.id}')
        audit(s,request,u,'visits','book','visit',v.id,{},res.facility_id); s.commit(); return {'id':v.id,'status':status}

@app.post('/api/visits/{vid}/cancel')
def cancel_visit(vid:int,request:Request):
    with SessionLocal() as s:
        u=current_user(request,s); require_roles(u,'family','staff','middle_manager','admin','hq'); v=s.get(VisitReservation,vid)
        if not v: raise HTTPException(404,'予約がありません')
        if u.role=='family' and v.family_user_id!=u.id: raise HTTPException(403,'権限がありません')
        facility_scope(u,v.facility_id); v.status='cancelled'
        recipients=users_by_roles(s,['admin'],v.facility_id) if u.role=='family' else [v.family_user_id]
        emit_notification_event(s,'visit.cancelled',recipients,'面会予約がキャンセルされました','ITTOKAIで面会予約の状態をご確認ください。','visits',vid,v.facility_id,u.id,f'visit.cancelled:{vid}:{u.role}')
        audit(s,request,u,'visits','cancel','visit',vid,{},v.facility_id); s.commit(); return {'ok':True}


class VisitStatusIn(BaseModel):
    status:str

@app.post('/api/visits/{vid}/status')
def update_visit_status(vid:int,inp:VisitStatusIn,request:Request):
    with SessionLocal() as s:
        u=current_user(request,s); require_roles(u,'admin','middle_manager'); v=s.get(VisitReservation,vid)
        if not v: raise HTTPException(404,'予約がありません')
        facility_scope(u,v.facility_id)
        if inp.status not in ('confirmed','cancelled'): raise HTTPException(400,'状態が不正です')
        if v.status not in ('pending','confirmed'): raise HTTPException(400,'この予約は変更できません')
        v.status=inp.status
        et='visit.confirmed' if inp.status=='confirmed' else 'visit.cancelled'; title='面会予約が確定しました' if inp.status=='confirmed' else '面会予約がキャンセルされました'
        emit_notification_event(s,et,[v.family_user_id],title,'ITTOKAIで面会予約の詳細をご確認ください。','visits',v.id,v.facility_id,u.id,f'{et}:{v.id}:facility')
        audit(s,request,u,'visits','status','visit',v.id,{'status':inp.status},v.facility_id); s.commit(); return {'ok':True,'status':v.status}

class VisitRuleIn(BaseModel):
    place_name:str='面会室A'; weekday:int; start_time:str='14:00'; end_time:str='17:00'; slot_minutes:int=30; capacity:int=1; max_visitors:int=4; booking_days_ahead:int=30; cutoff_hours:int=24; approval_mode:str='auto'
@app.post('/api/visit-rules')
def create_visit_rule(inp:VisitRuleIn,request:Request):
    with SessionLocal() as s:
        u=current_user(request,s); require_roles(u,'admin'); r=VisitRule(facility_id=u.facility_id,**inp.model_dump()); s.add(r); s.flush(); audit(s,request,u,'visits','rule_create','visit_rule',r.id,inp.model_dump(),u.facility_id); s.commit(); return {'id':r.id,'ok':True}

@app.delete('/api/visit-rules/{rid}')
def delete_visit_rule(rid:int,request:Request):
    with SessionLocal() as s:
        u=current_user(request,s); require_roles(u,'admin'); r=s.get(VisitRule,rid)
        if not r or r.facility_id!=u.facility_id: raise HTTPException(404,'設定がありません')
        s.delete(r); audit(s,request,u,'visits','rule_delete','visit_rule',rid,{},u.facility_id); s.commit(); return {'ok':True}

class VisitBlockIn(BaseModel):
    block_date:str; start_time:str='00:00'; end_time:str='23:59'; reason:str=''
@app.post('/api/visit-blocks')
def create_visit_block(inp:VisitBlockIn,request:Request):
    with SessionLocal() as s:
        u=current_user(request,s); require_roles(u,'admin'); b=VisitBlock(facility_id=u.facility_id,**inp.model_dump()); s.add(b); s.flush(); audit(s,request,u,'visits','block_create','visit_block',b.id,inp.model_dump(),u.facility_id); s.commit(); return {'id':b.id,'ok':True}

class ProxyVisitIn(VisitIn):
    family_user_id:int
@app.post('/api/visits/proxy')
def proxy_visit(inp:ProxyVisitIn,request:Request):
    with SessionLocal() as s:
        u=current_user(request,s); require_roles(u,'admin'); res=s.get(Resident,inp.resident_id); fam=s.get(User,inp.family_user_id)
        if not res or res.facility_id!=u.facility_id or not fam or fam.facility_id!=u.facility_id: raise HTTPException(400,'利用者または家族が不正です')
        v=VisitReservation(facility_id=u.facility_id,resident_id=res.id,family_user_id=fam.id,visit_date=inp.visit_date,start_time=inp.start_time,end_time=inp.end_time,place_name=inp.place_name,visitor_count=inp.visitor_count,representative=inp.representative,phone=inp.phone,notes=inp.notes,status='confirmed'); s.add(v); s.flush(); emit_notification_event(s,'visit.confirmed',[fam.id],'面会予約を登録しました',f'{inp.visit_date} {inp.start_time} {inp.place_name}','visits',v.id,u.facility_id,u.id,f'visit.proxy_confirmed:{v.id}'); audit(s,request,u,'visits','proxy_book','visit',v.id,{},u.facility_id); s.commit(); return {'id':v.id,'status':'confirmed'}

# -------- procurement / product & inventory master v20.6 --------
def product_meta_for(s,pid):
    return s.query(ProductMeta).filter_by(product_id=pid).first()

def product_public_json(s,p):
    m=product_meta_for(s,p.id)
    return {'id':p.id,'sku':p.sku,'name':p.name,'category':p.category,'specification':p.specification,'price':(m.sale_price if m and m.sale_price else p.standard_price),'unit':p.unit,'short_message':m.short_message if m else '', 'recommended':m.recommended if m else 0,'active':p.active,'image_url':f'/api/products/{p.id}/image'}

def partner_public_products(s,limit=50):
    out=[]
    q=s.query(PartnerProduct,PartnerStore).join(PartnerStore,PartnerProduct.store_id==PartnerStore.id).filter(PartnerProduct.active==1,PartnerStore.active==1).order_by(PartnerProduct.recommended.desc(),PartnerProduct.id.desc()).limit(limit)
    for p,st in q:
        out.append({'id':p.id,'store_id':st.id,'store_name':st.store_name,'store_category':st.category,'sku':p.sku,'name':p.name,'short_message':p.short_message,'description':p.description,'category':p.category,'specification':p.specification,'unit':p.unit,'price':p.price,'image_url':p.image_url,'recommended':p.recommended})
    return out

def open_incoming_qty(s,fid,pid):
    q=(s.query(func.coalesce(func.sum(PurchaseOrderItem.qty-PurchaseOrderItem.received_qty),0))
       .join(PurchaseOrder,PurchaseOrderItem.order_id==PurchaseOrder.id)
       .filter(PurchaseOrder.facility_id==fid,PurchaseOrderItem.product_id==pid,PurchaseOrder.status.in_(['submitted','confirmed','shipped','partially_received'])))
    return float(q.scalar() or 0)

def stock_status(s,inv):
    incoming=open_incoming_qty(s,inv.facility_id,inv.product_id)
    if inv.on_hand<=0 and incoming<=0: return '欠品'
    if incoming>0: return '発注済'
    if inv.on_hand<inv.reorder_point: return '発注必要'
    if inv.on_hand<=inv.safety_stock: return '残りわずか'
    return '正常'

def suggested_reorder_qty(s,inv,p):
    pol=s.query(FacilityStockPolicy).filter_by(facility_id=inv.facility_id,product_id=p.id).first()
    incoming=open_incoming_qty(s,inv.facility_id,p.id)
    if incoming>0: return 0
    if pol and pol.reorder_qty>0: q=pol.reorder_qty
    else: q=max(p.min_order_qty, max(0,inv.target_stock-inv.on_hand))
    pack=max(1,int(p.pack_size or 1)); q=max(float(p.min_order_qty or 1),float(q))
    if pack>1: q=((int(q+pack-1)//pack)*pack)
    return float(q)

def _create_po(s,fid,sup,items,approved_by=None,status='submitted'):
    no='PO-'+datetime.now(JST).strftime('%Y%m%d%H%M%S%f')+'-'+secrets.token_hex(2).upper(); o=PurchaseOrder(order_no=no,facility_id=fid,supplier_id=sup.id,status=status,approved_by=approved_by,ordered_at=now_dt() if status=='submitted' else None); s.add(o); s.flush(); total=0
    for p,q in items:
        fp=s.query(FacilityPrice).filter_by(facility_id=fid,product_id=p.id).first(); price=fp.purchase_price if fp else (product_meta_for(s,p.id).purchase_price if product_meta_for(s,p.id) else p.standard_price); total+=q*price; s.add(PurchaseOrderItem(order_id=o.id,product_id=p.id,qty=q,unit_price=price))
    o.total_amount=total; return o

def evaluate_auto_orders(s,facility_id=None,product_id=None,actor_id=None):
    q=s.query(Inventory,Product).join(Product,Inventory.product_id==Product.id)
    if facility_id: q=q.filter(Inventory.facility_id==facility_id)
    if product_id: q=q.filter(Inventory.product_id==product_id)
    sup=s.query(Supplier).filter_by(code='CAREDO').first(); created=[]
    if not sup: return created
    for inv,p in q.all():
        pol=s.query(FacilityStockPolicy).filter_by(facility_id=inv.facility_id,product_id=p.id).first()
        if not pol or pol.auto_order_mode=='off' or inv.on_hand>=inv.reorder_point: continue
        if open_incoming_qty(s,inv.facility_id,p.id)>0: continue
        existing=s.query(AutoOrderEvent).filter_by(facility_id=inv.facility_id,product_id=p.id,status='pending_approval').first()
        if existing: continue
        qty=suggested_reorder_qty(s,inv,p)
        if qty<=0: continue
        if pol.auto_order_mode=='auto':
            o=_create_po(s,inv.facility_id,sup,[(p,qty)],approved_by=actor_id,status='submitted'); ev=AutoOrderEvent(facility_id=inv.facility_id,product_id=p.id,order_id=o.id,mode='auto',status='ordered',suggested_qty=qty,on_hand_at_trigger=inv.on_hand,reorder_point=inv.reorder_point,resolved_at=now_dt()); pol.last_ordered_at=now_dt(); created.append({'event_id':None,'order_id':o.id,'mode':'auto'})
        else:
            ev=AutoOrderEvent(facility_id=inv.facility_id,product_id=p.id,mode='approval',status='pending_approval',suggested_qty=qty,on_hand_at_trigger=inv.on_hand,reorder_point=inv.reorder_point); created.append({'event_id':None,'mode':'approval'})
        s.add(ev); s.flush(); created[-1]['event_id']=ev.id
        recipients=users_by_roles(s,['admin'],inv.facility_id)+users_by_roles(s,['caredo'])
        if pol.auto_order_mode=='auto':
            emit_notification_event(s,'inventory.auto_order_created',recipients,'自動発注が作成されました',f'{p.name} の自動発注が作成されました。','procurement',o.id,inv.facility_id,actor_id,f'inventory.auto_order_created:{ev.id}')
        else:
            emit_notification_event(s,'inventory.auto_order_candidate',recipients,'自動発注候補があります',f'{p.name} が発注点を下回りました。発注候補をご確認ください。','procurement',ev.id,inv.facility_id,actor_id,f'inventory.auto_order_candidate:{ev.id}')
    return created

def procurement_snapshot(s,fid):
    inv=s.query(Inventory,Product,Supplier).join(Product,Inventory.product_id==Product.id).join(Supplier,Product.supplier_id==Supplier.id).filter(Inventory.facility_id==fid).all()
    inventory=[]; suggestions=[]
    for i,p,sup in inv:
        pol=s.query(FacilityStockPolicy).filter_by(facility_id=fid,product_id=p.id).first()
        incoming=open_incoming_qty(s,fid,p.id); status=stock_status(s,i); q=suggested_reorder_qty(s,i,p) if i.on_hand<i.reorder_point else 0
        row={'inventory_id':i.id,'product_id':p.id,'sku':p.sku,'name':p.name,'category':p.category,'supplier':sup.name,'supplier_type':sup.supplier_type,'on_hand':i.on_hand,'safety_stock':i.safety_stock,'reorder_point':i.reorder_point,'target_stock':i.target_stock,'avg_daily_usage':i.avg_daily_usage,'unit':p.unit,'incoming_qty':incoming,'status':status,'reorder_qty':pol.reorder_qty if pol else 0,'max_stock':pol.max_stock if pol else 0,'auto_order_mode':pol.auto_order_mode if pol else 'off','last_ordered_at':dt_iso(pol.last_ordered_at) if pol else None,'last_received_at':dt_iso(pol.last_received_at) if pol else None}; inventory.append(row)
        if i.on_hand < i.reorder_point and incoming<=0: suggestions.append(dict(row,suggested_qty=q))
    orders=[order_json(s,o) for o in s.query(PurchaseOrder).filter_by(facility_id=fid).order_by(PurchaseOrder.id.desc()).limit(100)]
    invoices=[invoice_json(s,x) for x in s.query(Invoice).filter_by(facility_id=fid).order_by(Invoice.id.desc()).limit(100)]
    events=[auto_event_json(s,e) for e in s.query(AutoOrderEvent).filter_by(facility_id=fid).order_by(AutoOrderEvent.id.desc()).limit(100)]
    return {'inventory':inventory,'suggestions':suggestions,'orders':orders,'invoices':invoices,'auto_order_events':events}

def order_json(s,o):
    sup=s.get(Supplier,o.supplier_id); items=[]
    for it in s.query(PurchaseOrderItem).filter_by(order_id=o.id).all():
        p=s.get(Product,it.product_id); items.append({'id':it.id,'product_id':it.product_id,'sku':p.sku if p else '','name':p.name if p else '','qty':it.qty,'unit_price':it.unit_price,'received_qty':it.received_qty})
    return {'id':o.id,'order_no':o.order_no,'facility_id':o.facility_id,'supplier_id':o.supplier_id,'supplier':sup.name if sup else '','status':o.status,'total_amount':o.total_amount,'ordered_at':dt_iso(o.ordered_at),'confirmed_at':dt_iso(o.confirmed_at),'shipped_at':dt_iso(o.shipped_at),'received_at':dt_iso(o.received_at),'items':items}

def invoice_json(s,inv): return {'id':inv.id,'invoice_no':inv.invoice_no,'facility_id':inv.facility_id,'order_id':inv.order_id,'invoice_date':inv.invoice_date,'total_amount':inv.total_amount,'status':inv.status,'match_status':inv.match_status,'differences':json.loads(inv.difference_json or '{}')}

def auto_event_json(s,e):
    f=s.get(Facility,e.facility_id); p=s.get(Product,e.product_id)
    return {'id':e.id,'facility_id':e.facility_id,'facility_name':f.name if f else '','product_id':e.product_id,'sku':p.sku if p else '','product_name':p.name if p else '','order_id':e.order_id,'mode':e.mode,'status':e.status,'suggested_qty':e.suggested_qty,'on_hand_at_trigger':e.on_hand_at_trigger,'reorder_point':e.reorder_point,'created_at':dt_iso(e.created_at),'resolved_at':dt_iso(e.resolved_at)}

def caredo_product_json(s,p):
    m=product_meta_for(s,p.id); buy=m.purchase_price if m and m.purchase_price else p.standard_price; sale=m.sale_price if m and m.sale_price else p.standard_price; margin=sale-buy; rate=(margin/sale*100) if sale else 0
    return {'id':p.id,'sku':p.sku,'name':p.name,'specification':p.specification,'category':p.category,'unit':p.unit,'pack_size':p.pack_size,'tax_rate':p.tax_rate,'standard_price':p.standard_price,'purchase_price':buy,'sale_price':sale,'gross_profit':round(margin,2),'gross_margin_rate':round(rate,1),'min_order_qty':p.min_order_qty,'lead_time_days':p.lead_time_days,'inventory_managed':p.inventory_managed,'active':p.active,'short_message':m.short_message if m else '','description':m.description if m else '','recommended':m.recommended if m else 0,'sale_start':m.sale_start if m else '','sale_end':m.sale_end if m else '','display_order':m.display_order if m else 0,'image_url':f'/api/products/{p.id}/image'}

def customer_order_json(s,o,merchant_view=False):
    fac=s.get(Facility,o.facility_id); items=[]
    for it in s.query(CustomerOrderItem).filter_by(order_id=o.id).order_by(CustomerOrderItem.id).all():
        items.append({'id':it.id,'product_kind':it.product_kind,'product_id':it.product_id,'sku':it.sku,'name':it.name,'qty':it.qty,'unit_price':it.unit_price,'amount':round(float(it.qty or 0)*float(it.unit_price or 0),2)})
    out={'id':o.id,'order_no':o.order_no,'merchant_type':o.merchant_type,'status':o.status,'total_amount':o.total_amount,'gift_message':o.gift_message,'delivery_date':o.delivery_date,'facility_id':o.facility_id,'facility_name':fac.name if fac else '','created_at':dt_iso(o.created_at),'accepted_at':dt_iso(o.accepted_at),'shipped_at':dt_iso(o.shipped_at),'delivered_at':dt_iso(o.delivered_at),'items':items}
    if not merchant_view:
        r=s.get(Resident,o.resident_id) if o.resident_id else None; out['resident_id']=o.resident_id; out['resident_name']=r.name if r else ''
    return out

def customer_orders_for_family(s,user_id):
    return [customer_order_json(s,o,False) for o in s.query(CustomerOrder).filter_by(family_user_id=user_id).order_by(CustomerOrder.id.desc()).limit(100)]

def customer_orders_for_caredo(s):
    return [customer_order_json(s,o,True) for o in s.query(CustomerOrder).filter_by(merchant_type='caredo').order_by(CustomerOrder.id.desc()).limit(200)]

def customer_orders_for_sponsor(s,user_id):
    stores=[x.id for x in s.query(PartnerStore).filter_by(owner_user_id=user_id).all()]
    if not stores: return []
    return [customer_order_json(s,o,True) for o in s.query(CustomerOrder).filter(CustomerOrder.merchant_type=='sponsor',CustomerOrder.store_id.in_(stores)).order_by(CustomerOrder.id.desc()).limit(200)]

def caredo_snapshot(s):
    sup=s.query(Supplier).filter_by(code='CAREDO').first(); orders=[order_json(s,o) for o in s.query(PurchaseOrder).filter_by(supplier_id=sup.id).order_by(PurchaseOrder.id.desc()).limit(300)] if sup else []
    facilities=[facility_json(f) for f in s.query(Facility).filter_by(active=1).all()]
    all_inventory=[]
    for f in s.query(Facility).filter_by(active=1).all(): all_inventory.extend(procurement_snapshot(s,f.id)['inventory'])
    # restore facility labels after aggregation
    for row in all_inventory:
        inv=s.get(Inventory,row['inventory_id']); fac=s.get(Facility,inv.facility_id) if inv else None; row['facility_id']=inv.facility_id if inv else None; row['facility_name']=fac.name if fac else ''
    events=[auto_event_json(s,e) for e in s.query(AutoOrderEvent).order_by(AutoOrderEvent.id.desc()).limit(300)]
    return {'supplier':{'id':sup.id,'name':sup.name,'site_url':sup.site_url} if sup else None,'orders':orders,'products':[caredo_product_json(s,p) for p in s.query(Product).filter_by(supplier_id=sup.id).order_by(Product.id).all()] if sup else [],'facilities':facilities,'facility_inventory':all_inventory,'auto_order_events':events,'customer_orders':customer_orders_for_caredo(s)}

def sponsor_snapshot(s,u):
    stores=s.query(PartnerStore).filter_by(owner_user_id=u.id).order_by(PartnerStore.id).all()
    return {'message':'協賛会社ポータル。利用者個人情報・介護情報は表示しません。','stores':[{'id':x.id,'company_name':x.company_name,'store_name':x.store_name,'category':x.category,'image_url':x.image_url,'address':x.address,'phone':x.phone,'hours':x.hours,'closed_days':x.closed_days,'delivery_area':x.delivery_area,'min_order_amount':x.min_order_amount,'order_cutoff':x.order_cutoff,'delivery_methods':x.delivery_methods,'message':x.message,'active':x.active} for x in stores], 'products':[{'id':p.id,'store_id':p.store_id,'sku':p.sku,'name':p.name,'short_message':p.short_message,'description':p.description,'category':p.category,'specification':p.specification,'unit':p.unit,'price':p.price,'image_url':p.image_url,'recommended':p.recommended,'active':p.active,'sale_start':p.sale_start,'sale_end':p.sale_end} for p in s.query(PartnerProduct).filter(PartnerProduct.store_id.in_([x.id for x in stores] or [-1])).order_by(PartnerProduct.id).all()],'customer_orders':customer_orders_for_sponsor(s,u.id)}

class CustomerOrderIn(BaseModel):
    product_kind:str; product_id:int; qty:float=1; resident_id:Optional[int]=None; gift_message:str=''; delivery_date:str=''

@app.post('/api/shop/orders')
def create_customer_order(inp:CustomerOrderIn,request:Request):
    with SessionLocal() as s:
        u=current_user(request,s); require_roles(u,'family')
        if inp.qty<=0 or inp.qty>100: raise HTTPException(400,'数量が不正です')
        resident_id=inp.resident_id
        if resident_id:
            fl=s.query(FamilyLink).filter_by(user_id=u.id,resident_id=resident_id).first()
            if not fl: raise HTTPException(403,'この利用者への注文権限がありません')
            res=s.get(Resident,resident_id); facility_id=res.facility_id
        else:
            fl=s.query(FamilyLink).filter_by(user_id=u.id,is_primary=1).first() or s.query(FamilyLink).filter_by(user_id=u.id).first()
            if not fl: raise HTTPException(400,'配送先利用者が設定されていません')
            resident_id=fl.resident_id; res=s.get(Resident,resident_id); facility_id=res.facility_id
        merchant_type=''; supplier_id=None; store_id=None; recipients=[]
        if inp.product_kind=='caredo':
            p=s.get(Product,inp.product_id); sup=s.query(Supplier).filter_by(code='CAREDO').first(); m=product_meta_for(s,p.id) if p else None
            if not p or not sup or p.supplier_id!=sup.id or not p.active: raise HTTPException(404,'商品がありません')
            price=float((m.sale_price if m and m.sale_price else p.standard_price) or 0); sku=p.sku; name=p.name; merchant_type='caredo'; supplier_id=sup.id; recipients=users_by_roles(s,['caredo'])
        elif inp.product_kind=='sponsor':
            p=s.get(PartnerProduct,inp.product_id); st=s.get(PartnerStore,p.store_id) if p else None
            if not p or not st or not p.active or not st.active: raise HTTPException(404,'ギフト商品がありません')
            price=float(p.price or 0); sku=p.sku; name=p.name; merchant_type='sponsor'; store_id=st.id; owner=s.get(User,st.owner_user_id); recipients=[owner] if owner and owner.active else []
        else: raise HTTPException(400,'商品種別が不正です')
        no='CO-'+datetime.now(JST).strftime('%Y%m%d%H%M%S%f')+'-'+secrets.token_hex(2).upper()
        o=CustomerOrder(order_no=no,family_user_id=u.id,facility_id=facility_id,resident_id=resident_id,merchant_type=merchant_type,supplier_id=supplier_id,store_id=store_id,status='submitted',total_amount=round(price*inp.qty,2),gift_message=inp.gift_message[:1000],delivery_date=inp.delivery_date[:10]); s.add(o); s.flush()
        s.add(CustomerOrderItem(order_id=o.id,product_kind=inp.product_kind,product_id=inp.product_id,sku=sku,name=name,qty=inp.qty,unit_price=price))
        emit_notification_event(s,'shop.order_submitted',recipients,'新しい注文が入りました',f'注文番号 {o.order_no}。ITTOKAIで内容をご確認ください。','shop',o.id,facility_id,u.id,f'shop.order_submitted:{o.id}',{'merchant_type':merchant_type})
        emit_notification_event(s,'shop.order_status',[u.id],'ご注文を受け付けました',f'{name} の注文を受け付けました。','shop',o.id,facility_id,None,f'shop.order_ack:{o.id}')
        audit(s,request,u,'shop','customer_order_create','customer_order',o.id,{'merchant_type':merchant_type,'amount':o.total_amount},facility_id); s.commit(); return customer_order_json(s,o,False)

class CustomerOrderStatusIn(BaseModel):
    status:str

@app.post('/api/shop/orders/{oid}/status')
def customer_order_status(oid:int,inp:CustomerOrderStatusIn,request:Request):
    with SessionLocal() as s:
        u=current_user(request,s); require_roles(u,'caredo','sponsor','admin'); o=s.get(CustomerOrder,oid)
        if not o: raise HTTPException(404,'注文がありません')
        if u.role=='caredo' and o.merchant_type!='caredo': raise HTTPException(403,'この注文は操作できません')
        if u.role=='sponsor':
            st=s.get(PartnerStore,o.store_id) if o.store_id else None
            if o.merchant_type!='sponsor' or not st or st.owner_user_id!=u.id: raise HTTPException(403,'この注文は操作できません')
        if u.role=='admin' and o.facility_id!=u.facility_id: raise HTTPException(403,'他施設の注文は操作できません')
        allowed={'submitted':{'accepted','cancelled'},'accepted':{'shipped','cancelled'},'shipped':{'delivered'}}
        if inp.status not in allowed.get(o.status,set()): raise HTTPException(400,'現在の状態から変更できません')
        o.status=inp.status; o.updated_at=now_dt()
        if inp.status=='accepted': o.accepted_at=now_dt()
        elif inp.status=='shipped': o.shipped_at=now_dt()
        elif inp.status=='delivered': o.delivered_at=now_dt()
        elif inp.status=='cancelled': o.cancelled_at=now_dt()
        labels={'accepted':'注文を確認しました','shipped':'商品を発送しました','delivered':'商品のお届けが完了しました','cancelled':'注文がキャンセルされました'}
        emit_notification_event(s,'shop.order_status',[o.family_user_id],labels.get(inp.status,'注文状況が更新されました'),'ITTOKAIで注文状況をご確認ください。','shop',o.id,o.facility_id,u.id,f'shop.order_status:{o.id}:{inp.status}')
        audit(s,request,u,'shop','customer_order_status','customer_order',o.id,{'status':inp.status},o.facility_id); s.commit(); return {'ok':True,'status':o.status}

class InventoryAdjust(BaseModel): product_id:int; qty:float; txn_type:str='adjust'; note:str=''
@app.post('/api/inventory/adjust')
def inventory_adjust(inp:InventoryAdjust,request:Request):
    with SessionLocal() as s:
        u=current_user(request,s); require_roles(u,'staff','middle_manager','admin'); inv=s.query(Inventory).filter_by(facility_id=u.facility_id,product_id=inp.product_id).first()
        if not inv: raise HTTPException(404,'在庫がありません')
        inv.on_hand += inp.qty; inv.updated_at=now_dt(); s.add(InventoryTxn(facility_id=u.facility_id,product_id=inp.product_id,txn_type=inp.txn_type,qty=inp.qty,note=inp.note,actor_id=u.id)); created=evaluate_auto_orders(s,u.facility_id,inp.product_id,u.id); audit(s,request,u,'procurement','inventory_adjust','product',inp.product_id,{'qty':inp.qty,'auto_order_events':created},u.facility_id); s.commit(); return {'ok':True,'on_hand':inv.on_hand,'auto_order_events':created}

class OrderIn(BaseModel): items:list[dict]
@app.post('/api/orders')
def create_order(inp:OrderIn,request:Request):
    with SessionLocal() as s:
        u=current_user(request,s); require_roles(u,'middle_manager','admin'); sup=s.query(Supplier).filter_by(code='CAREDO').first()
        if not sup: raise HTTPException(500,'ケア・ドゥ取引先マスタがありません')
        if not inp.items: raise HTTPException(400,'発注明細がありません')
        validated=[]
        for x in inp.items:
            try: pid=int(x['product_id']); q=float(x['qty'])
            except Exception: raise HTTPException(400,'発注明細が不正です')
            p=s.get(Product,pid)
            if not p or not p.active or p.supplier_id!=sup.id: raise HTTPException(400,'ケア・ドゥの商品ではありません')
            if q<=0: raise HTTPException(400,'発注数量は1以上にしてください')
            if q<p.min_order_qty: raise HTTPException(400,f'{p.name}は最低{p.min_order_qty}{p.unit}から発注できます')
            validated.append((p,q))
        no='PO-'+datetime.now(JST).strftime('%Y%m%d%H%M%S%f')+'-'+secrets.token_hex(2).upper(); o=PurchaseOrder(order_no=no,facility_id=u.facility_id,supplier_id=sup.id,status='submitted',approved_by=u.id,ordered_at=now_dt()); s.add(o); s.flush(); total=0
        for p,q in validated:
            fp=s.query(FacilityPrice).filter_by(facility_id=u.facility_id,product_id=p.id).first(); price=fp.purchase_price if fp else p.standard_price; total+=q*price; s.add(PurchaseOrderItem(order_id=o.id,product_id=p.id,qty=q,unit_price=price))
        o.total_amount=total
        emit_notification_event(s,'procurement.order_submitted',users_by_roles(s,['caredo']),'施設から発注が届きました',f'注文番号 {o.order_no}。ITTOKAIで受注内容をご確認ください。','procurement',o.id,u.facility_id,u.id,f'procurement.order_submitted:{o.id}')
        audit(s,request,u,'procurement','order_submit','purchase_order',o.id,{'total':total},u.facility_id); s.commit(); return {'id':o.id,'order_no':o.order_no,'total_amount':total}

class OrderStatusIn(BaseModel): status:str
@app.post('/api/orders/{oid}/status')
def order_status(oid:int,inp:OrderStatusIn,request:Request):
    with SessionLocal() as s:
        u=current_user(request,s); require_roles(u,'caredo','admin'); o=s.get(PurchaseOrder,oid)
        if not o: raise HTTPException(404,'注文がありません')
        if u.role=='admin': facility_scope(u,o.facility_id)
        allowed={'submitted':{'confirmed','cancelled'},'confirmed':{'shipped','cancelled'}} if u.role=='caredo' else {'submitted':{'cancelled'},'confirmed':{'cancelled'}}
        if inp.status not in allowed.get(o.status,set()): raise HTTPException(400,'現在の状態からその変更はできません')
        o.status=inp.status
        if inp.status=='confirmed': o.confirmed_at=now_dt()
        if inp.status=='shipped': o.shipped_at=now_dt()
        label={'confirmed':'ケア・ドゥが注文を確認しました','shipped':'ケア・ドゥが商品を出荷しました','cancelled':'発注がキャンセルされました'}.get(inp.status,'発注状況が更新されました')
        emit_notification_event(s,'procurement.order_status',users_by_roles(s,['admin','middle_manager'],o.facility_id),label,'ITTOKAIで発注状況をご確認ください。','procurement',o.id,o.facility_id,u.id,f'procurement.order_status:{o.id}:{inp.status}')
        audit(s,request,u,'procurement','order_status','purchase_order',o.id,{'status':inp.status},o.facility_id); s.commit(); return {'ok':True}

class InvoiceIn(BaseModel): order_id:int; invoice_no:str; invoice_date:str; items:list[dict]
@app.post('/api/invoices')
def create_invoice(inp:InvoiceIn,request:Request):
    with SessionLocal() as s:
        u=current_user(request,s); require_roles(u,'caredo'); o=s.get(PurchaseOrder,inp.order_id)
        if not o: raise HTTPException(404,'注文がありません')
        if o.status!='received': raise HTTPException(400,'検品完了後に請求書を作成してください')
        if s.query(Invoice).filter_by(order_id=o.id).first(): raise HTTPException(400,'この注文の請求書は作成済みです')
        if s.query(Invoice).filter_by(invoice_no=inp.invoice_no).first(): raise HTTPException(400,'請求書番号が重複しています')
        total=0; inv=Invoice(invoice_no=inp.invoice_no,supplier_id=o.supplier_id,facility_id=o.facility_id,order_id=o.id,invoice_date=inp.invoice_date,status='received'); s.add(inv); s.flush()
        for x in inp.items:
            pid=int(x['product_id']); qty=float(x['qty']); price=float(x['unit_price']); amount=qty*price; total+=amount; s.add(InvoiceItem(invoice_id=inv.id,product_id=pid,qty=qty,unit_price=price,amount=amount))
        inv.total_amount=total; perform_match(s,inv,o); s.flush(); inv.pdf_blob=make_invoice_pdf(s,inv)
        emit_notification_event(s,'procurement.invoice_created',users_by_roles(s,['admin','middle_manager'],o.facility_id),'請求書が作成されました','ITTOKAIで請求・3点照合結果をご確認ください。','procurement',inv.id,o.facility_id,u.id,f'procurement.invoice_created:{inv.id}')
        audit(s,request,u,'procurement','invoice_create','invoice',inv.id,{'match':inv.match_status},o.facility_id); s.commit(); return invoice_json(s,inv)

def perform_match(s,inv,o):
    diffs=[]; oi={x.product_id:x for x in s.query(PurchaseOrderItem).filter_by(order_id=o.id).all()}
    for ii in s.query(InvoiceItem).filter_by(invoice_id=inv.id).all():
        x=oi.get(ii.product_id)
        if not x: diffs.append({'product_id':ii.product_id,'type':'not_ordered'}); continue
        if abs(ii.qty-x.received_qty)>0.0001: diffs.append({'product_id':ii.product_id,'type':'qty','ordered':x.qty,'received':x.received_qty,'invoice':ii.qty})
        if abs(ii.unit_price-x.unit_price)>0.01: diffs.append({'product_id':ii.product_id,'type':'price','ordered_price':x.unit_price,'invoice_price':ii.unit_price})
    inv.difference_json=json.dumps({'items':diffs},ensure_ascii=False); inv.match_status='matched' if not diffs else 'difference'; o.status='invoiced' if not diffs else o.status



class ReceiveIn(BaseModel):
    items:list[dict]
@app.post('/api/orders/{oid}/receive')
def receive_order(oid:int,inp:ReceiveIn,request:Request):
    with SessionLocal() as s:
        u=current_user(request,s); require_roles(u,'admin','middle_manager'); o=s.get(PurchaseOrder,oid)
        if not o: raise HTTPException(404,'注文がありません')
        facility_scope(u,o.facility_id)
        if o.status not in ('shipped','partially_received'): raise HTTPException(400,'出荷済みの注文だけ検品できます')
        byid={x.id:x for x in s.query(PurchaseOrderItem).filter_by(order_id=o.id).all()}
        for x in inp.items:
            it=byid.get(int(x.get('item_id',0)))
            if not it: continue
            new=float(x.get('received_qty',0))
            if new<0 or new>it.qty: raise HTTPException(400,'受領数量が不正です')
            delta=new-it.received_qty; it.received_qty=new
            if delta:
                inv=s.query(Inventory).filter_by(facility_id=o.facility_id,product_id=it.product_id).first()
                if inv: inv.on_hand+=delta; inv.updated_at=now_dt()
                pol=s.query(FacilityStockPolicy).filter_by(facility_id=o.facility_id,product_id=it.product_id).first()
                if pol: pol.last_received_at=now_dt(); pol.updated_at=now_dt()
                s.add(InventoryTxn(facility_id=o.facility_id,product_id=it.product_id,txn_type='receipt',qty=delta,reference_type='purchase_order',reference_id=str(o.id),actor_id=u.id))
        all_items=list(byid.values()); complete=all(abs(x.received_qty-x.qty)<0.0001 for x in all_items)
        o.status='received' if complete else 'partially_received'; o.received_at=now_dt() if complete else None
        emit_notification_event(s,'procurement.order_status',users_by_roles(s,['caredo']),'施設で受領処理が行われました',('全数の受領が完了しました。' if complete else '一部受領として登録されました。'),'procurement',o.id,o.facility_id,u.id,f'procurement.received:{o.id}:{o.status}')
        audit(s,request,u,'procurement','receive_inspect','purchase_order',o.id,{'complete':complete},o.facility_id); s.commit(); return {'ok':True,'status':o.status}

class ProductIn(BaseModel):
    sku:str; name:str; specification:str=''; category:str=''; unit:str='個'; standard_price:float=0; purchase_price:float=0; sale_price:float=0; short_message:str=''; description:str=''; recommended:int=0; min_order_qty:int=1; lead_time_days:int=2; pack_size:int=1; inventory_managed:int=1; active:int=1; sale_start:str=''; sale_end:str=''

class ProductUpdate(BaseModel):
    name:Optional[str]=None; specification:Optional[str]=None; category:Optional[str]=None; unit:Optional[str]=None; standard_price:Optional[float]=None; purchase_price:Optional[float]=None; sale_price:Optional[float]=None; short_message:Optional[str]=None; description:Optional[str]=None; recommended:Optional[int]=None; min_order_qty:Optional[int]=None; lead_time_days:Optional[int]=None; pack_size:Optional[int]=None; inventory_managed:Optional[int]=None; active:Optional[int]=None; sale_start:Optional[str]=None; sale_end:Optional[str]=None
@app.get('/api/products/{product_id}/image')
def get_product_image(product_id:int):
    with SessionLocal() as s:
        img=s.query(ProductImage).filter_by(product_id=product_id,is_primary=1).order_by(ProductImage.id.desc()).first()
        if not img: raise HTTPException(404,'商品画像がありません')
        return Response(content=bytes(img.image_blob),media_type=img.mime_type,headers={'Cache-Control':'public, max-age=3600'})

@app.post('/api/products/{product_id}/image')
async def upload_product_image(product_id:int,request:Request,file:UploadFile=File(...)):
    raw=await file.read()
    if len(raw)>5*1024*1024: raise HTTPException(400,'商品画像は5MB以下にしてください')
    if file.content_type not in ('image/jpeg','image/png','image/webp'): raise HTTPException(400,'JPEG/PNG/WebPを選択してください')
    try:
        im=PILImage.open(io.BytesIO(raw)); im.verify()
    except Exception: raise HTTPException(400,'画像ファイルを確認できません')
    with SessionLocal() as s:
        u=current_user(request,s); require_roles(u,'caredo','hq'); p=s.get(Product,product_id)
        if not p: raise HTTPException(404,'商品がありません')
        for old in s.query(ProductImage).filter_by(product_id=product_id,is_primary=1).all(): old.is_primary=0
        img=ProductImage(product_id=product_id,mime_type=file.content_type,image_blob=raw,image_hash=sha256_bytes(raw),is_primary=1); s.add(img); audit(s,request,u,'procurement','product_image_upload','product',product_id,{'hash':img.image_hash}); s.commit(); return {'ok':True,'image_url':f'/api/products/{product_id}/image'}

@app.post('/api/products')
def create_product(inp:ProductIn,request:Request):
    with SessionLocal() as s:
        u=current_user(request,s); require_roles(u,'caredo','hq'); sup=s.query(Supplier).filter_by(code='CAREDO').first()
        if s.query(Product).filter_by(sku=inp.sku).first(): raise HTTPException(400,'SKUが重複しています')
        d=inp.model_dump(); meta_keys={'purchase_price','sale_price','short_message','description','recommended','sale_start','sale_end'}; pd={k:v for k,v in d.items() if k not in meta_keys}; p=Product(supplier_id=sup.id,**pd); s.add(p); s.flush(); m=ProductMeta(product_id=p.id,purchase_price=inp.purchase_price or inp.standard_price,sale_price=inp.sale_price or inp.standard_price,short_message=inp.short_message,description=inp.description,recommended=inp.recommended,sale_start=inp.sale_start,sale_end=inp.sale_end); s.add(m); audit(s,request,u,'procurement','product_create','product',p.id,{'sku':p.sku}); s.commit(); return {'id':p.id,'ok':True}

@app.patch('/api/products/{product_id}')
def update_product(product_id:int,inp:ProductUpdate,request:Request):
    with SessionLocal() as s:
        u=current_user(request,s); require_roles(u,'caredo','hq'); p=s.get(Product,product_id)
        if not p: raise HTTPException(404,'商品がありません')
        d=inp.model_dump(exclude_none=True); meta_keys={'purchase_price','sale_price','short_message','description','recommended','sale_start','sale_end'}
        for k,v in d.items():
            if k not in meta_keys and hasattr(p,k): setattr(p,k,v)
        m=product_meta_for(s,p.id) or ProductMeta(product_id=p.id); s.add(m)
        for k in meta_keys:
            if k in d: setattr(m,k,d[k])
        m.updated_at=now_dt(); audit(s,request,u,'procurement','product_update','product',p.id,d); s.commit(); return {'ok':True,'product':caredo_product_json(s,p)}

class ProductBatchIn(BaseModel):
    product_ids:list[int]; changes:dict
@app.post('/api/products/batch')
def batch_products(inp:ProductBatchIn,request:Request):
    with SessionLocal() as s:
        u=current_user(request,s); require_roles(u,'caredo','hq'); n=0
        for pid in inp.product_ids:
            p=s.get(Product,pid)
            if not p: continue
            m=product_meta_for(s,p.id) or ProductMeta(product_id=p.id); s.add(m)
            for k,v in inp.changes.items():
                if hasattr(p,k): setattr(p,k,v)
                elif hasattr(m,k): setattr(m,k,v)
            m.updated_at=now_dt(); n+=1
        audit(s,request,u,'procurement','product_batch','product','',{'count':n,'changes':inp.changes}); s.commit(); return {'ok':True,'count':n}

class StockPolicyIn(BaseModel):
    facility_id:int; product_id:int; reorder_point:float; target_stock:float=0; reorder_qty:float=0; max_stock:float=0; auto_order_mode:str='approval'
@app.post('/api/inventory/policy')
def update_stock_policy(inp:StockPolicyIn,request:Request):
    with SessionLocal() as s:
        u=current_user(request,s); require_roles(u,'admin','hq','caredo')
        if u.role=='admin' and inp.facility_id!=u.facility_id: raise HTTPException(403,'他施設は変更できません')
        if inp.auto_order_mode not in ('off','approval','auto'): raise HTTPException(400,'自動発注モードが不正です')
        inv=s.query(Inventory).filter_by(facility_id=inp.facility_id,product_id=inp.product_id).first()
        if not inv: inv=Inventory(facility_id=inp.facility_id,product_id=inp.product_id,location='主倉庫',on_hand=0); s.add(inv)
        inv.reorder_point=inp.reorder_point; inv.target_stock=inp.target_stock; inv.updated_at=now_dt()
        pol=s.query(FacilityStockPolicy).filter_by(facility_id=inp.facility_id,product_id=inp.product_id).first() or FacilityStockPolicy(facility_id=inp.facility_id,product_id=inp.product_id); pol.reorder_qty=inp.reorder_qty; pol.max_stock=inp.max_stock; pol.auto_order_mode=inp.auto_order_mode; pol.updated_at=now_dt(); s.add(pol); s.flush()
        created=evaluate_auto_orders(s,inp.facility_id,inp.product_id,u.id); audit(s,request,u,'procurement','stock_policy_update','product',inp.product_id,inp.model_dump(),inp.facility_id); s.commit(); return {'ok':True,'auto_order_events':created}

@app.post('/api/inventory/evaluate-auto-orders')
def run_auto_orders(request:Request):
    with SessionLocal() as s:
        u=current_user(request,s); require_roles(u,'admin','hq','caredo'); fid=u.facility_id if u.role=='admin' else None; created=evaluate_auto_orders(s,fid,None,u.id); audit(s,request,u,'procurement','auto_order_evaluate','facility',fid or 'all',{'created':created},fid); s.commit(); return {'ok':True,'created':created}

@app.post('/api/auto-orders/{event_id}/approve')
def approve_auto_order(event_id:int,request:Request):
    with SessionLocal() as s:
        u=current_user(request,s); require_roles(u,'admin','caredo','hq'); ev=s.get(AutoOrderEvent,event_id)
        if not ev or ev.status!='pending_approval': raise HTTPException(404,'承認待ちの自動発注候補がありません')
        if u.role=='admin' and ev.facility_id!=u.facility_id: raise HTTPException(403,'他施設は承認できません')
        p=s.get(Product,ev.product_id); sup=s.query(Supplier).filter_by(code='CAREDO').first()
        if open_incoming_qty(s,ev.facility_id,ev.product_id)>0: ev.status='skipped_duplicate'; ev.resolved_at=now_dt(); s.commit(); return {'ok':True,'status':ev.status}
        o=_create_po(s,ev.facility_id,sup,[(p,ev.suggested_qty)],approved_by=u.id,status='submitted'); ev.order_id=o.id; ev.status='ordered'; ev.resolved_at=now_dt(); pol=s.query(FacilityStockPolicy).filter_by(facility_id=ev.facility_id,product_id=ev.product_id).first();
        if pol: pol.last_ordered_at=now_dt()
        rec=users_by_roles(s,['caredo'])+users_by_roles(s,['admin'],ev.facility_id)
        rec=[x for x in rec if x.id!=u.id]
        emit_notification_event(s,'inventory.auto_order_created',rec,'自動発注が承認・作成されました',f'{p.name} の発注が作成されました。','procurement',o.id,ev.facility_id,u.id,f'inventory.auto_order_approved:{ev.id}')
        audit(s,request,u,'procurement','auto_order_approve','purchase_order',o.id,{'event_id':ev.id},ev.facility_id); s.commit(); return {'ok':True,'order_id':o.id,'order_no':o.order_no}

# ---- Excel master import/export ----
def _xlsx_response(wb,filename):
    bio=io.BytesIO(); wb.save(bio); bio.seek(0); return StreamingResponse(bio,media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',headers={'Content-Disposition':f'attachment; filename="{filename}"'})

def _xlsx_rows(raw):
    wb=load_workbook(io.BytesIO(raw),data_only=True); ws=wb.active; headers=[str(x.value or '').strip() for x in ws[1]]; out=[]
    for row in ws.iter_rows(min_row=2,values_only=True):
        if not any(v not in (None,'') for v in row): continue
        out.append({headers[i]:row[i] for i in range(min(len(headers),len(row)))})
    return out

@app.get('/api/caredo/products.xlsx')
def export_caredo_products_xlsx(request:Request):
    with SessionLocal() as s:
        u=current_user(request,s); require_roles(u,'caredo','hq'); sup=s.query(Supplier).filter_by(code='CAREDO').first(); wb=Workbook(); ws=wb.active; ws.title='商品マスタ'; headers=['sku','name','short_message','description','specification','category','unit','pack_size','purchase_price','sale_price','gross_margin_rate','min_order_qty','lead_time_days','inventory_managed','recommended','active','sale_start','sale_end']; ws.append(headers)
        for p in s.query(Product).filter_by(supplier_id=sup.id).order_by(Product.sku).all():
            x=caredo_product_json(s,p); ws.append([x.get(h,'') for h in headers])
        return _xlsx_response(wb,'ITTOKAI_CareDo_ProductMaster_v20.7.xlsx')

@app.post('/api/caredo/products/import-xlsx')
async def import_caredo_products_xlsx(request:Request,file:UploadFile=File(...),commit:bool=False):
    raw=await file.read(); rows=_xlsx_rows(raw); preview={'new':0,'update':0,'unchanged':0,'errors':[]}; parsed=[]
    with SessionLocal() as s:
        u=current_user(request,s); require_roles(u,'caredo','hq'); sup=s.query(Supplier).filter_by(code='CAREDO').first()
        for idx,r in enumerate(rows,start=2):
            sku=str(r.get('sku') or '').strip(); name=str(r.get('name') or '').strip()
            if not sku or not name: preview['errors'].append({'row':idx,'error':'sku/nameは必須'}); continue
            p=s.query(Product).filter_by(sku=sku).first(); preview['update' if p else 'new']+=1; parsed.append((r,p))
        if commit and not preview['errors']:
            for r,p in parsed:
                if not p: p=Product(supplier_id=sup.id,sku=str(r.get('sku')).strip(),name=str(r.get('name')).strip()); s.add(p); s.flush()
                p.name=str(r.get('name') or p.name); p.specification=str(r.get('specification') or ''); p.category=str(r.get('category') or ''); p.unit=str(r.get('unit') or '個'); p.pack_size=int(float(r.get('pack_size') or 1)); p.standard_price=float(r.get('sale_price') or 0); p.min_order_qty=int(float(r.get('min_order_qty') or 1)); p.lead_time_days=int(float(r.get('lead_time_days') or 2)); p.inventory_managed=int(float(r.get('inventory_managed') if r.get('inventory_managed') is not None else 1)); p.active=int(float(r.get('active') if r.get('active') is not None else 1))
                m=product_meta_for(s,p.id) or ProductMeta(product_id=p.id); s.add(m); m.short_message=str(r.get('short_message') or ''); m.description=str(r.get('description') or ''); m.purchase_price=float(r.get('purchase_price') or 0); m.sale_price=float(r.get('sale_price') or 0); m.recommended=int(float(r.get('recommended') or 0)); m.sale_start=str(r.get('sale_start') or ''); m.sale_end=str(r.get('sale_end') or ''); m.updated_at=now_dt()
            audit(s,request,u,'procurement','product_xlsx_import','product','',preview); s.commit()
        return {'ok':True,'commit':commit,**preview}

@app.get('/api/caredo/facility-inventory.xlsx')
def export_facility_inventory_xlsx(request:Request):
    with SessionLocal() as s:
        u=current_user(request,s); require_roles(u,'caredo','hq','admin'); wb=Workbook(); ws=wb.active; ws.title='施設別在庫'; headers=['facility_code','facility_name','sku','name','on_hand','safety_stock','reorder_point','target_stock','reorder_qty','max_stock','auto_order_mode','status','incoming_qty','last_ordered_at','last_received_at']; ws.append(headers)
        fs=[s.get(Facility,u.facility_id)] if u.role=='admin' else s.query(Facility).filter_by(active=1).all()
        for f in fs:
            for x in procurement_snapshot(s,f.id)['inventory']: ws.append([f.code,f.name,x['sku'],x['name'],x['on_hand'],x['safety_stock'],x['reorder_point'],x['target_stock'],x['reorder_qty'],x['max_stock'],x['auto_order_mode'],x['status'],x['incoming_qty'],x['last_ordered_at'],x['last_received_at']])
        return _xlsx_response(wb,'ITTOKAI_FacilityInventory_v20.7.xlsx')

@app.post('/api/caredo/facility-inventory/import-xlsx')
async def import_facility_inventory_xlsx(request:Request,file:UploadFile=File(...),commit:bool=False):
    raw=await file.read(); rows=_xlsx_rows(raw); preview={'new':0,'update':0,'errors':[]}; parsed=[]
    with SessionLocal() as s:
        u=current_user(request,s); require_roles(u,'caredo','hq','admin')
        for idx,r in enumerate(rows,start=2):
            f=s.query(Facility).filter_by(code=str(r.get('facility_code') or '').strip()).first(); p=s.query(Product).filter_by(sku=str(r.get('sku') or '').strip()).first()
            if not f or not p: preview['errors'].append({'row':idx,'error':'施設コードまたはSKU不明'}); continue
            if u.role=='admin' and f.id!=u.facility_id: preview['errors'].append({'row':idx,'error':'他施設は更新不可'}); continue
            inv=s.query(Inventory).filter_by(facility_id=f.id,product_id=p.id,location='主倉庫').first(); preview['update' if inv else 'new']+=1; parsed.append((r,f,p,inv))
        if commit and not preview['errors']:
            for r,f,p,inv in parsed:
                if not inv: inv=Inventory(facility_id=f.id,product_id=p.id,location='主倉庫'); s.add(inv)
                for k in ('on_hand','safety_stock','reorder_point','target_stock'): setattr(inv,k,float(r.get(k) or 0))
                inv.updated_at=now_dt(); pol=s.query(FacilityStockPolicy).filter_by(facility_id=f.id,product_id=p.id).first() or FacilityStockPolicy(facility_id=f.id,product_id=p.id); pol.reorder_qty=float(r.get('reorder_qty') or 0); pol.max_stock=float(r.get('max_stock') or 0); pol.auto_order_mode=str(r.get('auto_order_mode') or 'approval'); pol.updated_at=now_dt(); s.add(pol); s.flush(); evaluate_auto_orders(s,f.id,p.id,u.id)
            audit(s,request,u,'procurement','facility_inventory_xlsx_import','inventory','',preview); s.commit()
        return {'ok':True,'commit':commit,**preview}

class PartnerStoreIn(BaseModel):
    company_name:str=''; store_name:str; category:str='その他'; image_url:str=''; address:str=''; phone:str=''; hours:str=''; closed_days:str=''; delivery_area:str=''; min_order_amount:float=0; order_cutoff:str=''; delivery_methods:str='施設配送'; message:str=''; active:int=1
class PartnerProductIn(BaseModel):
    store_id:int; sku:str; name:str; short_message:str=''; description:str=''; category:str='ギフト'; specification:str=''; unit:str='個'; price:float=0; image_url:str=''; recommended:int=0; active:int=1; sale_start:str=''; sale_end:str=''

@app.post('/api/sponsor/stores')
def create_partner_store(inp:PartnerStoreIn,request:Request):
    with SessionLocal() as s:
        u=current_user(request,s); require_roles(u,'sponsor'); st=PartnerStore(owner_user_id=u.id,**inp.model_dump()); s.add(st); s.flush(); audit(s,request,u,'procurement','partner_store_create','partner_store',st.id,{'store_name':st.store_name}); s.commit(); return {'ok':True,'id':st.id}
@app.patch('/api/sponsor/stores/{sid}')
def update_partner_store(sid:int,inp:PartnerStoreIn,request:Request):
    with SessionLocal() as s:
        u=current_user(request,s); require_roles(u,'sponsor'); st=s.get(PartnerStore,sid)
        if not st or st.owner_user_id!=u.id: raise HTTPException(404,'店舗がありません')
        for k,v in inp.model_dump().items(): setattr(st,k,v)
        st.updated_at=now_dt(); audit(s,request,u,'procurement','partner_store_update','partner_store',st.id,{}); s.commit(); return {'ok':True}
@app.post('/api/sponsor/products')
def create_partner_product(inp:PartnerProductIn,request:Request):
    with SessionLocal() as s:
        u=current_user(request,s); require_roles(u,'sponsor'); st=s.get(PartnerStore,inp.store_id)
        if not st or st.owner_user_id!=u.id: raise HTTPException(403,'店舗権限がありません')
        if s.query(PartnerProduct).filter_by(store_id=st.id,sku=inp.sku).first(): raise HTTPException(400,'SKUが重複しています')
        p=PartnerProduct(**inp.model_dump()); s.add(p); s.flush(); audit(s,request,u,'procurement','partner_product_create','partner_product',p.id,{}); s.commit(); return {'ok':True,'id':p.id}
@app.patch('/api/sponsor/products/{pid}')
def update_partner_product(pid:int,inp:PartnerProductIn,request:Request):
    with SessionLocal() as s:
        u=current_user(request,s); require_roles(u,'sponsor'); p=s.get(PartnerProduct,pid); st=s.get(PartnerStore,p.store_id) if p else None
        if not p or not st or st.owner_user_id!=u.id: raise HTTPException(404,'商品がありません')
        d=inp.model_dump()
        if d['store_id']!=p.store_id: newst=s.get(PartnerStore,d['store_id']);
        if d['store_id']!=p.store_id and (not newst or newst.owner_user_id!=u.id): raise HTTPException(403,'店舗権限がありません')
        for k,v in d.items(): setattr(p,k,v)
        p.updated_at=now_dt(); audit(s,request,u,'procurement','partner_product_update','partner_product',p.id,{}); s.commit(); return {'ok':True}

@app.get('/api/sponsor/master.xlsx')
def export_partner_master_xlsx(request:Request):
    with SessionLocal() as s:
        u=current_user(request,s); require_roles(u,'sponsor'); wb=Workbook(); ws=wb.active; ws.title='店舗マスタ'; ws.append(['store_id','company_name','store_name','category','address','phone','hours','closed_days','delivery_area','min_order_amount','order_cutoff','delivery_methods','message','active'])
        stores=s.query(PartnerStore).filter_by(owner_user_id=u.id).order_by(PartnerStore.id).all()
        for st in stores: ws.append([st.id,st.company_name,st.store_name,st.category,st.address,st.phone,st.hours,st.closed_days,st.delivery_area,st.min_order_amount,st.order_cutoff,st.delivery_methods,st.message,st.active])
        ws2=wb.create_sheet('商品マスタ'); ws2.append(['store_id','sku','name','short_message','description','category','specification','unit','price','image_url','recommended','active','sale_start','sale_end'])
        ids=[x.id for x in stores]
        for p in s.query(PartnerProduct).filter(PartnerProduct.store_id.in_(ids or [-1])).order_by(PartnerProduct.id).all(): ws2.append([p.store_id,p.sku,p.name,p.short_message,p.description,p.category,p.specification,p.unit,p.price,p.image_url,p.recommended,p.active,p.sale_start,p.sale_end])
        return _xlsx_response(wb,'ITTOKAI_Sponsor_Master_v20.7.xlsx')

@app.post('/api/sponsor/master/import-xlsx')
async def import_partner_master_xlsx(request:Request,file:UploadFile=File(...),commit:bool=False):
    raw=await file.read(); wb=load_workbook(io.BytesIO(raw),data_only=True); preview={'stores_new':0,'stores_update':0,'products_new':0,'products_update':0,'errors':[]}
    with SessionLocal() as s:
        u=current_user(request,s); require_roles(u,'sponsor'); store_map={}
        if '店舗マスタ' in wb.sheetnames:
            ws=wb['店舗マスタ']; hs=[str(c.value or '').strip() for c in ws[1]]
            for idx,row in enumerate(ws.iter_rows(min_row=2,values_only=True),start=2):
                if not any(v not in (None,'') for v in row): continue
                r={hs[i]:row[i] for i in range(min(len(hs),len(row)))}; sid=int(r.get('store_id') or 0); name=str(r.get('store_name') or '').strip()
                if not name: preview['errors'].append({'sheet':'店舗マスタ','row':idx,'error':'store_nameは必須'}); continue
                st=s.get(PartnerStore,sid) if sid else None
                if st and st.owner_user_id!=u.id: preview['errors'].append({'sheet':'店舗マスタ','row':idx,'error':'他社店舗は更新不可'}); continue
                preview['stores_update' if st else 'stores_new']+=1
                if commit and not preview['errors']:
                    if not st: st=PartnerStore(owner_user_id=u.id,store_name=name); s.add(st); s.flush()
                    for k in ['company_name','store_name','category','address','phone','hours','closed_days','delivery_area','order_cutoff','delivery_methods','message']:
                        if k in r: setattr(st,k,str(r.get(k) or ''))
                    st.min_order_amount=float(r.get('min_order_amount') or 0); st.active=int(float(r.get('active') if r.get('active') is not None else 1)); st.updated_at=now_dt(); store_map[sid or st.id]=st.id
                elif st: store_map[sid]=st.id
        if '商品マスタ' in wb.sheetnames:
            ws=wb['商品マスタ']; hs=[str(c.value or '').strip() for c in ws[1]]
            for idx,row in enumerate(ws.iter_rows(min_row=2,values_only=True),start=2):
                if not any(v not in (None,'') for v in row): continue
                r={hs[i]:row[i] for i in range(min(len(hs),len(row)))}; sid=int(r.get('store_id') or 0); sid=store_map.get(sid,sid); st=s.get(PartnerStore,sid) if sid else None; sku=str(r.get('sku') or '').strip(); name=str(r.get('name') or '').strip()
                if not st or st.owner_user_id!=u.id or not sku or not name: preview['errors'].append({'sheet':'商品マスタ','row':idx,'error':'store_id/sku/nameを確認してください'}); continue
                pp=s.query(PartnerProduct).filter_by(store_id=st.id,sku=sku).first(); preview['products_update' if pp else 'products_new']+=1
                if commit and not preview['errors']:
                    if not pp: pp=PartnerProduct(store_id=st.id,sku=sku,name=name); s.add(pp)
                    pp.name=name; pp.short_message=str(r.get('short_message') or ''); pp.description=str(r.get('description') or ''); pp.category=str(r.get('category') or 'ギフト'); pp.specification=str(r.get('specification') or ''); pp.unit=str(r.get('unit') or '個'); pp.price=float(r.get('price') or 0); pp.image_url=str(r.get('image_url') or ''); pp.recommended=int(float(r.get('recommended') or 0)); pp.active=int(float(r.get('active') if r.get('active') is not None else 1)); pp.sale_start=str(r.get('sale_start') or ''); pp.sale_end=str(r.get('sale_end') or ''); pp.updated_at=now_dt()
        if commit and not preview['errors']:
            audit(s,request,u,'procurement','partner_xlsx_import','partner_master','',preview); s.commit()
        return {'ok':True,'commit':commit,**preview}

def make_invoice_pdf(s,inv:Invoice):
    buf=io.BytesIO(); st=pdf_styles(); fac=s.get(Facility,inv.facility_id); sup=s.get(Supplier,inv.supplier_id); story=[Paragraph('請求書',st['title']),Paragraph(f'請求書番号: {inv.invoice_no}',st['body']),Paragraph(f'請求日: {inv.invoice_date}',st['body']),Spacer(1,4*mm),Paragraph(f'請求先: {fac.name if fac else ""}',st['body']),Paragraph(f'発行元: {sup.name if sup else ""}',st['body']),Spacer(1,6*mm)]
    data=[['SKU','商品','数量','単価','金額']]
    for it in s.query(InvoiceItem).filter_by(invoice_id=inv.id).all():
        p=s.get(Product,it.product_id); data.append([p.sku if p else '',p.name if p else '',str(it.qty),f'{it.unit_price:,.0f}',f'{it.amount:,.0f}'])
    data.append(['','','','','合計 ¥'+f'{inv.total_amount:,.0f}']); t=Table(data,colWidths=[28*mm,65*mm,18*mm,28*mm,35*mm]); t.setStyle(TableStyle([('FONTNAME',(0,0),(-1,-1),'HeiseiKakuGo-W5'),('FONTSIZE',(0,0),(-1,-1),8),('GRID',(0,0),(-1,-1),.3,colors.HexColor('#dddddd')),('BACKGROUND',(0,0),(-1,0),colors.HexColor('#fff0df')),('ALIGN',(2,1),(-1,-1),'RIGHT')])); story.append(t); story.append(Spacer(1,5*mm)); story.append(Paragraph(f'3点照合: {inv.match_status}',st['small'])); SimpleDocTemplate(buf,pagesize=A4,rightMargin=15*mm,leftMargin=15*mm,topMargin=15*mm,bottomMargin=15*mm).build(story); return buf.getvalue()

@app.get('/api/invoices/{iid}/pdf')
def invoice_pdf(iid:int,request:Request):
    with SessionLocal() as s:
        u=current_user(request,s); require_roles(u,'admin','hq','caredo'); inv=s.get(Invoice,iid)
        if not inv: raise HTTPException(404,'請求書がありません')
        if u.role=='admin': facility_scope(u,inv.facility_id)
        pdf=inv.pdf_blob or make_invoice_pdf(s,inv); return Response(content=bytes(pdf),media_type='application/pdf',headers={'Content-Disposition':f'inline; filename="invoice_{iid}.pdf"'})

# -------- CSV bulk import --------
def csv_rows(raw:bytes):
    text=raw.decode('utf-8-sig'); return list(csv.DictReader(io.StringIO(text)))

@app.post('/api/import/{kind}')
async def bulk_import(kind:str,request:Request,file:UploadFile=File(...)):
    raw=await file.read()
    rows=csv_rows(raw); count=0; warnings=[]
    with SessionLocal() as s:
        u=current_user(request,s)
        if kind=='facilities':
            require_roles(u,'hq')
            for r in rows:
                code=(r.get('facility_code') or '').strip(); name=(r.get('facility_name') or '').strip()
                if not code or not name: warnings.append('施設コード/施設名不足'); continue
                f=s.query(Facility).filter_by(code=code).first() or Facility(code=code,name=name); f.name=name; f.corporation=r.get('corporation') or '社会福祉法人 一燈会'; f.address=r.get('address') or ''; f.phone=r.get('phone') or ''; s.add(f); count+=1
        elif kind=='staff':
            require_roles(u,'hq','admin'); default_pw=os.environ.get('IMPORT_DEFAULT_PASSWORD','ChangeMe1234!')
            for r in rows:
                email=(r.get('email') or '').strip().lower(); code=(r.get('employee_code') or '').strip(); name=(r.get('name') or '').strip(); fcode=(r.get('facility_code') or '').strip(); dcode=(r.get('department_code') or '').strip(); role=(r.get('role') or 'staff').strip()
                f=s.query(Facility).filter_by(code=fcode).first();
                if not email or not f or role not in ('staff','middle_manager','admin'): warnings.append(f'{email or name}: データ不足'); continue
                if u.role=='admin' and f.id!=u.facility_id: warnings.append(f'{email}: 他施設'); continue
                d=s.query(Department).filter_by(facility_id=f.id,code=dcode).first() if dcode else None
                usr=s.query(User).filter(func.lower(User.email)==email).first()
                if not usr: usr=User(email=email,password_hash=hash_password(default_pw),created_at=now_dt())
                usr.employee_code=code; usr.name=name; usr.role=role; usr.facility_id=f.id; usr.department_id=d.id if d else None; usr.active=1; s.add(usr); count+=1
        elif kind=='residents_family':
            require_roles(u,'hq','admin'); default_pw=os.environ.get('IMPORT_DEFAULT_PASSWORD','ChangeMe1234!')
            for r in rows:
                f=s.query(Facility).filter_by(code=(r.get('facility_code') or '').strip()).first(); rcode=(r.get('resident_code') or '').strip(); rname=(r.get('resident_name') or '').strip(); email=(r.get('family_email') or '').strip().lower()
                if not f or not rcode or not rname or not email: warnings.append(f'{rcode or rname}: データ不足'); continue
                if u.role=='admin' and f.id!=u.facility_id: continue
                res=s.query(Resident).filter_by(resident_code=rcode).first() or Resident(resident_code=rcode,facility_id=f.id,name=rname); res.name=rname; res.facility_id=f.id; res.service_type=r.get('service_type') or '入所'; s.add(res); s.flush()
                fam=s.query(User).filter(func.lower(User.email)==email).first()
                if not fam: fam=User(name=r.get('family_name') or '家族',email=email,password_hash=hash_password(default_pw),role='family',facility_id=f.id,created_at=now_dt()); s.add(fam); s.flush()
                if not s.query(FamilyLink).filter_by(user_id=fam.id,resident_id=res.id).first(): s.add(FamilyLink(user_id=fam.id,resident_id=res.id,relation=r.get('relation') or '家族',is_primary=1))
                count+=1
        elif kind=='caredo_products':
            require_roles(u,'hq','caredo'); sup=s.query(Supplier).filter_by(code='CAREDO').first()
            for r in rows:
                sku=(r.get('sku') or '').strip(); name=(r.get('name') or '').strip()
                if not sku or not name: continue
                p=s.query(Product).filter_by(sku=sku).first() or Product(supplier_id=sup.id,sku=sku,name=name); p.name=name; p.specification=r.get('specification') or ''; p.category=r.get('category') or ''; p.unit=r.get('unit') or '個'; p.pack_size=int(float(r.get('pack_size') or 1)); p.standard_price=float(r.get('standard_price') or 0); p.min_order_qty=int(float(r.get('min_order_qty') or 1)); p.lead_time_days=int(float(r.get('lead_time_days') or 2)); p.substitute_sku=r.get('substitute_sku') or ''; p.inventory_managed=int(float(r.get('inventory_managed') or 1)); p.active=1; s.add(p); count+=1
        elif kind=='inventory':
            require_roles(u,'hq','admin')
            for r in rows:
                f=s.query(Facility).filter_by(code=(r.get('facility_code') or '').strip()).first(); p=s.query(Product).filter_by(sku=(r.get('sku') or '').strip()).first()
                if not f or not p: warnings.append('施設またはSKU不明'); continue
                if u.role=='admin' and f.id!=u.facility_id: continue
                loc=r.get('location') or '主倉庫'; inv=s.query(Inventory).filter_by(facility_id=f.id,product_id=p.id,location=loc).first() or Inventory(facility_id=f.id,product_id=p.id,location=loc)
                for field in ['on_hand','safety_stock','reorder_point','target_stock','avg_daily_usage']: setattr(inv,field,float(r.get(field) or 0))
                inv.updated_at=now_dt(); s.add(inv)
                fp=s.query(FacilityPrice).filter_by(facility_id=f.id,product_id=p.id).first() or FacilityPrice(facility_id=f.id,product_id=p.id)
                fp.purchase_price=float(r.get('purchase_price') or 0); fp.sale_price=float(r.get('sale_price') or 0); s.add(fp); count+=1
        elif kind=='leave_grants':
            require_roles(u,'hq','admin')
            for r in rows:
                usr=s.query(User).filter_by(employee_code=(r.get('employee_code') or '').strip()).first(); gd=safe_date(r.get('grant_date'))
                if not usr or not gd: warnings.append(f"{r.get('employee_code','')}: 職員または付与日不明"); continue
                if u.role=='admin' and usr.facility_id!=u.facility_id: continue
                exp=safe_date(r.get('expires_on')) or (add_years(gd,2)-timedelta(days=1)); granted=float(r.get('granted_days') or 0); remain=float(r.get('remaining_days') or granted); taken=float(r.get('taken_days') or max(0,granted-remain))
                g=s.query(LeaveGrant).filter_by(user_id=usr.id,grant_date=gd.isoformat()).first() or LeaveGrant(user_id=usr.id,grant_date=gd.isoformat())
                g.granted_days=granted; g.remaining_days=remain; g.taken_days=taken; g.expires_on=exp.isoformat(); g.source='chronos_import'; g.note=r.get('note') or ''; s.add(g); count+=1
        elif kind=='chronos_actual':
            require_roles(u,'hq','admin')
            for r in rows:
                usr=s.query(User).filter_by(employee_code=(r.get('employee_code') or '').strip()).first(); wd=safe_date(r.get('work_date'))
                if not usr or not wd: warnings.append(f"{r.get('employee_code','')}: 職員または勤務日不明"); continue
                if u.role=='admin' and usr.facility_id!=u.facility_id: continue
                rec=s.query(WorkTimeRecord).filter_by(user_id=usr.id,work_date=wd.isoformat(),source='chronos').first() or WorkTimeRecord(user_id=usr.id,work_date=wd.isoformat(),source='chronos')
                rec.overtime_minutes=int(float(r.get('overtime_minutes') or 0)); rec.holiday_minutes=int(float(r.get('holiday_minutes') or 0)); rec.imported_at=now_dt(); s.add(rec); count+=1
        else: raise HTTPException(404,'未対応の取込種別です')
        audit(s,request,u,'import','csv_import',kind,'',{'count':count,'warnings':warnings[:20]}); s.commit(); return {'ok':True,'count':count,'warnings':warnings[:50]}

# -------- integrations / dashboard --------
def hq_dashboard(s):
    return {'facilities':s.query(Facility).filter_by(active=1).count(),'users':s.query(User).filter_by(active=1).count(),'families':s.query(User).filter_by(role='family',active=1).count(),'staff':s.query(User).filter(User.role.in_(['staff','middle_manager','admin']),User.active==1).count(),'middle_managers':s.query(User).filter_by(role='middle_manager',active=1).count(),'admins':s.query(User).filter_by(role='admin',active=1).count(),'documents_pending':s.query(DocumentRecipient).filter(DocumentRecipient.status!='signed').count(),'workflow_pending':s.query(WorkflowRequest).filter(WorkflowRequest.status.in_(['pending_manager','pending_admin'])).count(),'visits_upcoming':s.query(VisitReservation).filter(VisitReservation.status.in_(['confirmed','pending'])).count(),'orders_open':s.query(PurchaseOrder).filter(PurchaseOrder.status.notin_(['closed','cancelled'])).count()}
def integration_list(s): return [{'system_name':x.system_name,'display_name':x.display_name,'mode':x.mode,'direction':x.direction,'status':x.status,'notes':x.notes} for x in s.query(IntegrationProfile).all()]

@app.get('/api/dashboard')
def dashboard(request:Request):
    with SessionLocal() as s:
        u=current_user(request,s)
        if u.role=='hq': return hq_dashboard(s)
        if u.role in ('admin','middle_manager'): return {'workflow_pending':len(workflow_list(s,u,approval=True)),'visits':len(visit_list_for_facility(s,u.facility_id)),'procurement':procurement_snapshot(s,u.facility_id) if u.role=='admin' else {}}
        return {}

@app.get('/api/integrations')
def integrations(request:Request):
    with SessionLocal() as s:
        u=current_user(request,s); require_roles(u,'hq','admin'); return {'items':integration_list(s)}

@app.get('/api/integrations/export/{system_name}')
def integration_export(system_name:str,request:Request):
    with SessionLocal() as s:
        u=current_user(request,s); require_roles(u,'hq','admin')
        out=io.StringIO(); w=csv.writer(out)
        if system_name=='chronos':
            w.writerow(['employee_code','request_type','request_date','start_time','end_time','status','approved_at'])
            q=s.query(WorkflowRequest,User).join(User,WorkflowRequest.applicant_id==User.id).filter(WorkflowRequest.status=='approved')
            if u.role=='admin': q=q.filter(WorkflowRequest.facility_id==u.facility_id)
            for r,usr in q.all(): w.writerow([usr.employee_code,r.request_type,r.request_date,r.start_time,r.end_time,r.status,dt_iso(r.updated_at)])
        elif system_name=='honobono':
            w.writerow(['resident_code','resident_name','facility_code','service_type'])
            q=s.query(Resident,Facility).join(Facility,Resident.facility_id==Facility.id)
            if u.role=='admin': q=q.filter(Resident.facility_id==u.facility_id)
            for r,f in q.all(): w.writerow([r.resident_code,r.name,f.code,r.service_type])
        elif system_name=='yayoi':
            w.writerow(['invoice_no','invoice_date','facility','total_amount','match_status'])
            q=s.query(Invoice,Facility).join(Facility,Invoice.facility_id==Facility.id)
            if u.role=='admin': q=q.filter(Invoice.facility_id==u.facility_id)
            for inv,f in q.all(): w.writerow([inv.invoice_no,inv.invoice_date,f.name,inv.total_amount,inv.match_status])
        else: raise HTTPException(404,'未対応の連携先です')
        data='\ufeff'+out.getvalue(); return Response(content=data.encode('utf-8'),media_type='text/csv; charset=utf-8',headers={'Content-Disposition':f'attachment; filename="ittokai_{system_name}_export.csv"'})

@app.get('/api/audit')
def audit_list(request:Request,module:Optional[str]=None):
    with SessionLocal() as s:
        u=current_user(request,s); require_roles(u,'admin','hq'); q=s.query(AuditLog)
        if u.role=='admin': q=q.filter(AuditLog.facility_id==u.facility_id)
        if module: q=q.filter(AuditLog.module==module)
        return {'items':[{'id':x.id,'actor_name':x.actor_name,'module':x.module,'action':x.action,'entity_type':x.entity_type,'entity_id':x.entity_id,'metadata':json.loads(x.metadata_json or '{}'),'ip':x.ip,'created_at':dt_iso(x.created_at)} for x in q.order_by(AuditLog.id.desc()).limit(300)]}

# Static / PWA
app.mount('/assets', StaticFiles(directory=os.path.join(STATIC,'assets')), name='assets')
@app.get('/manifest.webmanifest')
def manifest(): return FileResponse(os.path.join(STATIC,'manifest.webmanifest'),media_type='application/manifest+json')
@app.get('/sw.js')
def sw(): return FileResponse(os.path.join(STATIC,'sw.js'),media_type='application/javascript')
@app.get('/privacy')
def privacy(): return FileResponse(os.path.join(STATIC,'privacy.html'))
@app.get('/terms')
def terms(): return FileResponse(os.path.join(STATIC,'terms.html'))
@app.get('/{path:path}')
def spa(path:str):
    fp=os.path.join(STATIC,path)
    if path and os.path.isfile(fp): return FileResponse(fp)
    return FileResponse(os.path.join(STATIC,'index.html'))

if __name__=='__main__':
    import uvicorn
    port=int(os.environ.get('PORT','8000'))
    uvicorn.run(app,host='0.0.0.0',port=port)
