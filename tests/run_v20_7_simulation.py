import os, sys, json, base64, io, importlib
from pathlib import Path
from datetime import datetime, timedelta
from PIL import Image

ROOT=Path(__file__).resolve().parents[1]
DB=ROOT/'data'/'ittokai_v20_7_test.db'
if DB.exists(): DB.unlink()
os.environ['DATABASE_URL']=f'sqlite:///{DB}'
sys.path.insert(0,str(ROOT))
import server
from fastapi.testclient import TestClient

client=TestClient(server.app)
accounts={
 'family':('family@ittokai.local','demo1234'),
 'staff':('staff@ittokai.local','staff1234'),
 'middle_manager':('manager@ittokai.local','manager1234'),
 'admin':('admin@ittokai.local','admin1234'),
 'hq':('hq@ittokai.local','hq1234'),
 'caredo':('caredo@ittokai.local','caredo1234'),
 'sponsor':('sponsor@ittokai.local','sponsor1234'),
}
tokens={}
checks=0
failures=[]
scenarios=0

def ok(cond,msg):
 global checks
 checks+=1
 if not cond: raise AssertionError(msg)

def req(role,method,path,**kwargs):
 h=kwargs.pop('headers',{})
 h={**h,'Authorization':'Bearer '+tokens[role]}
 r=client.request(method,path,headers=h,**kwargs)
 return r

def expect(role,method,path,status=200,**kwargs):
 r=req(role,method,path,**kwargs)
 ok(r.status_code==status,f'{role} {method} {path}: expected {status}, got {r.status_code}: {r.text[:300]}')
 return r

# Login all seven role IDs and bootstrap
for role,(email,pw) in accounts.items():
 r=client.post('/api/login',json={'email':email,'password':pw})
 ok(r.status_code==200,f'login failed {role}: {r.text}')
 j=r.json(); tokens[role]=j['token']; ok(j['user']['role']==role,f'role mismatch {role}')
 b=expect(role,'GET','/api/bootstrap').json(); ok(b['user']['role']==role,f'bootstrap role mismatch {role}')
scenarios += 7

# Role UI/backend data expectations
bf=expect('family','GET','/api/bootstrap').json(); family_id=bf['user']['id']; resident_id=bf['residents'][0]['id']; facility_id=bf['user']['facility_id']
ba=expect('admin','GET','/api/bootstrap').json(); ok(len(ba.get('family_users',[]))>=1,'admin family list missing')
bs=expect('staff','GET','/api/bootstrap').json(); ok(len(bs.get('family_users',[]))>=1,'staff family list missing')
bm=expect('middle_manager','GET','/api/bootstrap').json(); ok(len(bm.get('family_users',[]))>=1,'middle manager family list missing')

# Chat permissions and round-trip. 120 operational cycles.
for i in range(20):
 body=f'家族テストメッセージ {i:03d}'
 expect('family','POST','/api/messages',json={'body':body})
 g=expect('admin','GET',f'/api/messages?user_id={family_id}').json(); ok(g['messages'][-1]['body']==body,f'admin did not receive family message {i}')
 reply=f'施設返信 {i:03d}'
 expect('admin','POST','/api/messages',json={'body':reply,'user_id':family_id})
 gf=expect('family','GET','/api/messages').json(); ok(gf['messages'][-1]['body']==reply,f'family did not receive admin reply {i}')
 rev=expect('family','GET','/api/messages/revision').json()['revision']; ok(bool(rev),f'chat revision missing {i}')
 scenarios += 1

# Staff and middle manager can use facility-family chat
expect('staff','GET',f'/api/messages?user_id={family_id}')
expect('middle_manager','GET',f'/api/messages?user_id={family_id}')
# CareDo / sponsor / HQ cannot open family chat
for role in ('caredo','sponsor','hq'):
 expect(role,'GET',f'/api/messages?user_id={family_id}',status=403)
 scenarios += 1

# Document cycles: create -> family view -> sign -> admin sees signed -> PDF
img=Image.new('RGB',(120,40),'white'); bio=io.BytesIO(); img.save(bio,format='PNG'); sig='data:image/png;base64,'+base64.b64encode(bio.getvalue()).decode()
for i in range(5):
 r=expect('admin','POST','/api/documents',json={'title':f'契約テスト{i}','category':'同意書','body':'テスト本文','recipient_user_ids':[family_id],'due_days':7,'requires_signature':True,'required_checks':['確認']})
 did=r.json()['id']
 expect('family','POST',f'/api/documents/{did}/view')
 expect('family','POST',f'/api/documents/{did}/sign',json={'signer_name':'相原 良太','signature_data':sig,'consent':True})
 famdocs=expect('family','GET','/api/bootstrap').json()['documents']; ok(any(d['id']==did and d['status']=='signed' for d in famdocs),f'family signed status missing {did}')
 admindocs=expect('admin','GET','/api/bootstrap').json()['documents']; ok(any(d['id']==did and d['status']=='signed' for d in admindocs),f'admin signed status missing {did}')
 pdf=expect('family','GET',f'/api/documents/{did}/pdf'); ok(pdf.content.startswith(b'%PDF'),f'bad signed pdf {did}')
 scenarios += 1

# Reject fake PDF upload
files={'file':('not.pdf',b'not a pdf','application/pdf')}
data={'title':'fake','category':'契約書','recipient_user_ids':str(family_id),'due_days':'7'}
r=req('admin','POST','/api/documents/upload',files=files,data=data); ok(r.status_code==400,'fake PDF was accepted')

# Workflow cycles: staff -> middle manager -> admin
for i in range(4):
 r=expect('staff','POST','/api/workflows',json={'request_type':'有給申請','request_date':(datetime.now()+timedelta(days=10+i)).date().isoformat(),'start_time':'','end_time':'','payload':{},'comment':f'テスト{i}'})
 wid=r.json()['id']; ok(r.json()['status']=='pending_manager','workflow initial status')
 r=expect('middle_manager','POST',f'/api/workflows/{wid}/action',json={'action':'approve','comment':''}); ok(r.json()['status']=='pending_admin','manager transition failed')
 r=expect('admin','POST',f'/api/workflows/{wid}/action',json={'action':'approve','comment':''}); ok(r.json()['status']=='approved','admin transition failed')
 # repeat action must be blocked
 expect('admin','POST',f'/api/workflows/{wid}/action',status=400,json={'action':'approve','comment':''})
 scenarios += 1

# Visit cycles: availability -> reserve -> cancel
future=(datetime.now()+timedelta(days=2)).date().isoformat()
for i in range(5):
 av=expect('family','GET',f'/api/visits/availability?facility_id={facility_id}&visit_date={future}').json()['slots']
 slot=next((x for x in av if x['remaining']>0),None); ok(slot is not None,'no visit slot available')
 r=expect('family','POST','/api/visits',json={'resident_id':resident_id,'visit_date':future,'start_time':slot['start_time'],'end_time':'23:59','place_name':slot['place_name'],'visitor_count':1,'representative':'テスト家族','phone':'','notes':''})
 vid=r.json()['id']; expect('family','POST',f'/api/visits/{vid}/cancel')
 scenarios += 1
# Sponsor cannot cancel a visit ID even if known
av=expect('family','GET',f'/api/visits/availability?facility_id={facility_id}&visit_date={future}').json()['slots']; slot=next(x for x in av if x['remaining']>0)
vid=expect('family','POST','/api/visits',json={'resident_id':resident_id,'visit_date':future,'start_time':slot['start_time'],'end_time':slot['end_time'],'place_name':slot['place_name'],'visitor_count':1,'representative':'権限テスト','phone':'','notes':''}).json()['id']
expect('sponsor','POST',f'/api/visits/{vid}/cancel',status=403)
expect('family','POST',f'/api/visits/{vid}/cancel')

# Procurement cycles: admin order -> CareDo confirm/ship -> admin receive -> CareDo invoice -> PDF
for i in range(5):
 snap=expect('admin','GET','/api/bootstrap').json()['procurement']
 product=snap['inventory'][0]; qty=max(1,int(product.get('target_stock',8)-product.get('on_hand',0)+1))
 # ensure minimum is not an issue; seed first product min 1
 r=expect('admin','POST','/api/orders',json={'items':[{'product_id':product['product_id'],'qty':qty}]}); oid=r.json()['id']
 expect('caredo','POST',f'/api/orders/{oid}/status',json={'status':'confirmed'})
 expect('caredo','POST',f'/api/orders/{oid}/status',json={'status':'shipped'})
 care=expect('caredo','GET','/api/bootstrap').json()['caredo']; order=next(o for o in care['orders'] if o['id']==oid)
 recv=[{'item_id':x['id'],'received_qty':x['qty']} for x in order['items']]
 rr=expect('admin','POST',f'/api/orders/{oid}/receive',json={'items':recv}); ok(rr.json()['status']=='received','receive not complete')
 inv_no=f'TEST-{i:03d}-{oid}'
 items=[{'product_id':x['product_id'],'qty':x['qty'],'unit_price':x['unit_price']} for x in order['items']]
 inv=expect('caredo','POST','/api/invoices',json={'order_id':oid,'invoice_no':inv_no,'invoice_date':datetime.now().date().isoformat(),'items':items}).json(); ok(inv['match_status']=='matched','3-way match failed')
 pdf=expect('admin','GET',f"/api/invoices/{inv['id']}/pdf"); ok(pdf.content.startswith(b'%PDF'),'invoice PDF invalid')
 scenarios += 1

# News flow
nid=expect('admin','POST','/api/news',json={'category':'テスト','title':'全ID運用テスト','body':'通知確認','important':False,'image_url':''}).json()['id']
news=expect('family','GET','/api/bootstrap').json()['news']; ok(any(n['id']==nid for n in news),'news not visible to family'); scenarios += 1

# HQ integrations / exports
for sysname in ('honobono','chronos','yayoi'):
 r=expect('hq','GET',f'/api/integrations/export/{sysname}'); ok('text/csv' in r.headers.get('content-type','') or r.status_code==200,f'export failed {sysname}')
scenarios += 3

# Cross-facility security regression: create second facility and admin/document directly.
with server.SessionLocal() as s:
 f2=server.Facility(code='F2',name='第二施設'); s.add(f2); s.flush()
 a2=server.User(name='第二管理者',email='admin2@test.local',password_hash=server.hash_password('x'),role='admin',facility_id=f2.id); fam2=server.User(name='第二家族',email='fam2@test.local',password_hash=server.hash_password('x'),role='family',facility_id=f2.id); s.add_all([a2,fam2]); s.flush()
 d2=server.Document(facility_id=f2.id,title='第二施設秘密書類',category='契約',body='secret',status='sent',created_by=a2.id,sent_at=server.now_dt()); d2.original_hash=server.sha256_text('secret'); s.add(d2); s.flush(); s.add(server.DocumentRecipient(document_id=d2.id,user_id=fam2.id,status='sent')); s.commit(); d2id=d2.id
expect('admin','GET',f'/api/documents/{d2id}/pdf',status=403); scenarios += 1

# v20.7 regression: contract card links resident and survives reissue
r=expect('admin','POST','/api/documents',json={'title':'重要事項説明書2026','category':'重要事項説明書','body':'契約カルテ検証','recipient_user_ids':[family_id],'resident_id':resident_id,'due_days':7,'requires_signature':True,'required_checks':['確認']})
did=r.json()['id']
bf2=expect('family','GET','/api/bootstrap').json(); cards=bf2.get('contract_cards',[]); ok(any(c['resident_id']==resident_id and any(d['id']==did for d in c['documents']) for c in cards),'contract card link missing')
r2=expect('admin','POST',f'/api/documents/{did}/reissue',json={'title':'重要事項説明書2026改訂','body':'改訂','due_days':7}); did2=r2.json()['id']
ba2=expect('admin','GET','/api/bootstrap').json(); cards=ba2.get('contract_cards',[]); card=next(c for c in cards if c['resident_id']==resident_id); ok(any(d['id']==did2 for d in card['documents']),'reissue did not keep resident link'); scenarios+=1

# v20.7 regression: product image upload/get
care=expect('caredo','GET','/api/bootstrap').json()['caredo']; pid=care['products'][0]['id']
img=Image.new('RGB',(64,64),'orange'); ib=io.BytesIO(); img.save(ib,format='PNG')
r=expect('caredo','POST',f'/api/products/{pid}/image',files={'file':('product.png',ib.getvalue(),'image/png')}); ok(r.json()['ok'],'product image upload failed')
r=client.get(f'/api/products/{pid}/image'); ok(r.status_code==200 and r.headers['content-type'].startswith('image/'),'product image get failed'); scenarios+=1

# v20.7 regression: previous/current grant lots, expiration prediction and old-lot-first consumption
bs2=expect('staff','GET','/api/bootstrap').json(); led=bs2['labor_summary']['leave_ledger']; ok(led['previous_grant'] is not None and led['current_grant'] is not None,'leave grant lots missing'); before_prev=led['previous_grant']['remaining_days']; before_cur=led['current_grant']['remaining_days']; ok(led['previous_grant']['expires_on']<'2100-01-01','leave expiry missing')
r=expect('staff','POST','/api/workflows',json={'request_type':'有給申請','request_date':(datetime.now()+timedelta(days=20)).date().isoformat(),'start_time':'','end_time':'','payload':{'leave_days':1},'comment':'残数検証'}); wid=r.json()['id']; expect('middle_manager','POST',f'/api/workflows/{wid}/action',json={'action':'approve','comment':''}); expect('admin','POST',f'/api/workflows/{wid}/action',json={'action':'approve','comment':''})
led2=expect('staff','GET','/api/bootstrap').json()['labor_summary']['leave_ledger']; ok(led2['previous_grant']['remaining_days']==round(before_prev-1,1) or (before_prev<1 and led2['current_grant']['remaining_days']<before_cur),'leave was not consumed oldest-first'); scenarios+=1

# v20.7 regression: Chronos actual import and legal alerts
staff_code=bs2['user']['employee_code']; today=datetime.now().date(); rows='employee_code,work_date,overtime_minutes,holiday_minutes\n'+f'{staff_code},{today.isoformat()},5400,600\n'
r=expect('admin','POST','/api/import/chronos_actual',files={'file':('chronos.csv',rows.encode('utf-8'),'text/csv')}); ok(r.json()['count']==1,'chronos actual import failed')
lab=expect('staff','GET','/api/bootstrap').json()['labor_summary']; ok(lab['labor_source'].startswith('クロノス実績'),'chronos source not used'); ok(lab['overtime_plus_holiday_hours']>=100,'combined hours calculation failed'); ok(any('100時間' in a['title'] for a in lab['alerts']),'100h alert missing'); scenarios+=1

# v20.7 regression: labor exports
for path in ('/api/labor/leave-ledger.csv?scope=mine','/api/labor/leave-ledger.csv?scope=facility','/api/labor/compliance.csv'):
 role='staff' if 'scope=mine' in path else 'admin'; rr=expect(role,'GET',path); ok('text/csv' in rr.headers.get('content-type',''),'labor csv missing'); scenarios+=1

# v20.7: CareDo editable product master with short message / margins
care6=expect('caredo','GET','/api/bootstrap').json()['caredo']; p6=care6['products'][0]
r=expect('caredo','PATCH',f"/api/products/{p6['id']}",json={'short_message':'毎日使うものだから、施設へまとめてお届け。','purchase_price':1000,'sale_price':1500,'active':1})
ok(r.json()['product']['short_message'].startswith('毎日使う'),'short message update failed'); ok(abs(r.json()['product']['gross_margin_rate']-33.3)<0.2,'gross margin failed'); scenarios+=1

# v20.7: Excel product master export -> preview import -> commit import
rx=expect('caredo','GET','/api/caredo/products.xlsx'); ok(len(rx.content)>1000,'product xlsx export empty')
from openpyxl import load_workbook
wb=load_workbook(io.BytesIO(rx.content)); ws=wb.active; headers=[c.value for c in ws[1]]; col={h:i+1 for i,h in enumerate(headers)}

row_target=next(r for r in range(2,ws.max_row+1) if ws.cell(r,col['sku']).value==p6['sku'])
ws.cell(row_target,col['short_message']).value='Excel更新キャッチ'; bio=io.BytesIO(); wb.save(bio); raw=bio.getvalue()
pre=expect('caredo','POST','/api/caredo/products/import-xlsx',files={'file':('products.xlsx',raw,'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')}).json(); ok(pre['update']>=1 and not pre['errors'],'product xlsx preview failed')
com=expect('caredo','POST','/api/caredo/products/import-xlsx?commit=true',files={'file':('products.xlsx',raw,'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')}).json(); ok(com['commit'] is True,'product xlsx commit failed')
care6=expect('caredo','GET','/api/bootstrap').json()['caredo']; ok(any(x['id']==p6['id'] and x['short_message']=='Excel更新キャッチ' for x in care6['products']),'xlsx change not applied'); scenarios+=1

# v20.7: facility stock policy + approval auto-order + duplicate prevention
admin6=expect('admin','GET','/api/bootstrap').json(); inv6=admin6['procurement']['inventory'][0]; fid6=admin6['user']['facility_id']; pid6=inv6['product_id']
# ensure no old open PO for this product from earlier tests by using a newly created product
newp=expect('caredo','POST','/api/products',json={'sku':'AUTO-V207','name':'自動発注検証品','category':'日用品','unit':'ケース','standard_price':1000,'purchase_price':800,'sale_price':1000,'short_message':'不足する前に自動補充','min_order_qty':1,'lead_time_days':2}).json(); pid6=newp['id']
with server.SessionLocal() as s:
    s.add(server.Inventory(facility_id=fid6,product_id=pid6,location='主倉庫',on_hand=2,safety_stock=3,reorder_point=5,target_stock=12)); s.commit()
pol=expect('caredo','POST','/api/inventory/policy',json={'facility_id':fid6,'product_id':pid6,'reorder_point':5,'target_stock':12,'reorder_qty':10,'max_stock':15,'auto_order_mode':'approval'}).json(); ok(len(pol['auto_order_events'])==1,'approval candidate not created')
# evaluate again must not duplicate same pending candidate
r2=expect('caredo','POST','/api/inventory/evaluate-auto-orders').json(); b6=expect('caredo','GET','/api/bootstrap').json()['caredo']; pend=[e for e in b6['auto_order_events'] if e['product_id']==pid6 and e['status']=='pending_approval']; ok(len(pend)==1,'duplicate pending auto-order created')
ae=pend[0]; appr=expect('caredo','POST',f"/api/auto-orders/{ae['id']}/approve").json(); ok(appr.get('order_id'),'approval did not create PO')
# third evaluation must not duplicate because incoming order exists
expect('caredo','POST','/api/inventory/evaluate-auto-orders'); b6=expect('caredo','GET','/api/bootstrap').json()['caredo']; ordered=[e for e in b6['auto_order_events'] if e['product_id']==pid6 and e['status']=='ordered']; ok(len(ordered)==1,'duplicate ordered PO created'); scenarios+=1

# v20.7: full-auto mode creates submitted order immediately
newp2=expect('caredo','POST','/api/products',json={'sku':'AUTO-V207-2','name':'完全自動発注検証品','category':'日用品','unit':'ケース','standard_price':800,'purchase_price':600,'sale_price':800,'short_message':'自動補充テスト','min_order_qty':1,'lead_time_days':2}).json(); pid62=newp2['id']
with server.SessionLocal() as s:
    s.add(server.Inventory(facility_id=fid6,product_id=pid62,location='主倉庫',on_hand=0,safety_stock=2,reorder_point=4,target_stock=10)); s.commit()
pol2=expect('caredo','POST','/api/inventory/policy',json={'facility_id':fid6,'product_id':pid62,'reorder_point':4,'target_stock':10,'reorder_qty':8,'max_stock':12,'auto_order_mode':'auto'}).json(); ok(any(x.get('order_id') for x in pol2['auto_order_events']),'full-auto PO missing'); scenarios+=1

# v20.7: facility inventory Excel export/import
ix=expect('caredo','GET','/api/caredo/facility-inventory.xlsx'); ok(len(ix.content)>1000,'inventory xlsx export empty')
wb=load_workbook(io.BytesIO(ix.content)); ws=wb.active; hs=[c.value for c in ws[1]]; ci={h:i+1 for i,h in enumerate(hs)}
# find AUTO-V207 row and change threshold to a safe value so import does not create another order
for row in range(2,ws.max_row+1):
    if ws.cell(row,ci['sku']).value=='AUTO-V207': ws.cell(row,ci['reorder_point']).value=1; ws.cell(row,ci['auto_order_mode']).value='off'
bio=io.BytesIO(); wb.save(bio); raw=bio.getvalue()
pre=expect('caredo','POST','/api/caredo/facility-inventory/import-xlsx',files={'file':('inventory.xlsx',raw,'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')}).json(); ok(not pre['errors'],'inventory xlsx preview failed')
expect('caredo','POST','/api/caredo/facility-inventory/import-xlsx?commit=true',files={'file':('inventory.xlsx',raw,'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')}); scenarios+=1

# v20.7: sponsor store/product master; sponsor is isolated from CareDo product editing
sp=expect('sponsor','GET','/api/bootstrap').json()['sponsor']; ok(len(sp['stores'])>=2 and len(sp['products'])>=3,'sponsor seed master missing')
store=expect('sponsor','POST','/api/sponsor/stores',json={'company_name':'テスト協賛','store_name':'記念日ケーキ店','category':'ケーキ屋','delivery_area':'対象施設','delivery_methods':'施設配送','message':'大切な日に甘い贈りもの','active':1}).json(); sid=store['id']
pp=expect('sponsor','POST','/api/sponsor/products',json={'store_id':sid,'sku':'GIFT-CAKE-01','name':'お祝いケーキ','short_message':'今年のお祝いを、施設へ直接お届け。','category':'記念日・ケーキ','specification':'4号','unit':'個','price':3800,'image_url':'','recommended':1,'active':1,'sale_start':'','sale_end':''}).json(); ok(pp['id']>0,'sponsor product create failed')
expect('sponsor','PATCH',f"/api/products/{p6['id']}",status=403,json={'name':'不正変更'})
spx=expect('sponsor','GET','/api/sponsor/master.xlsx'); ok(len(spx.content)>1000,'sponsor xlsx export empty'); scenarios+=1

# v20.7: family sees partner gift short messages but no sponsor private/person data exposure
fam6=expect('family','GET','/api/bootstrap').json(); gifts=fam6.get('gift_recommendations',[]); ok(any(g['name']=='お祝いケーキ' and '施設へ直接' in g['short_message'] for g in gifts),'family gift recommendation missing'); ok(all('owner_user_id' not in g for g in gifts),'private sponsor owner leaked'); scenarios+=1

# Health/version
h=client.get('/health'); ok(h.status_code==200 and h.json()['version']=='20.7','health version mismatch')

report={'version':'20.7','scenario_cycles':scenarios,'assertions':checks,'failures':failures,'result':'PASS'}
print(json.dumps(report,ensure_ascii=False,indent=2))
Path(ROOT/'TEST_RESULT_v20.7.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
