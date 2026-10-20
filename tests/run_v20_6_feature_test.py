import os, sys, io, json
from pathlib import Path
ROOT=Path(__file__).resolve().parents[1]
DB=ROOT/'data'/'ittokai_v20_6_feature_test.db'
if DB.exists(): DB.unlink()
os.environ['DATABASE_URL']=f'sqlite:///{DB}'
sys.path.insert(0,str(ROOT))
import server
from fastapi.testclient import TestClient
from openpyxl import load_workbook
c=TestClient(server.app)
accounts={
 'family':('family@ittokai.local','demo1234'),'staff':('staff@ittokai.local','staff1234'),
 'middle_manager':('manager@ittokai.local','manager1234'),'admin':('admin@ittokai.local','admin1234'),
 'hq':('hq@ittokai.local','hq1234'),'caredo':('caredo@ittokai.local','caredo1234'),
 'sponsor':('sponsor@ittokai.local','sponsor1234')}
tokens={}; checks=[]
def ck(x,msg):
 if not x: raise AssertionError(msg)
 checks.append(msg)
def req(role,method,path,expect=200,**kw):
 h=kw.pop('headers',{}); h['Authorization']='Bearer '+tokens[role]
 r=c.request(method,path,headers=h,**kw); ck(r.status_code==expect,f'{role} {method} {path}={r.status_code} expected {expect}: {r.text[:200]}'); return r
for role,(email,pw) in accounts.items():
 r=c.post('/api/login',json={'email':email,'password':pw}); ck(r.status_code==200,f'login {role}'); tokens[role]=r.json()['token']
 b=req(role,'GET','/api/bootstrap').json(); ck(b['version']=='20.6',f'bootstrap v20.6 {role}')
# CareDo master update + message/margin
care=req('caredo','GET','/api/bootstrap').json()['caredo']; ck(len(care['products'])>=3,'seed products'); p=care['products'][0]
r=req('caredo','PATCH',f"/api/products/{p['id']}",json={'short_message':'施設へ直接お届け。','purchase_price':1000,'sale_price':1500,'active':1}).json()['product']
ck(r['short_message']=='施設へ直接お届け。','short message'); ck(abs(r['gross_margin_rate']-33.3)<0.2,'margin')
# Product xlsx export/import
x=req('caredo','GET','/api/caredo/products.xlsx'); ck(len(x.content)>1000,'product xlsx export')
wb=load_workbook(io.BytesIO(x.content)); ws=wb.active; hs=[c.value for c in ws[1]]; ci={h:i+1 for i,h in enumerate(hs)}
row=next(i for i in range(2,ws.max_row+1) if ws.cell(i,ci['sku']).value==p['sku']); ws.cell(row,ci['short_message']).value='Excelで更新した一言'; bio=io.BytesIO(); wb.save(bio); raw=bio.getvalue()
pre=req('caredo','POST','/api/caredo/products/import-xlsx',files={'file':('p.xlsx',raw,'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')}).json(); ck(not pre['errors'] and pre['update']>=1,'product xlsx preview')
req('caredo','POST','/api/caredo/products/import-xlsx?commit=true',files={'file':('p.xlsx',raw,'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')})
care=req('caredo','GET','/api/bootstrap').json()['caredo']; ck(any(z['id']==p['id'] and z['short_message']=='Excelで更新した一言' for z in care['products']),'product xlsx commit')
# New auto-order item
admin=req('admin','GET','/api/bootstrap').json(); fid=admin['user']['facility_id']
np=req('caredo','POST','/api/products',json={'sku':'V206-AUTO','name':'自動発注品','category':'日用品','unit':'ケース','standard_price':1000,'purchase_price':800,'sale_price':1000,'short_message':'不足前に補充','min_order_qty':1,'lead_time_days':2}).json(); pid=np['id']
with server.SessionLocal() as s:
 s.add(server.Inventory(facility_id=fid,product_id=pid,location='主倉庫',on_hand=2,safety_stock=3,reorder_point=5,target_stock=12)); s.commit()
pol=req('caredo','POST','/api/inventory/policy',json={'facility_id':fid,'product_id':pid,'reorder_point':5,'target_stock':12,'reorder_qty':10,'max_stock':15,'auto_order_mode':'approval'}).json(); ck(len(pol['auto_order_events'])==1,'approval candidate')
req('caredo','POST','/api/inventory/evaluate-auto-orders')
care=req('caredo','GET','/api/bootstrap').json()['caredo']; pend=[e for e in care['auto_order_events'] if e['product_id']==pid and e['status']=='pending_approval']; ck(len(pend)==1,'no duplicate pending')
ap=req('caredo','POST',f"/api/auto-orders/{pend[0]['id']}/approve").json(); ck(ap.get('order_id'),'approval created PO')
req('caredo','POST','/api/inventory/evaluate-auto-orders'); care=req('caredo','GET','/api/bootstrap').json()['caredo']; ck(len([e for e in care['auto_order_events'] if e['product_id']==pid and e['status']=='ordered'])==1,'no duplicate order')
# Full auto item
np2=req('caredo','POST','/api/products',json={'sku':'V206-AUTO2','name':'完全自動品','category':'日用品','unit':'ケース','standard_price':900,'purchase_price':700,'sale_price':900,'short_message':'完全自動','min_order_qty':1,'lead_time_days':2}).json(); pid2=np2['id']
with server.SessionLocal() as s:
 s.add(server.Inventory(facility_id=fid,product_id=pid2,location='主倉庫',on_hand=0,safety_stock=2,reorder_point=4,target_stock=10)); s.commit()
pol2=req('caredo','POST','/api/inventory/policy',json={'facility_id':fid,'product_id':pid2,'reorder_point':4,'target_stock':10,'reorder_qty':8,'max_stock':12,'auto_order_mode':'auto'}).json(); ck(any(e.get('order_id') for e in pol2['auto_order_events']),'full auto PO')
# inventory xlsx
ix=req('caredo','GET','/api/caredo/facility-inventory.xlsx'); ck(len(ix.content)>1000,'inventory xlsx export')
wb=load_workbook(io.BytesIO(ix.content)); ws=wb.active; hs=[c.value for c in ws[1]]; ci={h:i+1 for i,h in enumerate(hs)}
for i in range(2,ws.max_row+1):
 if ws.cell(i,ci['sku']).value=='V206-AUTO': ws.cell(i,ci['reorder_point']).value=1; ws.cell(i,ci['auto_order_mode']).value='off'
bio=io.BytesIO(); wb.save(bio); raw=bio.getvalue(); pre=req('caredo','POST','/api/caredo/facility-inventory/import-xlsx',files={'file':('i.xlsx',raw,'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')}).json(); ck(not pre['errors'],'inventory xlsx preview'); req('caredo','POST','/api/caredo/facility-inventory/import-xlsx?commit=true',files={'file':('i.xlsx',raw,'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')})
# sponsor store/products and rights
sp=req('sponsor','GET','/api/bootstrap').json()['sponsor']; ck(len(sp['stores'])>=2 and len(sp['products'])>=3,'sponsor seed')
st=req('sponsor','POST','/api/sponsor/stores',json={'company_name':'協賛テスト','store_name':'記念日ケーキ','category':'ケーキ屋','delivery_area':'対象施設','delivery_methods':'施設配送','message':'大切な記念日に','active':1}).json(); sid=st['id']
pp=req('sponsor','POST','/api/sponsor/products',json={'store_id':sid,'sku':'GIFT-01','name':'お祝いケーキ','short_message':'施設へ直接お届けする記念日のケーキ。','description':'','category':'記念日・ケーキ','specification':'4号','unit':'個','price':3800,'image_url':'','recommended':1,'active':1,'sale_start':'','sale_end':''}).json(); ck(pp['id']>0,'sponsor product')
req('sponsor','PATCH',f"/api/products/{p['id']}",expect=403,json={'name':'unauthorized'})
sx=req('sponsor','GET','/api/sponsor/master.xlsx'); ck(len(sx.content)>1000,'sponsor xlsx export')
# family gift visibility/privacy
fam=req('family','GET','/api/bootstrap').json(); gifts=fam.get('gift_recommendations',[]); ck(any(g['name']=='お祝いケーキ' and '施設へ直接' in g['short_message'] for g in gifts),'family gift'); ck(all('owner_user_id' not in g for g in gifts),'no sponsor owner leak')
# regression: basic manual order -> confirm -> ship -> receive -> invoice
# use a seed product not the auto-order ones and choose one without incoming order
seedp=next(x for x in care['products'] if x['sku']=='CD-TISSUE')
o=req('admin','POST','/api/orders',json={'items':[{'product_id':seedp['id'],'qty':1}]}).json(); oid=o['id']; req('caredo','POST',f'/api/orders/{oid}/status',json={'status':'confirmed'}); req('caredo','POST',f'/api/orders/{oid}/status',json={'status':'shipped'})
care2=req('caredo','GET','/api/bootstrap').json()['caredo']; oo=next(x for x in care2['orders'] if x['id']==oid); req('admin','POST',f'/api/orders/{oid}/receive',json={'items':[{'item_id':it['id'],'received_qty':it['qty']} for it in oo['items']]})
inv=req('caredo','POST','/api/invoices',json={'order_id':oid,'invoice_no':'V206-INV-1','invoice_date':'2026-10-20','items':[{'product_id':it['product_id'],'qty':it['qty'],'unit_price':it['unit_price']} for it in oo['items']]}).json(); ck(inv['match_status']=='matched','3way match regression')
h=c.get('/health'); ck(h.status_code==200 and h.json()['version']=='20.6','health')
report={'version':'20.6','checks':len(checks),'result':'PASS'}
print(json.dumps(report,ensure_ascii=False,indent=2)); (ROOT/'TEST_RESULT_v20.6_FEATURES.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
